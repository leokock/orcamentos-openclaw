import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from openpyxl import load_workbook
p = r"G:\Drives compartilhados\03 CTN Projetos\2. Projetos em Andamento\_Entregas\Orçamento_executivo\Mussi Empreendimentos\Chelsea\CTN_MSS_CLS - Orçamento R02.xlsx"
wb = load_workbook(p, data_only=True, read_only=True)
print("SHEETS:", wb.sheetnames)
for nm in wb.sheetnames:
    if nm.lower().strip() in ('obra','resumo','capa','dados'):
        ws = wb[nm]
        print(f"\n== {nm} ({ws.max_row}x{ws.max_column}) ==")
        for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row,60), values_only=False):
            cells = [f"{c.coordinate}={c.value!r}" for c in r if c.value is not None]
            if cells: print(" | ".join(cells))
wb.close()

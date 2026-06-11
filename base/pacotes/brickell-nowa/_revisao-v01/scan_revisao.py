import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from openpyxl import load_workbook
p = r"G:\Drives compartilhados\03 CTN Projetos\2. Projetos em Andamento\_Parametrico_IA\REVISAO-OBRAS-131.xlsx"
wb = load_workbook(p, data_only=True, read_only=True)
print("SHEETS:", wb.sheetnames)
for wsname in wb.sheetnames:
    ws = wb[wsname]
    print(f"\n== {wsname}: {ws.max_row} rows x {ws.max_column} cols ==")
    rows = list(ws.iter_rows(min_row=1, max_row=3, values_only=True))
    for i, r in enumerate(rows, 1):
        print(f"R{i}: {r}")
wb.close()

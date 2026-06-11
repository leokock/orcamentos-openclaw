import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from openpyxl import load_workbook

p = r"G:\Drives compartilhados\03 CTN Projetos\2. Projetos em Andamento\_Parametrico_IA\REVISAO-OBRAS-131.xlsx"
wb = load_workbook(p, data_only=True, read_only=True)
ws = wb['Revisão Obras']
hdr = None
targets = {'adore-level-up','terrassa-amaro','mussi-empreendimentos-chelsea','gdi-playa-negra','neuhaus-origem'}
print("### LINHAS DOS 5 ESCOLHIDOS ###")
rows = list(ws.iter_rows(values_only=True))
hdr = rows[0]
for r in rows[1:]:
    if r[0] in targets:
        for k, v in zip(hdr, r):
            print(f"  {r[0]}.{k} = {v!r}")
        print("---")
print("\n### UNIVERSO: SC residencial vertical com R$/m² valido ###")
print("slug | cidade | padrao | tipologia | data_base | fonte_db | AC | UR | R$/m2 | total")
for r in rows[1:]:
    slug, cliente, cidade, uf, cubreg, padrao, tipo, db, de, ac, ur, rsm2, total = r[:13]
    fonte_db = r[16]
    if uf == 'SC' and rsm2 and tipo and 'residencial' in str(tipo):
        print(f"{slug} | {cidade} | {padrao} | {tipo} | {db} | {fonte_db} | {ac} | {ur} | {round(rsm2,2) if rsm2 else None} | {round(total,2) if total else None}")
print("\n### RESUMO SHEET ###")
ws2 = wb['Resumo']
for r in ws2.iter_rows(values_only=True):
    if any(v is not None for v in r):
        print(r)
wb.close()

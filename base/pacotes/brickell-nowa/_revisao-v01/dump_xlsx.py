import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from openpyxl import load_workbook

def dump(path, out):
    lines = []
    for data_only, tag in [(False,'FORMULA'),(True,'VALUE')]:
        wb = load_workbook(path, data_only=data_only)
        for ws in wb.worksheets:
            lines.append(f"\n===== SHEET [{tag}]: {ws.title} (dims {ws.dimensions}) =====")
            for row in ws.iter_rows():
                cells = []
                for c in row:
                    if c.value is not None:
                        cells.append(f"{c.coordinate}={c.value!r}")
                if cells:
                    lines.append(" | ".join(cells))
    with open(out, 'w', encoding='utf-8') as f:
        f.write("\n".join(lines))
    print(f"wrote {out}")

dump('brickell-nowa-parametrico-v00-final.xlsx', '_revisao-v01/dump-parametrico.txt')
dump('brickell-nowa-preliminar-v00.xlsx', '_revisao-v01/dump-preliminar.txt')
dump('brickell-nowa-comparativo-similares-v00.xlsx', '_revisao-v01/dump-comparativo.txt')

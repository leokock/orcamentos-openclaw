from __future__ import annotations

import csv
import json
import shutil
import subprocess
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border


PKG = Path(r"C:\Users\leona\orcamentos-openclaw\base\pacotes\flow-da-lagoa")
SOURCE = Path(
    r"G:\Drives compartilhados\03 CTN Projetos\2. Projetos em Andamento"
    r"\AMS Empreendimentos\04. Custo\04.1 Custo - Paramétrico\01 - Projetos"
)
FINAL_DIR = Path(r"C:\Users\leona\orcamentos\parametricos\flow-da-lagoa")
GENERATOR = Path(r"C:\Users\leona\orcamentos-openclaw\scripts\gerar_template_dinamico_v2.py")

RUN_DATE = "2026-06-09"
PROJECT = "Flow da Lagoa"
SLUG = "flow-da-lagoa"

PARAMS = {
    "nome": PROJECT,
    "cliente": "AMS Empreendimentos Lagoa SPE Ltda",
    "cidade": "Florianópolis",
    "estado": "SC",
    "regiao": "Grande Florianópolis",
    "ac": 15000.62,
    "area_computavel_m2": 8122.55,
    "area_privativa_m2": 8343.90,
    "ur": 236,
    "np": 6,
    "npt": 3,
    "elev": 4,
    "vag": 161,
    "vagas_bicicleta": 213,
    "prazo": 24,
    "cub": 3096.25,
    "bdi_referencia": 0.20,
    "briefing": {
        "laje": "Convencional",
        "subsolos": "0",
        "fundacao": "Hélice",
        "padrao_acabamento": "Médio-Alto",
        "fachada": "Textura",
        "pressurizacao": "Não",
        "n_torres": "3",
        "gerador": "Não",
        "entrega": "Completa",
        "tipologia": "1-2 Dormitórios",
        "pe_direito": "Padrão (3.00)",
        "n_banheiros": "1",
        "tipo_piso": "Porcelanato",
        "piscina": "Sim",
    },
}

PROJECT_EVIDENCE = [
    ("Área construída", "15.000,62 m²", "ARQ AP 008 CORTE/TABELAS R12, quadro Área Total a Construir"),
    ("Área computável", "8.122,55 m²", "ARQ AP 008 CORTE/TABELAS R12, quadro Área Total a Construir"),
    ("Área privativa", "8.343,90 m²", "ARQ AP 008 CORTE/TABELAS R12, quadro Área Privativa"),
    ("Pavimentos principais", "NA01 a NA06", "ARQ AP 008 CORTE/TABELAS R12 e cortes"),
    ("Níveis técnicos", "NA07 Barrilete, NA08 Reservatório Superior, NA09 Tampa", "ARQ AP 008 CORTE/TABELAS R12"),
    ("Vagas automóveis", "161", "Quadro de vagas: 4 comerciais + 149 privativas + 8 visitantes"),
    ("Vagas bicicleta", "213", "Quadro de vagas: 5 comerciais + 198 privativas + 10 visitantes"),
    ("Núcleos verticais", "4", "Leitura visual das plantas e cortes; elevadores a validar"),
    ("Unidades residenciais", "236", "Premissa por área privativa média de 35,36 m²/UR; quadro de unidades não encontrado"),
    ("Estrutura", "Lajes maciças/mistas e=20 cm", "IFC estrutural R06: slabs, beams e columns; sem indício de protensão"),
    ("Piscina", "Sim", "Há disciplina PISCINA com DWG/PDF/IFC"),
]

DISCIPLINE_NOTES = {
    "ARQ": "Arquitetura com IFC e pranchas PDF/DWG. A prancha de cortes/tabelas forneceu áreas, níveis e vagas. O IFC não traz IfcSpace útil para contagem automática de unidades.",
    "EST": "Estrutura com modelos RVT/IFC e pranchas. O IFC R06 possui pavimentos estruturais, lajes, vigas e pilares; usado para enquadrar laje convencional/maciça-mista.",
    "FUND": "Fundações com PDF/DWG. Sem IFC quantitativo; fundação em hélice foi mantida como premissa V1 até validar memorial/sondagem.",
    "HID": "Hidrossanitário com PDF/DWG/IFC. Usado para confirmar complexidade de instalações e muitas áreas molhadas.",
    "ELETRICA": "Elétrica com PDF/DWG/IFC. Usado para confirmar disciplina completa; gerador dedicado não identificado na leitura inicial.",
    "PCI": "Prevenção contra incêndio com PDF/DWG/IFC. Entrou em sistemas/instalações e premissa de pressurização = Não por baixa altura do conjunto.",
    "SPDA": "SPDA com PDF/DWG/IFC. Entrou em sistemas especiais.",
    "CLI": "Climatização com PDF/DWG. Entrou como climatização básica residencial/comercial.",
    "PISCINA": "Projeto específico de piscina presente; briefing V2 marcado como Piscina = Sim.",
    "PAISAGISMO": "Paisagismo com PDF/DWG. Entrou em complementares/urbanização.",
    "INTERIORES": "Interiores com PDF/DWG/IFC. Reforça entrega completa e padrão médio-alto.",
    "TEL": "Telecom com PDF/DWG. Entrou em sistemas especiais/instalações.",
    "_Compactadas": "Backups ZIP das disciplinas. Mantidos no inventário, mas não usados como fonte primária quando havia arquivos abertos.",
}

SUPABASE_MACRO = [
    ("Médio-Alto", "Alvenaria", 33, 121.14, 139.70, 167.52),
    ("Médio-Alto", "Climatização", 9, 29.37, 30.59, 68.20),
    ("Médio-Alto", "Complementares", 27, 129.36, 201.06, 293.93),
    ("Médio-Alto", "Esquadrias", 35, 256.91, 301.24, 475.17),
    ("Médio-Alto", "Fachada", 20, 84.26, 117.38, 166.19),
    ("Médio-Alto", "Gerenciamento", 36, 339.65, 474.74, 1223.46),
    ("Médio-Alto", "Impermeabilização", 33, 41.71, 60.00, 93.23),
    ("Médio-Alto", "Imprevistos", 15, 44.33, 49.98, 72.41),
    ("Médio-Alto", "Infraestrutura", 37, 141.55, 196.14, 240.53),
    ("Médio-Alto", "Instalações", 33, 281.74, 332.12, 411.27),
    ("Médio-Alto", "Louças e Metais", 24, 39.60, 59.73, 147.71),
    ("Médio-Alto", "Movimentação de Terra", 37, 82.18, 114.52, 147.52),
    ("Médio-Alto", "Pintura", 26, 103.60, 124.55, 148.22),
    ("Médio-Alto", "Pisos", 24, 150.22, 199.02, 233.50),
    ("Médio-Alto", "Rev. Interno Parede", 15, 45.05, 60.62, 362.14),
    ("Médio-Alto", "Sistemas Especiais", 35, 122.38, 172.54, 242.80),
    ("Médio-Alto", "Supraestrutura", 36, 461.10, 654.62, 761.97),
    ("Médio-Alto", "Teto", 21, 43.28, 60.10, 69.40),
    ("Alto", "Alvenaria", 21, 128.63, 181.90, 219.79),
    ("Alto", "Climatização", 9, 28.24, 36.21, 78.09),
    ("Alto", "Complementares", 15, 153.72, 262.06, 365.34),
    ("Alto", "Esquadrias", 23, 285.86, 395.49, 553.96),
    ("Alto", "Fachada", 10, 142.31, 206.67, 252.97),
    ("Alto", "Gerenciamento", 22, 535.79, 640.41, 1715.95),
    ("Alto", "Impermeabilização", 20, 61.39, 74.35, 128.04),
    ("Alto", "Imprevistos", 7, 48.03, 72.24, 75.44),
    ("Alto", "Infraestrutura", 23, 156.80, 230.40, 276.00),
    ("Alto", "Instalações", 21, 185.17, 320.13, 453.78),
    ("Alto", "Louças e Metais", 17, 56.79, 141.98, 222.16),
    ("Alto", "Movimentação de Terra", 23, 91.90, 124.19, 178.15),
    ("Alto", "Pintura", 16, 114.05, 159.16, 252.13),
    ("Alto", "Pisos", 16, 187.44, 237.05, 638.25),
    ("Alto", "Rev. Interno Parede", 12, 59.02, 82.02, 212.73),
    ("Alto", "Sistemas Especiais", 21, 126.25, 211.02, 304.34),
    ("Alto", "Supraestrutura", 23, 659.32, 718.24, 854.66),
    ("Alto", "Teto", 12, 52.54, 62.23, 85.03),
]

SUPABASE_PUS = [
    ("Concreto usinado bombeável fck 40 MPa", "m³", 37, 136, 444.15, 477.75, 551.25),
    ("Concreto usinado 40 MPa para estacas", "m³", 16, 51, 502.00, 517.65, 565.00),
    ("Esquadrias de alumínio - G2", "m²", 77, 40, 824.00, 939.74, 1049.07),
    ("Fornecimento e instalação de esquadrias de alumínio", "m²", 39, 29, 1050.00, 1088.02, 1250.00),
    ("Elevadores", "un", 43, 10, 197190.00, 261192.12, 326782.50),
    ("Elevador social", "un", 27, 9, 107927.75, 220000.00, 405000.00),
    ("TOTAL IMPERMEABILIZAÇÃO", "R$/m² AC", 88, 111, 47.92, 61.39, 87.54),
    ("Mão de obra porcelanato piso", "m²", 17, 40, 42.00, 48.00, 48.00),
    ("Pintura latex acrílica paredes internas", "m²", 32, 37, 6.12, 6.12, 11.23),
]

AMPLI_REFERENCES = [
    ("Elevador", "un", 10, 12, 155659.15, 289254.38, 445000.00),
    ("Elevador social", "un", 10, 10, 213125.00, 626385.81, 1400000.00),
    ("Impermeabilização argamassa polimérica - MO", "m²", 9, 39, 16.50, 27.31, 50.00),
    ("Impermeabilização manta asfáltica - MO", "m²", 3, 12, 48.00, 48.00, 48.00),
    ("Pintura latex acrílica paredes internas 3 demãos - MO", "m²", 5, 33, 12.65, 20.00, 35.00),
    ("Pintura latex acrílica teto 3 demãos - MO", "m²", 5, 33, 12.65, 17.46, 39.50),
    ("Porcelanato 60x60", "m²", 0, 0, 0.00, 80.49, 0.00),
    ("Porcelanato 80x80", "m²", 0, 0, 0.00, 129.22, 0.00),
    ("Esquadria alumínio/vidro", "m²", 0, 0, 958.00, 1088.00, 1250.00),
]

COMPARABLES = [
    ("lumis-live", "medio-alto", 14888.46, 0, 43772261.00, 2940.01),
    ("mussi-empreendimentos-chelsea", "medio-alto", 15149.39, 0, 52163486.00, 3443.27),
    ("neuhaus-origem", "alto", 14559.12, 0, 67590131.00, 4642.46),
    ("fonseca-empreendimentos-estoril", "medio-alto", 14491.98, 110, 48960377.00, 3378.45),
    ("pavcor", "medio-alto", 14283.29, 0, 49360279.00, 3455.81),
    ("pass-e-connect", "medio-alto", 13144.27, 0, 48758956.00, 3709.52),
]


def money(value: float) -> str:
    return f"R$ {value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def number(value: float) -> str:
    return f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def inventory_files() -> tuple[list[dict], dict[str, Counter], dict[str, int]]:
    rows: list[dict] = []
    by_disc: dict[str, Counter] = defaultdict(Counter)
    size_by_disc: dict[str, int] = defaultdict(int)
    for path in sorted(SOURCE.rglob("*")):
        if not path.is_file():
            continue
        rel = path.relative_to(SOURCE)
        disc = rel.parts[0] if rel.parts else "_root"
        ext = path.suffix.lower() or "(sem extensão)"
        stat = path.stat()
        rows.append(
            {
                "disciplina": disc,
                "extensao": ext,
                "tamanho_mb": round(stat.st_size / 1024 / 1024, 3),
                "modificado_em": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
                "arquivo_relativo": str(rel),
            }
        )
        by_disc[disc][ext] += 1
        size_by_disc[disc] += stat.st_size
    return rows, by_disc, size_by_disc


def write_inventory(rows: list[dict], by_disc: dict[str, Counter], size_by_disc: dict[str, int]) -> None:
    csv_path = PKG / "inventario-arquivos-flow-da-lagoa.csv"
    with csv_path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(
            fh,
            fieldnames=["disciplina", "extensao", "tamanho_mb", "modificado_em", "arquivo_relativo"],
        )
        writer.writeheader()
        writer.writerows(rows)

    summary = []
    for disc in sorted(by_disc):
        summary.append(
            {
                "disciplina": disc,
                "arquivos": sum(by_disc[disc].values()),
                "tamanho_mb": round(size_by_disc[disc] / 1024 / 1024, 2),
                "extensoes": dict(sorted(by_disc[disc].items())),
                "nota": DISCIPLINE_NOTES.get(disc, ""),
            }
        )
    (PKG / "inventario-resumo-flow-da-lagoa.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def write_config() -> Path:
    config = {
        **PARAMS,
        "_slug": SLUG,
        "_gerado_em": RUN_DATE,
        "_source_path": str(SOURCE),
        "_metodo": "Paramétrico V2 Híbrido + calibradores Supabase indices-cartesian + referências Ampli.",
        "_premissas_criticas": [
            "Contagem de unidades não estava em quadro legível; UR=236 por área privativa média de 35,36 m²/UR.",
            "Vagas de automóvel = 161; bicicletas separadas do parâmetro de vagas.",
            "Sem subsolo no enquadramento V1; NA01 é térreo/pilotis com estacionamento.",
            "Fachada em textura e fundação hélice são premissas V1 a validar com memoriais.",
            "Nº Torres do template limitado a 3; projeto possui 4 blocos/núcleos, registrado nas notas.",
        ],
    }
    path = PKG / "parametrico-v2-config.json"
    path.write_text(json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def run_generator(config_path: Path) -> Path:
    output = PKG / f"parametrico-{SLUG}-v01.xlsx"
    cmd = ["python", str(GENERATOR), "--config", str(config_path), "-o", str(output)]
    subprocess.run(cmd, check=True)
    shutil.copy2(output, PKG / f"parametrico-{SLUG}.xlsx")
    return output


def add_workbook_context(xlsx: Path) -> dict:
    wb = load_workbook(xlsx)

    # Force recalculation when opened in Excel/LibreOffice.
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except Exception:
        pass

    for sheet in ["FLOW_RESUMO", "FLOW_FONTES", "BDI_PREMISSAS"]:
        if sheet in wb.sheetnames:
            del wb[sheet]

    header_fill = PatternFill("solid", fgColor="245AE4")
    orange_fill = PatternFill("solid", fgColor="FD3400")
    light_fill = PatternFill("solid", fgColor="F4F4F4")
    white_font = Font(color="FFFFFF", bold=True)
    bold = Font(bold=True)
    thin = Side(style="thin", color="E7E7E9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ac = PARAMS["ac"]
    ur = PARAMS["ur"]
    cub = PARAMS["cub"]
    medio = sum(r[4] for r in SUPABASE_MACRO if r[0] == "Médio-Alto")
    alto = sum(r[4] for r in SUPABASE_MACRO if r[0] == "Alto")
    recommended_rsm2 = 3680.00
    recommended_total = recommended_rsm2 * ac

    ws = wb.create_sheet("FLOW_RESUMO", 0)
    ws["A1"] = "FLOW DA LAGOA - RESUMO DE ENQUADRAMENTO"
    ws["A1"].font = Font(bold=True, size=14, color="111111")
    ws.merge_cells("A1:D1")
    rows = [
        ("Data da versão", RUN_DATE, "Fonte", "Pacote v01"),
        ("Área construída", ac, "m²", "Quadro ARQ"),
        ("Área computável", PARAMS["area_computavel_m2"], "m²", "Quadro ARQ"),
        ("Área privativa", PARAMS["area_privativa_m2"], "m²", "Quadro ARQ"),
        ("UR", ur, "un", "Premissa por área média"),
        ("Área construída/UR", ac / ur, "m²/UR", "Controle"),
        ("Área privativa/UR", PARAMS["area_privativa_m2"] / ur, "m²/UR", "Controle"),
        ("Vagas automóveis", PARAMS["vag"], "un", "Quadro ARQ"),
        ("Vagas bicicletas", PARAMS["vagas_bicicleta"], "un", "Quadro ARQ"),
        ("Elevadores", PARAMS["elev"], "un", "Premissa visual"),
        ("Prazo", PARAMS["prazo"], "meses", "Premissa V1"),
        ("CUB-SC ref. maio/2026", cub, "R$/m²", "Sinduscon-Fpolis"),
        ("Supabase médio-alto mediana", medio, "R$/m²", "Soma macrogrupos"),
        ("Supabase alto mediana", alto, "R$/m²", "Soma macrogrupos"),
        ("Cenário recomendado v01", recommended_rsm2, "R$/m²", "Entre médio-alto e alto"),
        ("Custo direto recomendado v01", recommended_total, "R$", "Controle de plausibilidade"),
        ("R$/UR recomendado v01", recommended_total / ur, "R$/UR", "Controle"),
        ("R$/m² / CUB", recommended_rsm2 / cub, "x CUB", "Controle"),
    ]
    for idx, row in enumerate(rows, start=3):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(idx, col, value)
            cell.border = border
            if col in (1, 3):
                cell.fill = light_fill
                cell.font = bold
            if isinstance(value, float):
                cell.number_format = '#,##0.00'
    for col, width in {"A": 28, "B": 18, "C": 16, "D": 34}.items():
        ws.column_dimensions[col].width = width

    wsb = wb.create_sheet("BDI_PREMISSAS")
    wsb["A1"] = "BDI / BIFES - CENÁRIOS PARA VALIDAÇÃO"
    wsb["A1"].font = Font(bold=True, size=13, color="111111")
    wsb.merge_cells("A1:E1")
    headers = ["Cenário", "BDI", "Custo direto", "Preço com BDI", "Observação"]
    for col, header in enumerate(headers, start=1):
        cell = wsb.cell(3, col, header)
        cell.fill = header_fill
        cell.font = white_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center")
    scenarios = [
        ("Sem BDI", 0.00, recommended_total, "uso para custo direto/obra"),
        ("BDI 15%", 0.15, recommended_total * 1.15, "referência enxuta"),
        ("BDI 20%", 0.20, recommended_total * 1.20, "referência intermediária"),
        ("BDI 25%", 0.25, recommended_total * 1.25, "referência conservadora"),
    ]
    for row_idx, (label, bdi, price, note) in enumerate(scenarios, start=4):
        vals = [label, bdi, recommended_total, price, note]
        for col, value in enumerate(vals, start=1):
            cell = wsb.cell(row_idx, col, value)
            cell.border = border
            if col == 2:
                cell.number_format = "0.0%"
            if col in (3, 4):
                cell.number_format = 'R$ #,##0.00'
    for col, width in {"A": 18, "B": 12, "C": 18, "D": 18, "E": 42}.items():
        wsb.column_dimensions[col].width = width

    wsf = wb.create_sheet("FLOW_FONTES")
    wsf["A1"] = "FONTES E BENCHMARKS USADOS NO V01"
    wsf["A1"].font = Font(bold=True, size=13, color="111111")
    wsf.merge_cells("A1:H1")

    def section(title: str, row: int, headers: list[str]) -> int:
        wsf.cell(row, 1, title)
        wsf.cell(row, 1).fill = orange_fill
        wsf.cell(row, 1).font = white_font
        wsf.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        row += 1
        for col, header in enumerate(headers, start=1):
            cell = wsf.cell(row, col, header)
            cell.fill = header_fill
            cell.font = white_font
            cell.border = border
        return row + 1

    row = section("Dados extraídos do projeto", 3, ["Campo", "Valor", "Fonte"])
    for item in PROJECT_EVIDENCE:
        for col, value in enumerate(item, start=1):
            cell = wsf.cell(row, col, value)
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    row = section("Calibração Supabase por macrogrupo", row + 2, ["Padrão", "Macrogrupo", "N", "P25", "Mediana", "P75"])
    for item in SUPABASE_MACRO:
        for col, value in enumerate(item, start=1):
            cell = wsf.cell(row, col, value)
            cell.border = border
            if col >= 4:
                cell.number_format = '#,##0.00'
        row += 1

    row = section("PUs Supabase / Ampli", row + 2, ["Fonte", "Descrição", "Unidade", "Projetos", "Obs", "P25/Min", "Med/Avg", "P75/Max"])
    for item in SUPABASE_PUS:
        vals = ("Supabase",) + item
        for col, value in enumerate(vals, start=1):
            cell = wsf.cell(row, col, value)
            cell.border = border
            if col >= 6:
                cell.number_format = '#,##0.00'
        row += 1
    for item in AMPLI_REFERENCES:
        vals = ("Ampli",) + item
        for col, value in enumerate(vals, start=1):
            cell = wsf.cell(row, col, value)
            cell.border = border
            if col >= 6:
                cell.number_format = '#,##0.00'
        row += 1

    row = section("Comparáveis Supabase por área", row + 2, ["Slug", "Padrão", "AC", "UR", "Total", "R$/m²"])
    for item in COMPARABLES:
        for col, value in enumerate(item, start=1):
            cell = wsf.cell(row, col, value)
            cell.border = border
            if col in (3, 5, 6):
                cell.number_format = '#,##0.00'
        row += 1

    for col, width in {"A": 18, "B": 42, "C": 14, "D": 12, "E": 14, "F": 14, "G": 14, "H": 14}.items():
        wsf.column_dimensions[col].width = width

    wb.save(xlsx)
    shutil.copy2(xlsx, PKG / f"parametrico-{SLUG}.xlsx")
    return {
        "medio_alto_rsm2": medio,
        "alto_rsm2": alto,
        "recommended_rsm2": recommended_rsm2,
        "recommended_total": recommended_total,
    }


def md_table(rows: list[tuple]) -> str:
    if not rows:
        return ""
    cols = len(rows[0])
    header = "| " + " | ".join(str(x) for x in rows[0]) + " |"
    sep = "| " + " | ".join(["---"] * cols) + " |"
    body = ["| " + " | ".join(str(x) for x in row) + " |" for row in rows[1:]]
    return "\n".join([header, sep, *body])


def write_docs(inv_summary: list[dict], metrics: dict) -> None:
    ac = PARAMS["ac"]
    ur = PARAMS["ur"]
    rec_total = metrics["recommended_total"]
    rec_rsm2 = metrics["recommended_rsm2"]

    premissas = f"""# Flow da Lagoa - Premissas de origem do orçamento paramétrico v01

Gerado em {RUN_DATE}.

## Fonte principal

- Pasta: `{SOURCE}`
- Cliente/empreendimento identificado em prancha: AMS Empreendimentos Lagoa SPE Ltda.
- Método: Paramétrico V2 Híbrido, com briefing do projeto, base de índices Supabase `indices-cartesian`, referências Ampli e CUB-SC residencial médio.
- CUB usado: {money(PARAMS["cub"])}/m², referência maio/2026 para uso em junho/2026.

## Dados de projeto usados

{md_table([("Campo", "Valor", "Fonte"), *PROJECT_EVIDENCE])}

## Premissas críticas v01

- UR = 236 é premissa de contagem por área privativa média: {number(PARAMS["area_privativa_m2"] / ur)} m² privativos/UR. O quadro oficial de unidades não apareceu legível nos arquivos analisados.
- Vagas para parâmetro de orçamento = 161 automóveis. As 213 vagas de bicicleta foram registradas separadamente para não inflar a fórmula de garagem.
- Projeto com 4 blocos/núcleos verticais. O template V2 aceita até 3 torres no dropdown; por isso usei `Nº Torres = 3` e registrei o ajuste em notas.
- Elevadores = 4 por leitura visual dos núcleos. Validar com projeto de transporte vertical.
- Fundação = Hélice é premissa inicial. A disciplina FUND existe em PDF/DWG, mas não havia IFC quantitativo ou memorial consolidado para cravar o tipo sem revisão humana.
- Fachada = Textura é premissa inicial. Ajustar para cerâmica/ACM/pele de vidro se o memorial arquitetônico confirmar outro acabamento.
- BDI: tratado como cenário de validação em aba própria. O custo direto recomendado v01 é {money(rec_total)}, equivalente a {money(rec_rsm2)}/m².

## Inventário por disciplina

{md_table([("Disciplina", "Arquivos", "MB", "Extensões", "Leitura v01"), *[(d["disciplina"], d["arquivos"], d["tamanho_mb"], ", ".join(f"{k}:{v}" for k, v in d["extensoes"].items()), d["nota"]) for d in inv_summary]])}

## Arquivos gerados

- `parametrico-flow-da-lagoa-v01.xlsx`
- `parametrico-flow-da-lagoa.xlsx`
- `PREMISSAS-ORIGEM.md`
- `JUSTIFICATIVA-ITENS-ACIMA-DA-MEDIA.md`
- `ANALISE-PROJETOS-FLOW-DA-LAGOA.md`
- `inventario-arquivos-flow-da-lagoa.csv`
"""

    justificativa_rows = [
        ("Grupo/item", "Evidência do projeto", "Base de custo usada", "Tratamento v01"),
        (
            "Supraestrutura",
            "IFC estrutural R06 com 717 lajes, 1349 vigas e 1151 pilares; lajes maciças/mistas e=20 cm.",
            f"Supabase médio-alto mediana R$ 654,62/m²; alto R$ 718,24/m². Concreto 40 MPa mediana R$ 477,75/m³.",
            "Mantido acima da mediana médio-alto quando o V2 reagir a área e briefing; sem protensão.",
        ),
        (
            "Infraestrutura/fundações",
            "Disciplina FUND completa em PDF/DWG, sem IFC quantitativo.",
            "Supabase médio-alto mediana R$ 196,14/m²; alto R$ 230,40/m²; concreto de estacas mediana R$ 517,65/m³.",
            "Fundação hélice como premissa; validar após memorial/sondagem.",
        ),
        (
            "Esquadrias",
            "Padrão médio-alto, Lagoa, arquitetura com grandes vãos e interiores completos.",
            "Supabase: esquadrias G2 mediana R$ 939,74/m²; fornecimento+instalação mediana R$ 1.088,02/m². Ampli corroborou faixa ~R$ 958-1.250/m².",
            "Não usar faixa econômica; manter referência médio-alta.",
        ),
        (
            "Instalações",
            "HID/ELE/PCI/SPDA/TEL/CLI presentes, produto compacto com muitas áreas molhadas.",
            "Supabase médio-alto mediana R$ 332,12/m²; P75 R$ 411,27/m².",
            "Tratar como médio-alto com atenção a P75 se o quadro de unidades confirmar 236+ unidades.",
        ),
        (
            "Sistemas especiais/elevadores",
            "4 núcleos verticais e piscina.",
            "Supabase elevadores mediana R$ 261.192/un; Ampli elevador avg R$ 289.254/un e elevador social com alta dispersão.",
            "4 elevadores no v01; validar quantidade e padrão.",
        ),
        (
            "Impermeabilização",
            "Piscina, cobertura, pilotis/áreas comuns e muitos banheiros.",
            "Supabase total impermeabilização mediana R$ 61,39/m²; Ampli argamassa polimérica MO média R$ 27,31/m² e manta asfáltica MO R$ 48/m².",
            "Não reduzir para P25; manter mediana/P75 conforme áreas molhadas.",
        ),
        (
            "Pisos/revestimentos",
            "Interiores completos e briefing com porcelanato.",
            "Supabase pisos médio-alto mediana R$ 199,02/m²; Ampli porcelanato 60x60 ~R$ 80,49/m² e 80x80 ~R$ 129,22/m²; MO porcelanato mediana R$ 48/m².",
            "Porcelanato como piso predominante.",
        ),
        (
            "Gerenciamento/indiretos",
            "4 blocos, 15.000 m², obra horizontalizada em pilotis, prazo v01 24 meses.",
            "Supabase médio-alto mediana R$ 474,74/m² e P75 R$ 1.223,46/m², com alta dispersão.",
            "Planilha mantém gerenciamento no V2 e BDI em cenários separados.",
        ),
    ]
    justificativa = f"""# Flow da Lagoa - Justificativa dos itens e premissas v01

Este documento explica os itens que merecem maior atenção antes da validação do Leo e do cliente.

## Síntese de enquadramento

- Área construída: {number(ac)} m².
- Custo direto recomendado para controle v01: {money(rec_total)}.
- R$/m² recomendado v01: {money(rec_rsm2)}/m².
- Faixa Supabase: médio-alto {money(metrics["medio_alto_rsm2"])}/m²; alto {money(metrics["alto_rsm2"])}/m².
- O recomendado fica acima da mediana médio-alto e abaixo do alto, por localização, padrão e complexidade, sem assumir luxo.

## Justificativas

{md_table(justificativa_rows)}

## BDI / bifes

- A aba `BDI_PREMISSAS` inclui cenários 0%, 15%, 20% e 25%.
- O valor-base da planilha é custo direto paramétrico. BDI final deve ser validado com escopo comercial, impostos, administração central, risco e margem.

## Pendências para v02

- Confirmar quadro oficial de unidades.
- Confirmar quantidade e padrão dos elevadores.
- Confirmar tipo de fundação e contenções após memorial/sondagem.
- Confirmar acabamento de fachada.
- Validar se haverá gerador dedicado.
- Recalibrar prazo e indiretos quando houver cronograma físico preliminar.
"""

    analysis = f"""# Flow da Lagoa - Análise dos projetos v01

## Pasta analisada

`{SOURCE}`

As pastas `02 - Diretos` e `03 - Indiretos` no mesmo diretório de custo estavam vazias no momento da geração. Portanto, a v01 usa projetos como fonte técnica e Supabase/Ampli/CUB como fonte de custo.

## Inventário consolidado

{md_table([("Disciplina", "Arquivos", "MB", "Extensões"), *[(d["disciplina"], d["arquivos"], d["tamanho_mb"], ", ".join(f"{k}:{v}" for k, v in d["extensoes"].items())) for d in inv_summary]])}

## Leitura por disciplina

{chr(10).join(f'- **{d["disciplina"]}**: {d["nota"]}' for d in inv_summary)}

## Enquadramento adotado

- Padrão: médio-alto.
- Tipologia: 1-2 dormitórios/compactos.
- Área construída: {number(ac)} m².
- Área privativa: {number(PARAMS["area_privativa_m2"])} m².
- UR: {ur}, a validar.
- Pavimentos principais: 6.
- Prazo v01: 24 meses.
- Custo direto recomendado para controle: {money(rec_total)}.

## Limitações da leitura automática

- Não foi usado OCR externo; as principais tabelas foram lidas a partir dos renders já gerados da prancha arquitetônica.
- O IFC de arquitetura não trouxe `IfcSpace`, então não permitiu contagem automática confiável de unidades.
- DWG/PDF de disciplinas complementares foram inventariados e usados como presença/complexidade, mas a v01 ainda é paramétrica, não orçamento executivo por quantitativo.
"""

    (PKG / "PREMISSAS-ORIGEM.md").write_text(premissas, encoding="utf-8")
    (PKG / "JUSTIFICATIVA-ITENS-ACIMA-DA-MEDIA.md").write_text(justificativa, encoding="utf-8")
    (PKG / "ANALISE-PROJETOS-FLOW-DA-LAGOA.md").write_text(analysis, encoding="utf-8")

    for stem, text in [
        ("flow-da-lagoa-PREMISSAS-ORIGEM-v01.docx", premissas),
        ("flow-da-lagoa-JUSTIFICATIVA-ITENS-v01.docx", justificativa),
        ("flow-da-lagoa-ANALISE-PROJETOS-v01.docx", analysis),
    ]:
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        styles = doc.styles
        styles["Normal"].font.name = "Arial"
        styles["Normal"].font.size = Pt(10)
        for line in text.splitlines():
            if not line.strip():
                continue
            if line.startswith("# "):
                doc.add_heading(line[2:], 0)
            elif line.startswith("## "):
                doc.add_heading(line[3:], 1)
            elif line.startswith("- "):
                doc.add_paragraph(line[2:], style="List Bullet")
            elif line.startswith("|"):
                # Keep markdown tables as monospaced paragraphs in docx to avoid brittle table conversion.
                para = doc.add_paragraph()
                run = para.add_run(line)
                run.font.name = "Consolas"
                run.font.size = Pt(8)
            else:
                doc.add_paragraph(line)
        doc.save(PKG / stem)


def copy_final() -> None:
    FINAL_DIR.mkdir(parents=True, exist_ok=True)
    for name in [
        f"parametrico-{SLUG}-v01.xlsx",
        f"parametrico-{SLUG}.xlsx",
        "parametrico-v2-config.json",
        "PREMISSAS-ORIGEM.md",
        "JUSTIFICATIVA-ITENS-ACIMA-DA-MEDIA.md",
        "ANALISE-PROJETOS-FLOW-DA-LAGOA.md",
        "inventario-arquivos-flow-da-lagoa.csv",
        "inventario-resumo-flow-da-lagoa.json",
        "flow-da-lagoa-PREMISSAS-ORIGEM-v01.docx",
        "flow-da-lagoa-JUSTIFICATIVA-ITENS-v01.docx",
        "flow-da-lagoa-ANALISE-PROJETOS-v01.docx",
    ]:
        src = PKG / name
        if src.exists():
            shutil.copy2(src, FINAL_DIR / name)


def main() -> None:
    PKG.mkdir(parents=True, exist_ok=True)
    rows, by_disc, size_by_disc = inventory_files()
    write_inventory(rows, by_disc, size_by_disc)
    config_path = write_config()
    xlsx = run_generator(config_path)
    metrics = add_workbook_context(xlsx)

    inv_summary = json.loads((PKG / "inventario-resumo-flow-da-lagoa.json").read_text(encoding="utf-8"))
    write_docs(inv_summary, metrics)
    copy_final()

    report = {
        "package": str(PKG),
        "final_dir": str(FINAL_DIR),
        "xlsx": str(xlsx),
        "arquivos_inventariados": len(rows),
        "recommended_total": metrics["recommended_total"],
        "recommended_rsm2": metrics["recommended_rsm2"],
    }
    (PKG / "_RELATORIO-RUN.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

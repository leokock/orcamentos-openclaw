from __future__ import annotations

import shutil
from pathlib import Path
from statistics import mean, median

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PKG = Path(r"C:\Users\leona\orcamentos-openclaw\base\pacotes\flow-da-lagoa")
FINAL_DIR = Path(r"C:\Users\leona\orcamentos\parametricos\flow-da-lagoa")
INPUT = FINAL_DIR / "comparativo-flow-vs-similares-v02.xlsx"
OUTPUT = PKG / "comparativo-flow-vs-similares-v03-analise-detalhada.xlsx"
OUTPUT_MD = PKG / "ANALISE-DETALHADA-COMPARATIVO-FLOW-v03.md"


def brl(value: float) -> str:
    return f"R$ {value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def pct(value: float) -> str:
    return f"{value * 100:.1f}%".replace(".", ",")


def load_records():
    wb = load_workbook(INPUT, data_only=True)
    ws = wb["COMPARATIVO"]
    headers = [c.value for c in ws[1]]
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(headers, row))
        if rec.get("Empreendimento"):
            records.append(rec)
    return wb, records


def classify(rec: dict) -> str:
    name = rec["Empreendimento"]
    padrao = rec["Padrão"]
    tip = rec["Tipologia"]
    city = rec["Cidade/UF"]
    if name == "Flow da Lagoa":
        return "Projeto-base"
    if padrao == "alto":
        return "Limite superior"
    if "residencial_misto" in str(tip):
        return "Controle local misto"
    if "Florianópolis" in str(city):
        return "Comparável local direto"
    return "Comparável regional direto"


def add_analysis_sheet(wb, records):
    if "ANALISE_DETALHADA" in wb.sheetnames:
        del wb["ANALISE_DETALHADA"]

    ws = wb.create_sheet("ANALISE_DETALHADA", 1)
    blue = PatternFill("solid", fgColor="245AE4")
    orange = PatternFill("solid", fgColor="FD3400")
    gray = PatternFill("solid", fgColor="F4F4F4")
    white = Font(color="FFFFFF", bold=True)
    bold = Font(bold=True)
    thin = Side(style="thin", color="E7E7E9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    flow = next(r for r in records if r["Empreendimento"] == "Flow da Lagoa")
    comps = [r for r in records if r["Empreendimento"] != "Flow da Lagoa"]
    direct = [r for r in comps if classify(r) in ("Comparável local direto", "Comparável regional direto")]
    premium = [r for r in comps if classify(r) == "Limite superior"]
    local = [r for r in comps if "Florianópolis" in str(r["Cidade/UF"])]

    ws["A1"] = "ANÁLISE DETALHADA - FLOW DA LAGOA VS 10 COMPARÁVEIS"
    ws["A1"].font = Font(bold=True, size=14, color="111111")
    ws.merge_cells("A1:H1")

    summary_rows = [
        ("Flow R$/m²", flow["R$/m²"], "Custo direto recomendado v01 do paramétrico Flow."),
        ("Média comparáveis diretos", mean(r["R$/m²"] for r in direct), "Exclui padrão alto e residencial misto."),
        ("Mediana comparáveis diretos", median(r["R$/m²"] for r in direct), "Melhor leitura contra outliers."),
        ("Média todos os 10 comparáveis", mean(r["R$/m²"] for r in comps), "Inclui limites altos e misto local."),
        ("Mediana todos os 10 comparáveis", median(r["R$/m²"] for r in comps), "Centro da amostra ampliada."),
        ("Média limites superiores", mean(r["R$/m²"] for r in premium), "Neuhaus Origem e Nobria."),
        ("Média Florianópolis/locais", mean(r["R$/m²"] for r in local), "Lumis, Neuhaus, Cota e Hacasa."),
    ]
    ws.append(["Indicador", "Valor", "Leitura"])
    for row in summary_rows:
        ws.append(row)

    start = ws.max_row + 3
    ws.cell(start, 1, "Leitura por empreendimento")
    ws.cell(start, 1).fill = orange
    ws.cell(start, 1).font = white
    ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=8)
    start += 1
    headers = ["Empreendimento", "Grupo", "AC", "R$/m²", "Dif. vs Flow", "Total norm.", "Leitura", "Peso recomendado"]
    for col, header in enumerate(headers, start=1):
        ws.cell(start, col, header)
    for rec in comps:
        group = classify(rec)
        diff = rec["Dif. % R$/m² vs Flow"]
        if group == "Limite superior":
            leitura = "Usar como teto/alerta de padrão; não puxar Flow para cima sem memorial premium."
            peso = "Médio-baixo"
        elif group == "Controle local misto":
            leitura = "Útil por Florianópolis, mas tipologia mista reduz comparabilidade direta."
            peso = "Médio"
        elif rec["Empreendimento"] == "pass-e-connect":
            leitura = "R$/m² praticamente colado ao Flow; bom teste de plausibilidade para produto compacto."
            peso = "Alto"
        elif diff < -0.15:
            leitura = "Mais barato que Flow; provável diferença de data-base, escopo ou padrão. Usar com cautela."
            peso = "Médio"
        elif diff > 0.10:
            leitura = "Acima do Flow; usar como pressão de alta somente se escopo confirmar padrão superior."
            peso = "Médio"
        else:
            leitura = "Comparável direto para faixa de custo do Flow."
            peso = "Alto"
        ws.append([
            rec["Empreendimento"],
            group,
            rec["AC (m²)"],
            rec["R$/m²"],
            diff,
            rec["Total normalizado na AC Flow"],
            leitura,
            peso,
        ])

    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in (2, start):
        for cell in ws[row]:
            cell.fill = blue
            cell.font = white
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=3):
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = gray
    for row in range(3, 10):
        ws.cell(row, 2).number_format = 'R$ #,##0.00'
    for row in range(start + 1, ws.max_row + 1):
        ws.cell(row, 3).number_format = '#,##0.00'
        ws.cell(row, 4).number_format = 'R$ #,##0.00'
        ws.cell(row, 5).number_format = '0.0%'
        ws.cell(row, 6).number_format = 'R$ #,##0.00'
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        width = max(len(str(ws.cell(row, col).value or "")) for row in range(1, ws.max_row + 1))
        ws.column_dimensions[letter].width = min(max(width + 2, 12), 58)

    return {
        "flow": flow,
        "direct": direct,
        "premium": premium,
        "local": local,
        "all": comps,
    }


def write_markdown(groups):
    flow = groups["flow"]
    direct = groups["direct"]
    premium = groups["premium"]
    local = groups["local"]
    comps = groups["all"]

    direct_avg = mean(r["R$/m²"] for r in direct)
    direct_median = median(r["R$/m²"] for r in direct)
    all_avg = mean(r["R$/m²"] for r in comps)
    all_median = median(r["R$/m²"] for r in comps)
    premium_avg = mean(r["R$/m²"] for r in premium)
    local_avg = mean(r["R$/m²"] for r in local)

    lines = [
        "# Análise detalhada - Flow da Lagoa vs comparáveis v03",
        "",
        "## Conclusão executiva",
        "",
        f"- O Flow está em {brl(flow['R$/m²'])}/m² e {brl(flow['Total original'])} de custo direto v01.",
        f"- A mediana dos comparáveis diretos é {brl(direct_median)}/m²; o Flow fica {pct(flow['R$/m²'] / direct_median - 1)} acima dessa mediana.",
        f"- A média dos comparáveis diretos é {brl(direct_avg)}/m²; o Flow fica {pct(flow['R$/m²'] / direct_avg - 1)} acima dessa média.",
        f"- A média dos 10 comparáveis é {brl(all_avg)}/m² e a mediana é {brl(all_median)}/m².",
        f"- Os limites superiores de padrão alto ficam em média em {brl(premium_avg)}/m².",
        f"- Os comparáveis locais de Florianópolis ficam em média em {brl(local_avg)}/m², mas misturam médio-alto, alto e tipologia mista.",
        "",
        "## Leitura de faixa",
        "",
        "- A faixa médio-alto direta fica principalmente entre Lumis Live, Mussi Chelsea, Fonseca Estoril, Pavcor, Pass-e Connect, Mabrem Gran Torino e Cota 365.",
        "- Pass-e Connect é o comparável mais aderente em R$/m², ficando praticamente no mesmo patamar do Flow.",
        "- Lumis Live e Mabrem Gran Torino puxam a média para baixo; devem ser lidos com cautela por data-base/escopo/padrão.",
        "- Neuhaus Origem e Nobria não são comparáveis diretos de padrão; funcionam como teto de referência.",
        "- Hacasa Brisa da Armação é útil por ser Florianópolis, mas a tipologia residencial mista reduz o peso.",
        "",
        "## Recomendação para a próxima validação",
        "",
        "- Manter o Flow v01 como faixa conservadora médio-alta, não como orçamento baixo.",
        "- Antes de reduzir custo, validar fachada, fundação, elevadores, gerador e contagem de UR.",
        "- Se o cliente pressionar por referência mais baixa, usar Lumis/Mabrem apenas como piso e explicar diferença de escopo/data-base.",
        "- Se o memorial confirmar padrão alto/local premium, usar Neuhaus como teto e recalibrar o Flow para cima.",
        "",
        "## Comparáveis por empreendimento",
        "",
    ]
    for rec in comps:
        lines.extend([
            f"- **{rec['Empreendimento']}**",
            f"  - AC: {rec['AC (m²)']:,.2f} m² | R$/m²: {brl(rec['R$/m²'])} | dif. vs Flow: {pct(rec['Dif. % R$/m² vs Flow'])}",
            f"  - Total normalizado na AC do Flow: {brl(rec['Total normalizado na AC Flow'])}",
        ])
    OUTPUT_MD.write_text("\n".join(lines), encoding="utf-8")


def main():
    wb, records = load_records()
    groups = add_analysis_sheet(wb, records)
    wb.save(OUTPUT)
    write_markdown(groups)
    FINAL_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUTPUT, FINAL_DIR / OUTPUT.name)
    shutil.copy2(OUTPUT_MD, FINAL_DIR / OUTPUT_MD.name)
    print(OUTPUT)
    print(FINAL_DIR / OUTPUT.name)


if __name__ == "__main__":
    main()

from __future__ import annotations

from pathlib import Path
import json
import shutil

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PKG = Path(r"C:\Users\leona\orcamentos-openclaw\base\pacotes\flow-da-lagoa")
FINAL_DIR = Path(r"C:\Users\leona\orcamentos\parametricos\flow-da-lagoa")
SOURCE_REVIEW = Path(
    r"G:\Drives compartilhados\03 CTN Projetos\2. Projetos em Andamento"
    r"\_Parametrico_IA\REVISAO-OBRAS-131.xlsx"
)
DRIVE_ROOT = Path(r"G:\Drives compartilhados\03 CTN Projetos\2. Projetos em Andamento")

OUTPUT = PKG / "comparativo-flow-vs-similares-v02.xlsx"
OUTPUT_MD = PKG / "COMPARATIVO-FLOW-VS-SIMILARES-v02.md"

FLOW = {
    "slug": "flow-da-lagoa",
    "cliente": "AMS",
    "empreendimento": "Flow da Lagoa",
    "cidade": "Florianópolis",
    "uf": "SC",
    "padrao": "medio-alto",
    "tipologia": "residencial_vertical_medio_alto",
    "data_base": "2026-06",
    "ac": 15000.62,
    "ur": 236,
    "rsm2": 3680.00,
    "total": 55202281.60,
    "fonte": r"C:\Users\leona\orcamentos\parametricos\flow-da-lagoa\parametrico-flow-da-lagoa-v01.xlsx",
    "justificativa": "Projeto-base desta análise. Custo direto recomendado v01 do paramétrico Flow.",
}

SELECTED_SLUGS = [
    "lumis-live",
    "mussi-empreendimentos-chelsea",
    "fonseca-empreendimentos-estoril",
    "pavcor",
    "neuhaus-origem",
    "pass-e-connect",
    "nobria",
    "mabrem-gran-torino",
    "cota-365",
    "hacasa-brisa-da-armacao",
]

SOURCE_FOLDERS = {
    "lumis-live": DRIVE_ROOT / "Lumis - 10.2021" / "2021 - Live" / "03. Custo",
    "mussi-empreendimentos-chelsea": DRIVE_ROOT / "Mussi Empreendimentos" / "Chelsea Residence 11.2022",
    "fonseca-empreendimentos-estoril": DRIVE_ROOT / "Fonseca" / "Residencial Estoril - Orçamento e Planejamento",
    "pavcor": DRIVE_ROOT / "Pavcor" / "CTN & Pavcor - Cinque",
    "neuhaus-origem": DRIVE_ROOT / "Neuhaus" / "2023.08 - Origem 3300",
    "pass-e-connect": DRIVE_ROOT / "Pass-e Empreendimentos" / "Connect",
    "nobria": DRIVE_ROOT / "Nobria Construtora - 03.2023" / "03. Custo",
    "mabrem-gran-torino": DRIVE_ROOT / "Mabrem 10.2020" / "Gran Torino",
    "cota-365": DRIVE_ROOT / "COTA" / "365 (antigo Afonso Pena)",
    "hacasa-brisa-da-armacao": DRIVE_ROOT / "Hacasa" / "2024 - Brisa da Armação",
}

JUSTIFICATIVAS = {
    "lumis-live": "AC quase idêntica ao Flow (diferença de 112 m²), Florianópolis/SC, residencial vertical médio-alto e orçamento consolidado com R$/m² válido.",
    "mussi-empreendimentos-chelsea": "AC muito próxima (diferença de 149 m²), residencial vertical em SC, padrão médio-alto e custo total/R$/m² válidos no consolidado.",
    "fonseca-empreendimentos-estoril": "Residencial vertical médio-alto em SC, AC 14.492 m², UR informado e data-base recente no consolidado; bom comparável de produto não-luxo.",
    "pavcor": "Residencial vertical médio-alto em SC, AC 14.283 m² e R$/m² próximo do cenário Flow; entra como comparável regional de porte semelhante.",
    "neuhaus-origem": "AC próxima e Florianópolis/SC, porém padrão alto. Usei como limite superior para testar se o Flow está abaixo de um produto mais premium.",
    "pass-e-connect": "Residencial vertical médio-alto em SC, AC 13.144 m², UR informado e R$/m² muito próximo do Flow; bom controle de produto compacto em Itajaí.",
    "nobria": "Residencial vertical alto em Bombinhas/SC, AC 12.880 m²; entra como comparável de porte médio e padrão superior fora de Florianópolis.",
    "mabrem-gran-torino": "Residencial vertical médio-alto em Piçarras/SC, AC 12.519 m²; escolhido por porte próximo e R$/m² válido no consolidado.",
    "cota-365": "Residencial vertical médio-alto em Florianópolis/SC, AC 17.506 m²; entra por cidade e padrão, com área um pouco maior que o Flow.",
    "hacasa-brisa-da-armacao": "Empreendimento residencial misto médio-alto em Florianópolis/SC, AC 12.828 m² e UR informado; útil como controle local com tipologia menos direta.",
}


def brl(value: float) -> str:
    return f"R$ {value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def pct(value: float) -> str:
    return f"{value * 100:.1f}%".replace(".", ",")


def load_review_rows() -> dict[str, dict]:
    wb = load_workbook(SOURCE_REVIEW, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        slug, cliente, cidade, uf, cub_regiao, padrao, tipologia, data_base, data_entrega, ac, ur, rsm2, total = row[:13]
        if not slug:
            continue
        rows[str(slug)] = {
            "slug": slug,
            "cliente": cliente,
            "empreendimento": slug,
            "cidade": cidade,
            "uf": uf,
            "cub_regiao": cub_regiao,
            "padrao": padrao,
            "tipologia": tipologia,
            "data_base": data_base,
            "data_entrega": data_entrega,
            "ac": float(ac) if ac is not None else None,
            "ur": ur,
            "rsm2": float(rsm2) if rsm2 is not None else None,
            "total": float(total) if total is not None else None,
        }
    return rows


def candidate_record(row: dict) -> dict:
    ac = row["ac"]
    rsm2 = row["rsm2"]
    total_at_flow = rsm2 * FLOW["ac"]
    return {
        **row,
        "fonte": str(SOURCE_REVIEW),
        "pasta_drive": str(SOURCE_FOLDERS.get(row["slug"], "")),
        "pasta_existe": SOURCE_FOLDERS.get(row["slug"], Path()).exists() if row["slug"] in SOURCE_FOLDERS else False,
        "justificativa": JUSTIFICATIVAS[row["slug"]],
        "diff_ac_m2": ac - FLOW["ac"],
        "diff_ac_pct": (ac / FLOW["ac"]) - 1,
        "diff_rsm2": rsm2 - FLOW["rsm2"],
        "diff_rsm2_pct": (rsm2 / FLOW["rsm2"]) - 1,
        "total_normalizado_flow_ac": total_at_flow,
        "diff_total_normalizado": total_at_flow - FLOW["total"],
        "diff_total_normalizado_pct": (total_at_flow / FLOW["total"]) - 1,
    }


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="245AE4")
    orange_fill = PatternFill("solid", fgColor="FD3400")
    gray_fill = PatternFill("solid", fgColor="F4F4F4")
    white_font = Font(color="FFFFFF", bold=True)
    bold = Font(bold=True, color="111111")
    thin = Side(style="thin", color="E7E7E9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = gray_fill
    return header_fill, orange_fill, gray_fill, white_font, bold


def autosize(ws, max_width: int = 58) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min(max_width, max(10, max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2))
        ws.column_dimensions[letter].width = width


def write_workbook(candidates: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "COMPARATIVO"

    headers = [
        "Empreendimento",
        "Cliente",
        "Cidade/UF",
        "Padrão",
        "Tipologia",
        "Data base",
        "AC (m²)",
        "UR",
        "R$/m²",
        "Total original",
        "Total normalizado na AC Flow",
        "Dif. R$/m² vs Flow",
        "Dif. % R$/m² vs Flow",
        "Dif. total norm. vs Flow",
        "Pasta Drive",
    ]
    ws.append(headers)
    all_rows = [FLOW, *candidates]
    for rec in all_rows:
        city_uf = f"{rec.get('cidade') or ''}/{rec.get('uf') or ''}".strip("/")
        diff_rsm2 = 0 if rec["slug"] == FLOW["slug"] else rec["diff_rsm2"]
        diff_rsm2_pct = 0 if rec["slug"] == FLOW["slug"] else rec["diff_rsm2_pct"]
        diff_total = 0 if rec["slug"] == FLOW["slug"] else rec["diff_total_normalizado"]
        total_norm = rec["total"] if rec["slug"] == FLOW["slug"] else rec["total_normalizado_flow_ac"]
        ws.append(
            [
                rec["empreendimento"],
                rec["cliente"],
                city_uf,
                rec["padrao"],
                rec["tipologia"],
                rec.get("data_base"),
                rec["ac"],
                rec.get("ur"),
                rec["rsm2"],
                rec["total"],
                total_norm,
                diff_rsm2,
                diff_rsm2_pct,
                diff_total,
                rec.get("pasta_drive", rec.get("fonte")),
            ]
        )

    style_sheet(ws)
    for row in ws.iter_rows(min_row=2):
        row[6].number_format = '#,##0.00'
        row[8].number_format = 'R$ #,##0.00'
        row[9].number_format = 'R$ #,##0.00'
        row[10].number_format = 'R$ #,##0.00'
        row[11].number_format = 'R$ #,##0.00'
        row[12].number_format = '0.0%'
        row[13].number_format = 'R$ #,##0.00'
    ws.freeze_panes = "A2"
    autosize(ws)

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "R$/m² - Flow vs comparáveis"
    chart.y_axis.title = "Empreendimento"
    chart.x_axis.title = "R$/m²"
    data = Reference(ws, min_col=9, min_row=1, max_row=1 + len(all_rows))
    cats = Reference(ws, min_col=1, min_row=2, max_row=1 + len(all_rows))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 8
    chart.width = 16
    ws.add_chart(chart, "Q2")

    wsj = wb.create_sheet("JUSTIFICATIVAS")
    wsj.append(["Empreendimento", "Por que entrou", "Risco/uso no comparativo"])
    wsj.append([
        FLOW["empreendimento"],
        "Base do comparativo, gerada na v01 do paramétrico Flow.",
        "Premissas ainda a validar: UR, elevadores, fundação, fachada e gerador.",
    ])
    for rec in candidates:
        uso = "Comparável direto" if rec["padrao"] == "medio-alto" else "Limite superior por padrão alto"
        wsj.append([rec["empreendimento"], rec["justificativa"], uso])
    style_sheet(wsj)
    autosize(wsj)

    wsf = wb.create_sheet("CRITERIO_E_FONTES")
    wsf.append(["Item", "Descrição"])
    criteria = [
        ("Fonte principal", str(SOURCE_REVIEW)),
        ("Critério 1", "R$/m² preenchido no consolidado de obras revisadas."),
        ("Critério 2", "Residencial vertical em SC."),
        ("Critério 3", "Padrão médio-alto como preferência; alto aceito como limite superior."),
        ("Critério 4", "Área construída mais próxima de 15.000,62 m²."),
        ("Normalização", "Além do total original, calculei o total de cada R$/m² aplicado sobre a AC do Flow para comparação justa."),
        ("Correção temporal", "Não apliquei correção por CUB entre data-bases; a planilha expõe a data-base quando disponível."),
    ]
    for item in criteria:
        wsf.append(item)
    wsf.append(["", ""])
    wsf.append(["Slug", "Pasta Drive encontrada"])
    for rec in candidates:
        wsf.append([rec["slug"], rec["pasta_drive"]])
    style_sheet(wsf)
    autosize(wsf)

    wsd = wb.create_sheet("DADOS_BRUTOS")
    raw_headers = [
        "slug",
        "cliente",
        "cidade",
        "uf",
        "padrao",
        "tipologia",
        "data_base",
        "ac",
        "ur",
        "rsm2",
        "total",
        "diff_ac_pct",
        "diff_rsm2_pct",
        "total_normalizado_flow_ac",
        "pasta_existe",
    ]
    wsd.append(raw_headers)
    for rec in candidates:
        wsd.append([rec.get(h) for h in raw_headers])
    style_sheet(wsd)
    autosize(wsd)

    wb.save(OUTPUT)


def write_markdown(candidates: list[dict]) -> None:
    lines = [
        "# Flow da Lagoa - Comparativo com empreendimentos semelhantes v02",
        "",
        f"Fonte principal: `{SOURCE_REVIEW}`.",
        "",
        "## Empreendimentos escolhidos",
        "",
    ]
    for rec in candidates:
        lines.extend(
            [
                f"- **{rec['empreendimento']}**: {rec['justificativa']}",
                f"  - AC: {rec['ac']:,.2f} m² | R$/m²: {brl(rec['rsm2'])} | Total: {brl(rec['total'])}",
                f"  - Diferença vs Flow: {pct(rec['diff_rsm2_pct'])} em R$/m²; total normalizado: {brl(rec['total_normalizado_flow_ac'])}",
                "",
            ]
        )
    lines.extend(
        [
            "## Observações",
            "",
            "- Comparei por R$/m² e por custo normalizado na área do Flow, porque total bruto varia com a área construída.",
            "- Não apliquei correção temporal por CUB entre data-bases; a planilha mantém a data-base de cada obra para validação posterior.",
            "- Neuhaus Origem e Nobria entram como limites superiores por padrão alto, não como comparáveis médio-alto diretos.",
            "- Hacasa Brisa da Armação entra como controle local, mas a tipologia é residencial misto, então deve ter peso menor que os residenciais verticais diretos.",
        ]
    )
    OUTPUT_MD.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    rows = load_review_rows()
    candidates = []
    for slug in SELECTED_SLUGS:
        if slug not in rows:
            raise RuntimeError(f"Slug não encontrado no consolidado: {slug}")
        candidates.append(candidate_record(rows[slug]))

    write_workbook(candidates)
    write_markdown(candidates)
    FINAL_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUTPUT, FINAL_DIR / OUTPUT.name)
    shutil.copy2(OUTPUT_MD, FINAL_DIR / OUTPUT_MD.name)
    report = {
        "output": str(OUTPUT),
        "final_output": str(FINAL_DIR / OUTPUT.name),
        "candidatos": [c["slug"] for c in candidates],
    }
    (PKG / "_RELATORIO-COMPARATIVO.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

# -*- coding: utf-8 -*-
"""
Template Cartesian — Ata de Reunião em Excel
=============================================

Gera uma ata no padrão visual Cartesian (paleta V2):
- Header da tabela e labels REUNIÃO/OBRA em azul oficial #245AE4
- Logo Cartesian no canto superior esquerdo
- Tabela com colunas: ITEM, ASSUNTO ABORDADO, DISCUSSÃO, RESPONSÁVEL, DATA, STATUS
- Layout calcado na ata "Diagrama de Rede 20/10/2025" do projeto Electra Towers

USO:
  1. Editar constantes no topo: DATA, NOME_REUNIAO_LINHA, NOME_REUNIAO_META, OBRA
  2. Preencher lista ITENS — cada tupla: (ASSUNTO, DISCUSSÃO, RESPONSÁVEL, DATA_RETORNO, STATUS)
  3. Ajustar path do logo se necessário (default: C:/Temp/cartesian-icon.png)
  4. Rodar: PYTHONIOENCODING=utf-8 python gerar_ata_excel.py
  5. Excel sai em C:/Temp/ata_thozen_30-04.xlsx — copiar pro Drive
  6. Pra gerar PDF: Excel COM via PowerShell (Workbook.ExportAsFixedFormat(0, path))

REFERÊNCIAS:
  - Paleta oficial Cartesian V2: ../references/identidade-visual.md
  - Brand guide PDF: ../../../referencias/cartesian/branding/cartesian-brand-guide-v2.pdf
  - Exemplo real gerado: ~/orcamentos/.../Electra Towers/.../CTN-TZN_ELT - Ata Entrega Parcial Orçamento_2026-04-30.xlsx

Para extrair só o ícone do logo (canto esquerdo do Logo-azul.png):
  python -c "from PIL import Image; img = Image.open('skills/cartesian-presentation/assets/Logo-azul.png'); img.crop((0, 0, 360, 448)).save('C:/Temp/cartesian-icon.png')"
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as XLImage

# Dados da reuniao
DATA = "30 de abril de 2026"
NOME_REUNIAO_LINHA = "CTN & THOZEN - ELECTRA TOWERS - ENTREGA PARCIAL ORÇAMENTO EXECUTIVO - 2026/04/30"
NOME_REUNIAO_META = "CTN & THOZEN - Electra Towers - Entrega Parcial Orçamento Executivo - 2026/04/30"
OBRA = "Electra Towers"

ITENS = [
    ("PARTICIPANTES",
     "Reunião presencial realizada no escritório da Thozen Construtora e Incorporadora às 16h30 BRT. "
     "Participaram pela Cartesian: Leonardo Kock Adriano (Coordenação do projeto). "
     "Pela Thozen: Nicholas (Gerente da obra) e Vinicius (Pesquisa e Desenvolvimento).",
     "INFORMAÇÃO", "", "Registrado"),
    ("QUANTITATIVOS DAS INSTALAÇÕES",
     "A Thozen solicitará aos projetistas a lista atualizada de quantitativos das disciplinas de "
     "PPCI, elétrica e hidrossanitário, para incorporação ao orçamento executivo. Essa atualização "
     "é necessária porque os quantitativos atuais estão incompletos e impactam a precificação das instalações.",
     "NICHOLAS (THOZEN)", "", "Definido"),
    ("QUANTITATIVOS DOS DEMAIS PROJETOS",
     "Os projetos disponibilizados não contemplam lista de quantitativos formal. A Thozen solicitará "
     "aos projetistas o envio dessas listas para subsidiar a precificação das disciplinas restantes.",
     "NICHOLAS (THOZEN)", "", "Definido"),
    ("PROJETO DE INTERIORES",
     "A Thozen encaminhará à Cartesian o projeto de interiores para cotação dos itens de mobiliário "
     "e decoração, viabilizando a inclusão desses custos no orçamento executivo.",
     "NICHOLAS (THOZEN)", "", "Definido"),
    ("FUNDAÇÃO – CUSTOS REAIS EXECUTADOS",
     "A obra já executou a fundação. A Thozen disponibilizará os custos reais executados, que serão "
     "acessados pela Cartesian por meio da plataforma Ampli, substituindo a estimativa orçamentária "
     "desta disciplina. Resolve as pendências prévias P1 (tipo de fundação) e P2 (tipo de contenção).",
     "NICHOLAS (THOZEN)", "", "Definido"),
    ("COMPARTILHAMENTO DA ENTREGA PARCIAL",
     "A Cartesian encaminhará a Nicholas e Vinicius o orçamento publicado no software Visus, "
     "juntamente com a planilha Excel atualizada do executivo parcial, para conferência.",
     "LEONARDO KOCK (CARTESIAN)", "", "Em andamento"),
    ("CICLO DE REVISÃO E CONSOLIDAÇÃO",
     "A Thozen revisará a entrega parcial e devolverá observações e ajustes à Cartesian, que então "
     "consolidará a versão final do orçamento executivo a partir do retorno do cliente.",
     "NICHOLAS E VINICIUS (THOZEN)", "", "Em andamento"),
]

# Cores e estilos
COR_HEADER_TABELA = "245AE4"   # azul oficial Cartesian
COR_TEXTO_BRANCO = "FFFFFF"
COR_META_LABEL = "245AE4"      # azul Cartesian nos labels REUNIAO/OBRA
COR_BORDA = "808080"

fonte_titulo1 = Font(name="Arial", size=14, bold=True, color="000000")
fonte_titulo2 = Font(name="Arial", size=11, bold=True, color="000000")
fonte_data = Font(name="Arial", size=11, bold=True, color="000000")
fonte_meta_label = Font(name="Arial", size=10, bold=True, color="FFFFFF")
fonte_header_tabela = Font(name="Arial", size=10, bold=True, color=COR_TEXTO_BRANCO)
fonte_item_assunto = Font(name="Arial", size=9, bold=True, color="000000")
fonte_item_normal = Font(name="Arial", size=9, color="000000")

fill_header = PatternFill(start_color=COR_HEADER_TABELA, end_color=COR_HEADER_TABELA, fill_type="solid")
fill_meta_label = PatternFill(start_color=COR_META_LABEL, end_color=COR_META_LABEL, fill_type="solid")

borda_fina = Border(
    left=Side(style="thin", color=COR_BORDA),
    right=Side(style="thin", color=COR_BORDA),
    top=Side(style="thin", color=COR_BORDA),
    bottom=Side(style="thin", color=COR_BORDA),
)

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

# Workbook
wb = Workbook()
ws = wb.active
ws.title = "Ata 30-04-2026"

# Larguras (A=ITEM, B=ASSUNTO, C=DISCUSSAO, D=RESPONSAVEL, E=DATA, F=STATUS)
ws.column_dimensions["A"].width = 11   # cabe REUNIAO e OBRA
ws.column_dimensions["B"].width = 24
ws.column_dimensions["C"].width = 58
ws.column_dimensions["D"].width = 20
ws.column_dimensions["E"].width = 20   # cabe DATA DE RETORNO/DEFINICAO em 2 linhas
ws.column_dimensions["F"].width = 14

# Cabecalho linhas 1-4
ws.row_dimensions[1].height = 22
ws.row_dimensions[2].height = 18
ws.row_dimensions[3].height = 18
ws.row_dimensions[4].height = 8

ws.merge_cells("A1:A4")
img = XLImage("C:/Temp/cartesian-icon.png")
img.width = 70
img.height = 86
img.anchor = "A1"
ws.add_image(img)

ws.merge_cells("B1:F1")
ws["B1"] = "CARTESIAN & THOZEN – ELECTRA TOWERS"
ws["B1"].font = fonte_titulo1
ws["B1"].alignment = align_center

ws.merge_cells("B2:F2")
ws["B2"] = NOME_REUNIAO_LINHA
ws["B2"].font = fonte_titulo2
ws["B2"].alignment = align_center

ws.merge_cells("B3:F3")
ws["B3"] = f"DATA: {DATA}"
ws["B3"].font = fonte_data
ws["B3"].alignment = align_center

# Linha META 5
ws.row_dimensions[5].height = 30
ws["A5"] = "REUNIÃO"
ws["A5"].font = fonte_meta_label
ws["A5"].fill = fill_meta_label
ws["A5"].alignment = align_center
ws["A5"].border = borda_fina

ws.merge_cells("B5:C5")
ws["B5"] = NOME_REUNIAO_META
ws["B5"].font = fonte_item_normal
ws["B5"].alignment = align_left
for col in ["B5", "C5"]:
    ws[col].border = borda_fina

ws["D5"] = "OBRA"
ws["D5"].font = fonte_meta_label
ws["D5"].fill = fill_meta_label
ws["D5"].alignment = align_center
ws["D5"].border = borda_fina

ws.merge_cells("E5:F5")
ws["E5"] = OBRA
ws["E5"].font = fonte_item_normal
ws["E5"].alignment = align_left
for col in ["E5", "F5"]:
    ws[col].border = borda_fina

# Linha 6 vazia
ws.row_dimensions[6].height = 10

# Header tabela
ws.row_dimensions[7].height = 42
headers = ["ITEM", "ASSUNTO ABORDADO", "DISCUSSÃO", "RESPONSÁVEL", "DATA DE RETORNO/DEFINIÇÃO", "STATUS"]
for i, h in enumerate(headers, start=1):
    c = ws.cell(row=7, column=i, value=h)
    c.font = fonte_header_tabela
    c.fill = fill_header
    c.alignment = align_center
    c.border = borda_fina

# Itens
for idx, (assunto, discussao, responsavel, data_retorno, status) in enumerate(ITENS, start=1):
    row = 7 + idx
    altura = max(60, min(180, 18 + len(discussao) // 4))
    ws.row_dimensions[row].height = altura

    ws.cell(row=row, column=1, value=idx).font = fonte_item_normal
    ws.cell(row=row, column=1).alignment = align_center

    ws.cell(row=row, column=2, value=assunto).font = fonte_item_assunto
    ws.cell(row=row, column=2).alignment = align_left

    ws.cell(row=row, column=3, value=discussao).font = fonte_item_normal
    ws.cell(row=row, column=3).alignment = align_left

    ws.cell(row=row, column=4, value=responsavel).font = fonte_item_normal
    ws.cell(row=row, column=4).alignment = align_center

    ws.cell(row=row, column=5, value=data_retorno).font = fonte_item_normal
    ws.cell(row=row, column=5).alignment = align_center

    ws.cell(row=row, column=6, value=status).font = fonte_item_normal
    ws.cell(row=row, column=6).alignment = align_center

    for col in range(1, 7):
        ws.cell(row=row, column=col).border = borda_fina

# Page setup
ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_options.horizontalCentered = True
ws.page_margins.left = 0.4
ws.page_margins.right = 0.4
ws.page_margins.top = 0.5
ws.page_margins.bottom = 0.5

saida = "C:/Temp/ata_thozen_30-04.xlsx"
wb.save(saida)
print(f"OK: {saida}")

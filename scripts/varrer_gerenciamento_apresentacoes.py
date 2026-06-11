#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Varrer SÓ os xlsx "Apresentação Orçamento" — formato padronizado da Cartesian:
aba ORÇAMENTO_EXECUTIVO contém linha por macrogrupo (ETAPA | VALOR | % | R$/m²)
+ linha TOTAL + ÁREA CONSTRUÍDA em bloco adjacente.

Gera CSV: cliente, projeto, ac_m2, total_obra, total_gerenciamento, pct, rsm2.

Uso:
    PYTHONIOENCODING=utf-8 py -3.10 -X utf8 scripts/varrer_gerenciamento_apresentacoes.py
"""
from __future__ import annotations

import csv
import io
import os
import re
import sys
import time
import warnings
from pathlib import Path
from typing import Optional

import openpyxl

warnings.filterwarnings("ignore")

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

DRIVE_ROOT = Path(
    r"G:\Drives compartilhados\03 CTN Projetos\2. Projetos em Andamento\_Entregas\Orçamento_executivo"
)
OUTPUT_CSV = Path(r"C:\Users\leona\orcamentos-openclaw\base\executivos-gerenciamento-apresentacoes-2026-04-22.csv")
LOG_TXT = Path(r"C:\Users\leona\orcamentos-openclaw\base\executivos-gerenciamento-apresentacoes-2026-04-22.log.txt")

TRANS = str.maketrans("áàâãäéèêëíìîïóòôõöúùûüç", "aaaaaeeeeiiiiooooouuuuc")


def norm(s) -> str:
    if s is None:
        return ""
    return str(s).lower().strip().translate(TRANS)


# Tokens que identificam linha de Gerenciamento/Canteiro/Indiretos
GEREN_ROW_TOKENS = [
    "gerenciamento",
    "gerencimento",  # typo comum
    "administr",
    "canteiro e indireto",
    "canteiro_indireto",
    "despesas indiretas",
    "desp indireta",
    "custos indiretos",
    "indiretos",
    "bdi direto",
    "servicos preliminares e adm",
]

# Palavras que desqualificam (não é Gerenciamento principal)
EXCLUDE_ROW_TOKENS = ["pos-obra", "pos obra", "pre-obra", "posterga"]

# Tokens AC
AC_TOKENS = [
    "area construida",
    "area construída",
    "area total construida",
    "area total de construcao",
    "act",
    "area real",
]


def to_num(v) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if v != 0 else None
    s = str(v).strip()
    if not s:
        return None
    s = re.sub(r"[R\$\s]", "", s)
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        n = float(s)
        return n if n != 0 else None
    except ValueError:
        return None


def find_ap_sheet(wb) -> Optional[str]:
    """Acha a aba do tipo ORÇAMENTO_EXECUTIVO / ORÇAMENTO SINTÉTICO / APRESENTAÇÃO."""
    candidates = []
    for sn in wb.sheetnames:
        n = norm(sn)
        # priority
        if "orcamento_executivo" in n or "orcamento executivo" in n:
            candidates.append((0, sn))
        elif "ger_executivo" in n:
            candidates.append((1, sn))
        elif "orcamento" in n and "executivo" in n:
            candidates.append((1, sn))
        elif n in ("resumo", "sintese", "síntese", "orcamento sintetico", "apresentacao"):
            candidates.append((2, sn))
        elif "orcamento" in n and len(n) < 30:
            candidates.append((3, sn))
        elif "resumo" in n or "sintese" in n:
            candidates.append((4, sn))
    if not candidates:
        return None
    candidates.sort()
    return candidates[0][1]


def parse_resumo_sheet(ws) -> dict:
    """Parse table with ETAPA | VALOR | % | R$/m² | ... | ÁREA CONSTRUÍDA value."""
    result = {
        "ac_m2": None,
        "total_obra": None,
        "total_gerenciamento": None,
        "linha_gerenciamento": "",
        "linhas_macrogrupo": 0,
        "header_row": None,
        "total_row": None,
    }

    max_row = min(ws.max_row or 0, 200)
    max_col = min(ws.max_column or 0, 20)
    if max_row == 0:
        return result

    # Encontra a linha com "ETAPA" e "VALOR" (header)
    header_row = None
    etapa_col = None
    valor_col = None
    pct_col = None
    rsm2_col = None
    for r in range(1, min(max_row + 1, 30)):
        cells_norm = [norm(ws.cell(row=r, column=c).value) for c in range(1, max_col + 1)]
        # precisa ter etapa (ou macrogrupo) + valor
        has_etapa = any(x in ("etapa", "macrogrupo", "descricao", "item") for x in cells_norm)
        has_valor = any("valor" in x and x for x in cells_norm)
        if has_etapa and has_valor:
            header_row = r
            for idx, v in enumerate(cells_norm, 1):
                if v in ("etapa", "macrogrupo", "descricao", "item"):
                    etapa_col = idx
                elif "valor orcado" in v or v == "valor orçado" or "valor" in v and "total" not in v:
                    if valor_col is None:
                        valor_col = idx
                elif v == "%" or "percent" in v:
                    pct_col = idx
                elif "m²" in v or "m2" in v or "rs/m" in v or "r$/m" in v or "/m²" in v:
                    rsm2_col = idx
            break

    if header_row is None:
        return result
    if etapa_col is None or valor_col is None:
        return result

    result["header_row"] = header_row

    # Varre linhas após o header
    total_row = None
    geren_val = None
    geren_label = ""
    macro_count = 0

    for r in range(header_row + 1, max_row + 1):
        label = ws.cell(row=r, column=etapa_col).value
        val = ws.cell(row=r, column=valor_col).value
        if label is None and val is None:
            # linha vazia — pode ser fim da tabela
            # mas deixa varrer mais um pouco pra achar TOTAL
            continue
        nl = norm(label) if label else ""
        if not nl:
            continue

        # Bloco de AC geralmente aparece em coluna separada (F/G) — vamos varrer tudo
        # linha de total?
        if nl.startswith("total") and "geral" not in nl[:5] and val is not None:
            total_val = to_num(val)
            if total_val and total_val > 100000:
                total_row = r
                result["total_obra"] = total_val
                result["total_row"] = r
                break
        elif "total geral" in nl or nl == "total":
            total_val = to_num(val)
            if total_val and total_val > 100000:
                total_row = r
                result["total_obra"] = total_val
                result["total_row"] = r
                break

        # É macrogrupo?
        num = to_num(val)
        if num and num > 1000:
            macro_count += 1
            # É linha de gerenciamento?
            if any(t in nl for t in GEREN_ROW_TOKENS) and not any(x in nl for x in EXCLUDE_ROW_TOKENS):
                if geren_val is None or num > geren_val:  # pega primeiro ou maior
                    geren_val = num
                    geren_label = str(label)

    result["total_gerenciamento"] = geren_val
    result["linha_gerenciamento"] = geren_label
    result["linhas_macrogrupo"] = macro_count

    # Procura AC em colunas adjacentes (E-J tipicamente)
    # "ÁREA CONSTRUÍDA" em col F, valor em col G
    for r in range(1, min(max_row + 1, 50)):
        for c in range(1, max_col + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            nv = norm(v)
            if any(t in nv for t in AC_TOKENS):
                # procura número à direita
                for dc in range(1, 5):
                    c2 = c + dc
                    if c2 > max_col:
                        break
                    n = to_num(ws.cell(row=r, column=c2).value)
                    if n and 100 < n < 200000:
                        result["ac_m2"] = n
                        break
                if result["ac_m2"]:
                    break
        if result["ac_m2"]:
            break

    return result


def process(path: Path) -> dict:
    rel = path.relative_to(DRIVE_ROOT)
    parts = rel.parts
    cliente = parts[0] if len(parts) >= 1 else ""
    projeto = parts[1] if len(parts) >= 2 else ""

    out = {
        "path": str(rel),
        "cliente": cliente,
        "projeto": projeto,
        "xlsx_name": path.name,
        "aba_usada": "",
        "linha_gerenciamento": "",
        "linhas_macrogrupo": 0,
        "ac_m2": None,
        "total_obra": None,
        "total_gerenciamento": None,
        "pct_gerenciamento": None,
        "rsm2_gerenciamento": None,
        "status": "ok",
        "erro": "",
    }

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except Exception as e:
        out["status"] = "erro_abertura"
        out["erro"] = f"{type(e).__name__}: {str(e)[:100]}"
        return out

    aba = find_ap_sheet(wb)
    if aba is None:
        out["status"] = "sem_aba_resumo"
        out["erro"] = "|".join(wb.sheetnames[:10])
        return out
    out["aba_usada"] = aba

    ws = wb[aba]
    try:
        parsed = parse_resumo_sheet(ws)
    except Exception as e:
        out["status"] = "erro_parse"
        out["erro"] = f"{type(e).__name__}: {str(e)[:100]}"
        return out

    out["ac_m2"] = parsed["ac_m2"]
    out["total_obra"] = parsed["total_obra"]
    out["total_gerenciamento"] = parsed["total_gerenciamento"]
    out["linha_gerenciamento"] = parsed["linha_gerenciamento"]
    out["linhas_macrogrupo"] = parsed["linhas_macrogrupo"]

    if out["total_gerenciamento"] and out["total_obra"]:
        out["pct_gerenciamento"] = out["total_gerenciamento"] / out["total_obra"]
    if out["total_gerenciamento"] and out["ac_m2"]:
        out["rsm2_gerenciamento"] = out["total_gerenciamento"] / out["ac_m2"]

    try:
        wb.close()
    except Exception:
        pass
    return out


def main():
    print(f"[INFO] Drive: {DRIVE_ROOT}", flush=True)

    # Pega apenas xlsx com "apresenta" no nome
    all_xlsx = []
    for p in DRIVE_ROOT.rglob("*.xlsx"):
        if p.name.startswith("~$"):
            continue
        if "apresenta" not in p.name.lower():
            continue
        all_xlsx.append(p)
    print(f"[INFO] {len(all_xlsx)} xlsx Apresentação", flush=True)

    OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    rows = []
    log = []
    t0 = time.time()
    for i, p in enumerate(all_xlsx, 1):
        if i % 10 == 0 or i == 1:
            print(f"[{i}/{len(all_xlsx)}] t={time.time()-t0:.0f}s", flush=True)
        try:
            r = process(p)
        except Exception as e:
            log.append(f"CRASH {p}: {e}")
            continue
        rows.append(r)
        if r["status"] != "ok":
            log.append(f"{r['status']} {p.name}: {r['erro']}")

    cols = list(rows[0].keys()) if rows else []
    with open(OUTPUT_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)

    with open(LOG_TXT, "w", encoding="utf-8") as f:
        f.write("\n".join(log))

    ok = sum(1 for r in rows if r["total_obra"] and r["total_gerenciamento"])
    com_ac = sum(1 for r in rows if r["ac_m2"] and r["total_gerenciamento"])
    print(f"\n[DONE] {len(rows)} linhas. Validos(Ger+GT)={ok}, Com AC={com_ac}", flush=True)
    print(f"[DONE] CSV={OUTPUT_CSV}", flush=True)


if __name__ == "__main__":
    main()

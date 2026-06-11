"""Utilitários compartilhados do pipeline de precificação 3 fontes.

Lógica portada de ~/openclaw/scripts/alfa_precificar_step1_preprocess.py e step5.
Mantida idêntica pra preservar reprodutibilidade do pacote alfa-colinas-claude/05-.

Inclui:
  - norm_text / norm_unit / units_compatible — normalização e compatibilidade
  - WORKSPACE_ROOT — raiz do workspace orcamentos
  - paths helpers — executivos/[slug], tmp dir, projetos
  - schema helpers — leitura padronizada de quantitativos por disciplina
"""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path

WORKSPACE_ROOT = Path(__file__).resolve().parents[2]  # /c/Users/leona/orcamentos
EXECUTIVOS_DIR = WORKSPACE_ROOT / "executivos"
PROJETOS_DIR = WORKSPACE_ROOT / "projetos"

# --- Normalização (idêntica ao step5_cartesian_pool, mais permissiva que step1) ---

UNIT_NORM = {
    "un": "un", "unidade": "un", "unidades": "un", "und": "un",
    "pç": "un", "pc": "un", "peca": "un", "peça": "un",
    "m": "m", "metro": "m", "metros": "m", "ml": "m", "mt": "m",
    "m2": "m2", "m²": "m2",
    "m3": "m3", "m³": "m3",
    "kg": "kg",
    "l": "l", "lt": "l", "litro": "l",
    "vb": "vb",
    "cj": "cj", "conjunto": "cj", "kit": "cj",
    "par": "par",
    "h": "h", "hora": "h", "dia": "dia", "mes": "mes", "mês": "mes",
    "saco": "saco", "sc": "saco",
    "rolo": "rolo", "rl": "rolo",
    "barra": "barra", "br": "barra",
    "t": "t", "ton": "t", "tonelada": "t",
}


def norm_text(s) -> str:
    """Normaliza descrição: minúsculas, sem acento, sem pontuação especial."""
    if not s:
        return ""
    s = str(s).lower()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^a-z0-9\s\.,]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def norm_unit(u) -> str:
    """Normaliza unidade de medida via tabela UNIT_NORM."""
    if not u:
        return ""
    k = str(u).lower().strip()
    k = unicodedata.normalize("NFKD", k)
    k = "".join(c for c in k if not unicodedata.combining(c))
    return UNIT_NORM.get(k, k)


def units_compatible(a: str, b: str) -> bool:
    """un/cj/par compatíveis entre si; m/ml/mt compatíveis. Outras: igualdade."""
    if not a or not b:
        return True  # permissivo quando faltar
    a, b = norm_unit(a), norm_unit(b)
    if a == b:
        return True
    eq_un = {"un", "cj", "par"}
    eq_m = {"m", "ml", "mt"}
    if a in eq_un and b in eq_un:
        return True
    if a in eq_m and b in eq_m:
        return True
    return False


# --- Schemas de quantitativos por disciplina ---
#
# REGRA #3 — schema unificado de 10 cols (idêntico pra todas disciplinas):
#   1=disciplina, 2=grupo, 3=sistema, 4=item, 5=unidade, 6=quantidade,
#   7=pavimento, 8=fonte, 9=status, 10=observacao
#
# Schema legacy do Claude (`alfa-colinas-claude/`) suportado pra regressão histórica:
#   PCI: 8 cols — Sistema | Pavimento | Rotulo | Mult | Item | Unidade | QtdUnit | QtdTotal
#   SAN/TEL: 6 cols — Grupo | Subgrupo | Descrição | Bitola | Unidade | Qtd

SCHEMAS = {
    # REGRA #3 (default daqui pra frente)
    "regra3": {
        "qty_col_idx": 6,         # F — quantidade
        "desc_cols": [4],          # D — item
        "extra_cols": [2, 3, 7],   # B,C,G — grupo, sistema, pavimento (contexto fuzzy)
        "unit_col_idx": 5,         # E — unidade
    },
    # Legacy só pra regressão contra alfa-colinas-claude/
    "pci_legacy": {
        "qty_col_idx": 8,
        "desc_cols": [5],
        "extra_cols": [1, 2],
        "unit_col_idx": 6,
    },
    "santel_legacy": {
        "qty_col_idx": 6,
        "desc_cols": [3, 4],
        "extra_cols": [1, 2],
        "unit_col_idx": 5,
    },
}


def detect_schema(disc_slug: str, xlsx_path: Path | None = None) -> str:
    """Default: 'regra3'. Se a planilha existir e o header[0] não for 'disciplina', cai pra legacy."""
    if xlsx_path is not None and xlsx_path.exists():
        try:
            import openpyxl
            wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
            if "Consolidado" in wb.sheetnames:
                ws = wb["Consolidado"]
                h0 = (next(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0] or "")
                wb.close()
                if isinstance(h0, str) and h0.strip().lower() == "disciplina":
                    return "regra3"
                # Fallback: legacy
                return "pci_legacy" if disc_slug.lower() in ("pci", "incendio", "pci-incendio") else "santel_legacy"
        except Exception:
            pass
    return "regra3"


# --- Paths helpers ---

def executivo_dir(slug: str) -> Path:
    return EXECUTIVOS_DIR / slug


def tmp_dir(slug: str) -> Path:
    """Pasta de intermediários do pipeline (gitignored)."""
    p = EXECUTIVOS_DIR / slug / "_tmp" / "precificacao"
    p.mkdir(parents=True, exist_ok=True)
    return p


def quantitativos_xlsx(slug: str, disc_idx: int, disc_slug: str) -> Path:
    """`executivos/[slug]/{disc_idx:02d}-{disc_slug}/quantitativos-{disc_slug}.xlsx`."""
    return EXECUTIVOS_DIR / slug / f"{disc_idx:02d}-{disc_slug}" / f"quantitativos-{disc_slug}.xlsx"


def discover_disciplinas(slug: str) -> list[tuple[int, str]]:
    """Lê executivos/[slug]/ e retorna [(idx, disc_slug), ...] ordenado.

    Procura subpastas numeradas tipo `01-pci`, `02-sanitario`, etc — exclui `00-projeto`
    e qualquer `XX-precificacao*`. Cada uma deve ter `quantitativos-*.xlsx` válido.
    """
    base = executivo_dir(slug)
    out = []
    if not base.exists():
        return out
    for child in sorted(base.iterdir()):
        if not child.is_dir():
            continue
        m = re.match(r"^(\d{2})-([a-z0-9\-]+)$", child.name)
        if not m:
            continue
        idx, disc = int(m.group(1)), m.group(2)
        if disc in ("projeto",) or "precificacao" in disc:
            continue
        qxlsx = child / f"quantitativos-{disc}.xlsx"
        if qxlsx.exists():
            out.append((idx, disc))
    return out


def next_precificacao_idx(slug: str) -> int:
    """Próximo número de pasta após a última disciplina (alfa: 04, bela vida: 06)."""
    discs = discover_disciplinas(slug)
    if not discs:
        return 1
    return max(idx for idx, _ in discs) + 1


def precificacao_out_dir(slug: str, source_tag: str = "banco-de-dados") -> Path:
    """`executivos/[slug]/XX-precificacao-{source_tag}/` — cria se não existir."""
    idx = next_precificacao_idx(slug)
    name = f"{idx:02d}-precificacao-{source_tag}"
    p = executivo_dir(slug) / name
    p.mkdir(parents=True, exist_ok=True)
    return p

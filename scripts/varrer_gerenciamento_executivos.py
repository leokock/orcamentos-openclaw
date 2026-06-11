#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Varrer todos os orçamentos executivos entregues pra extrair o macrogrupo
Gerenciamento (ou equivalentes: CI, Canteiro e Indiretos, Adm, BDI Direto).

Gera CSV com: cliente, projeto, xlsx, abas, total_gerenciamento, grand_total,
ac_m2, pct_gerenciamento, rsm2_gerenciamento.

Uso:
    PYTHONIOENCODING=utf-8 py -3.10 -X utf8 scripts/varrer_gerenciamento_executivos.py
"""
from __future__ import annotations

import csv
import io
import os
import re
import sys
import time
import traceback
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

# Forçar UTF-8 em stdout (Windows cp1252 crasha com acentos)
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

DRIVE_ROOT = Path(
    "/g/Drives compartilhados/03 CTN Projetos/2. Projetos em Andamento/_Entregas/Orçamento_executivo"
)
# No Windows, reescrever pra G:\
if os.name == "nt":
    DRIVE_ROOT = Path(
        r"G:\Drives compartilhados\03 CTN Projetos\2. Projetos em Andamento\_Entregas\Orçamento_executivo"
    )

OUTPUT_CSV = Path(r"C:\Users\leona\orcamentos-openclaw\base\executivos-gerenciamento-varredura-2026-04-22.csv")
LOG_TXT = Path(r"C:\Users\leona\orcamentos-openclaw\base\executivos-gerenciamento-varredura-2026-04-22.log.txt")

# Tokens pra identificar a aba de Gerenciamento (normalizado: lowercase, sem acento)
GEREN_TOKENS = [
    "gerenciamento", "geren",
    "canteiro e indiretos", "canteiro_indiretos", "canteiro",
    "indiretos", "despesas indiretas", "desp indiretas", "desp_indiretas",
    "administracao da obra", "adm da obra", "adm obra", "administracao",
    "bdi direto", "ci ", " ci", "custos indiretos",
]

# Tokens pra identificar aba de resumo/grand total
RESUMO_TOKENS = [
    "resumo", "sintese", "sintético", "sintetico",
    "total geral", "orcamento sintetico", "capa",
    "apresentacao", "orcamento", "orçamento sintetico"
]


def normalize(s: str) -> str:
    """Lowercase + remove acentos comuns."""
    if s is None:
        return ""
    s = str(s).lower().strip()
    trans = str.maketrans("áàâãäéèêëíìîïóòôõöúùûüç", "aaaaaeeeeiiiiooooouuuuc")
    return s.translate(trans)


def is_geren_sheet(name: str) -> bool:
    n = normalize(name)
    # exclui abas que contém memorial, memoria, etc
    if "memorial" in n or "memoria" in n or "justific" in n:
        return False
    for t in GEREN_TOKENS:
        if t in n:
            return True
    return False


def is_resumo_sheet(name: str) -> bool:
    n = normalize(name)
    for t in RESUMO_TOKENS:
        if t in n:
            return True
    return False


def to_number(v) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        if v == 0:
            return None
        return float(v)
    s = str(v).strip()
    if not s:
        return None
    # Remove R$, espaços
    s = re.sub(r"[R\$\s]", "", s)
    # Se tem vírgula e ponto, assumir ponto = milhar, vírgula = decimal
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        n = float(s)
        return n if n != 0 else None
    except ValueError:
        return None


def find_total_in_sheet(ws) -> Optional[float]:
    """Encontra o maior valor numa linha que tem 'TOTAL' em uma das primeiras colunas.
    Retorna None se não achar."""
    max_total = None
    max_row = ws.max_row or 0
    max_col = min(ws.max_column or 0, 30)
    if max_row == 0:
        return None
    # Varre do fim pro começo procurando linha com TOTAL
    for row in range(max_row, max(0, max_row - 500), -1):
        # Pega as primeiras ~5 cols de texto
        for col in range(1, min(6, max_col + 1)):
            cell = ws.cell(row=row, column=col)
            val = cell.value
            if val is None:
                continue
            vs = normalize(str(val))
            if "total" in vs and not any(x in vs for x in ["subtotal", "por total", "pre total"]):
                # Encontrou - procura o maior número na mesma linha
                for c2 in range(col + 1, max_col + 1):
                    v2 = ws.cell(row=row, column=c2).value
                    n = to_number(v2)
                    if n is not None and n > 1000:  # ignora pequenos
                        if max_total is None or n > max_total:
                            max_total = n
                if max_total is not None:
                    return max_total
    return max_total


def find_grand_total_from_resumo(wb) -> Optional[float]:
    """Procura aba de resumo e extrai o maior total (provavelmente TOTAL GERAL)."""
    candidates = []
    for sheet_name in wb.sheetnames:
        if is_resumo_sheet(sheet_name):
            try:
                ws = wb[sheet_name]
            except Exception:
                continue
            t = find_total_in_sheet(ws)
            if t is not None:
                candidates.append(t)
    if not candidates:
        return None
    # Grand total geralmente é o maior
    return max(candidates)


def find_ac_m2(wb) -> Optional[float]:
    """Procura 'área construída', 'AC', 'ACT' no workbook."""
    patterns = [
        re.compile(r"area\s+construida", re.I),
        re.compile(r"\bact\b", re.I),
        re.compile(r"\bac\s*[=:]\s*", re.I),
        re.compile(r"area\s+total\s+construida", re.I),
    ]
    for sheet_name in wb.sheetnames:
        nm = normalize(sheet_name)
        if not any(t in nm for t in ["resumo", "capa", "sintese", "apresentacao", "dados", "briefing", "premissas"]):
            continue
        try:
            ws = wb[sheet_name]
        except Exception:
            continue
        max_row = min(ws.max_row or 0, 200)
        max_col = min(ws.max_column or 0, 15)
        for row in range(1, max_row + 1):
            for col in range(1, max_col + 1):
                v = ws.cell(row=row, column=col).value
                if v is None:
                    continue
                s = str(v)
                if any(p.search(s) for p in patterns):
                    # Procura número na mesma linha ou próxima célula
                    for dr in range(0, 3):
                        for dc in range(1, 5):
                            v2 = ws.cell(row=row + dr, column=col + dc).value
                            n = to_number(v2)
                            if n is not None and 100 < n < 100000:
                                return n
    return None


def process_xlsx(path: Path) -> dict:
    result = {
        "path": str(path),
        "cliente": "",
        "projeto": "",
        "xlsx_name": path.name,
        "abas_encontradas": "",
        "aba_gerenciamento_usada": "",
        "total_gerenciamento": None,
        "grand_total": None,
        "ac_m2": None,
        "pct_gerenciamento": None,
        "rsm2_gerenciamento": None,
        "status": "ok",
        "erro": "",
    }
    # Extrair cliente/projeto do path relativo
    try:
        rel = path.relative_to(DRIVE_ROOT)
        parts = rel.parts
        if len(parts) >= 1:
            result["cliente"] = parts[0]
        if len(parts) >= 2:
            result["projeto"] = parts[1]
    except Exception:
        pass

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except (InvalidFileException, KeyError) as e:
        result["status"] = "erro_abertura"
        result["erro"] = f"{type(e).__name__}: {str(e)[:100]}"
        return result
    except Exception as e:
        result["status"] = "erro_abertura"
        result["erro"] = f"{type(e).__name__}: {str(e)[:100]}"
        return result

    result["abas_encontradas"] = "|".join(wb.sheetnames[:20])

    # Procura aba de gerenciamento
    geren_candidates = []
    for sn in wb.sheetnames:
        if is_geren_sheet(sn):
            try:
                ws = wb[sn]
                t = find_total_in_sheet(ws)
                if t is not None:
                    geren_candidates.append((sn, t))
            except Exception:
                continue

    if geren_candidates:
        # Ordenar por valor decrescente (pega o maior - provavelmente total da aba, não subitem)
        geren_candidates.sort(key=lambda x: -x[1])
        result["aba_gerenciamento_usada"] = geren_candidates[0][0]
        result["total_gerenciamento"] = geren_candidates[0][1]

    # Grand total
    result["grand_total"] = find_grand_total_from_resumo(wb)

    # AC
    result["ac_m2"] = find_ac_m2(wb)

    # Derivados
    if result["total_gerenciamento"] and result["grand_total"]:
        result["pct_gerenciamento"] = result["total_gerenciamento"] / result["grand_total"]
    if result["total_gerenciamento"] and result["ac_m2"]:
        result["rsm2_gerenciamento"] = result["total_gerenciamento"] / result["ac_m2"]

    try:
        wb.close()
    except Exception:
        pass

    return result


def main():
    print(f"[INFO] DRIVE_ROOT = {DRIVE_ROOT}", flush=True)
    print(f"[INFO] Existe? {DRIVE_ROOT.exists()}", flush=True)

    # Listar todos xlsx
    all_xlsx = []
    for p in DRIVE_ROOT.rglob("*.xlsx"):
        if p.name.startswith("~$"):
            continue
        # Filtrar: descartar nomes com "memorial", "diagnose"
        nm_low = p.name.lower()
        if any(t in nm_low for t in ["memorial", "diagnose", "diagnostic"]):
            continue
        all_xlsx.append(p)

    print(f"[INFO] {len(all_xlsx)} xlsx encontrados (após filtro)", flush=True)

    OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    log_lines = []

    rows = []
    t_start = time.time()
    for i, path in enumerate(all_xlsx, 1):
        if i % 10 == 0 or i == 1:
            elapsed = time.time() - t_start
            print(f"[{i}/{len(all_xlsx)}] t={elapsed:.0f}s - {path.name[:60]}", flush=True)
        try:
            t0 = time.time()
            r = process_xlsx(path)
            dt = time.time() - t0
            if dt > 30:
                log_lines.append(f"SLOW {dt:.1f}s: {path}")
            rows.append(r)
            if r["status"] != "ok":
                log_lines.append(f"{r['status']} {path}: {r['erro']}")
        except Exception as e:
            log_lines.append(f"CRASH {path}: {type(e).__name__}: {e}")
            log_lines.append(traceback.format_exc()[:500])
            rows.append({
                "path": str(path),
                "cliente": path.parts[-3] if len(path.parts) >= 3 else "",
                "projeto": path.parts[-2] if len(path.parts) >= 2 else "",
                "xlsx_name": path.name,
                "abas_encontradas": "",
                "aba_gerenciamento_usada": "",
                "total_gerenciamento": None,
                "grand_total": None,
                "ac_m2": None,
                "pct_gerenciamento": None,
                "rsm2_gerenciamento": None,
                "status": "crash",
                "erro": str(e)[:100],
            })

    # CSV
    cols = [
        "path", "cliente", "projeto", "xlsx_name",
        "abas_encontradas", "aba_gerenciamento_usada",
        "total_gerenciamento", "grand_total", "ac_m2",
        "pct_gerenciamento", "rsm2_gerenciamento",
        "status", "erro"
    ]
    with open(OUTPUT_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow(r)

    with open(LOG_TXT, "w", encoding="utf-8") as f:
        f.write("\n".join(log_lines))

    print(f"\n[DONE] {len(rows)} linhas em {OUTPUT_CSV}", flush=True)
    print(f"[DONE] log em {LOG_TXT}", flush=True)

    # Resumo rápido
    ok = sum(1 for r in rows if r["status"] == "ok")
    com_geren = sum(1 for r in rows if r["total_gerenciamento"])
    com_grand = sum(1 for r in rows if r["grand_total"])
    com_ac = sum(1 for r in rows if r["ac_m2"])
    com_pct = sum(1 for r in rows if r["pct_gerenciamento"])
    print(f"[RESUMO] ok={ok} com_geren={com_geren} com_grand_total={com_grand} com_ac={com_ac} com_pct={com_pct}", flush=True)


if __name__ == "__main__":
    main()

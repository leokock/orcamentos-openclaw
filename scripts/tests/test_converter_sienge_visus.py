import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
import converter_sienge_visus as c
import openpyxl

def test_derive_tipo():
    assert c.derive_tipo("01 - Mão-de-obra") == "Mao_obra"
    assert c.derive_tipo("02 - Materiais") == "Material"
    assert c.derive_tipo("03 - Equipamentos") == "Equipamento"
    assert c.derive_tipo("04 - Empreitada") == "Terceirizados"
    assert c.derive_tipo("05 - Verbas") == "Outro"
    assert c.derive_tipo("06 - Material de Expediente") == "Material"
    assert c.derive_tipo("99 - Desconhecido") == "Outro"   # fallback

def test_parse_ativo():
    assert c.parse_ativo("Sim") is True
    assert c.parse_ativo("Não") is False
    assert c.parse_ativo(None) is True   # default ativo

def test_normalize_desc():
    assert c.normalize_desc("  AÇO 6.3mm ") == "aco 6.3mm"
    assert c.normalize_desc("Metro Quadrado") == "metro quadrado"

def test_to_float():
    assert c.to_float("4,2") == 4.2
    assert c.to_float(5.5) == 5.5
    assert c.to_float("-") is None
    assert c.to_float("") is None
    assert c.to_float(None) is None

import pytest
DATA_DIR = r"G:\Drives compartilhados\03 CTN Projetos\2. Projetos em Andamento\RDO Empreendimentos\Villa Vauban\04. Custo\03.4 Custo - Fechamento\Importação Insumos e Composições Visus"
skip_no_data = pytest.mark.skipif(not os.path.isdir(DATA_DIR), reason="Drive folder unavailable")

@skip_no_data
def test_parse_insumos_counts():
    ins = c.parse_insumos(os.path.join(DATA_DIR, "insumos.xlsx"))
    assert len(ins) == 1340
    codes = [i.codigo for i in ins]
    assert len(codes) == len(set(codes)), "códigos de insumo duplicados"
    servente = next(i for i in ins if i.codigo == 1)
    assert servente.tipo == "Mao_obra"
    assert servente.unidade == "h"
    assert servente.valor == 4.2
    tipos = {i.tipo for i in ins}
    assert {"Mao_obra", "Material", "Equipamento"} <= tipos

@skip_no_data
def test_write_insumos_roundtrip(tmp_path):
    ins = c.parse_insumos(os.path.join(DATA_DIR, "insumos.xlsx"))
    out = str(tmp_path / "OUT_INSUMOS.xlsx")
    template = os.path.join(DATA_DIR, "Thozen - Electra_INSUMOS.xlsx")
    c.write_insumos(ins, template, out)
    wb = openpyxl.load_workbook(out, read_only=True, data_only=True)
    ws = wb.active
    rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0] is not None]
    assert len(rows) == 1340
    assert ws["A1"].value == "<codigo>"   # header preserved
    first = rows[0]
    assert first[0] == 1 and first[5] == "Mao_obra"   # codigo, tipo
    wb.close()

@skip_no_data
def test_parse_servicos():
    srv = c.parse_servicos(os.path.join(DATA_DIR, "serviços.xlsx"))
    assert len(srv) == 440
    s = srv[3315]
    assert s["unidade"] == "m3"
    assert round(s["total"], 2) == 407.11

@skip_no_data
def test_parse_composicoes():
    comps = c.parse_composicoes(os.path.join(DATA_DIR, "composições serviços.xlsx"))
    assert len(comps) == 308
    multi = next(x for x in comps if x.codigo == 5893)
    assert len(multi.itens) == 3
    assert multi.itens[0].codigo_insumo == 1116
    assert multi.itens[0].multiplicador == 1.609
    assert multi.grupo_servico

@skip_no_data
def test_build_catalog():
    ins = c.parse_insumos(os.path.join(DATA_DIR, "insumos.xlsx"))
    srv = c.parse_servicos(os.path.join(DATA_DIR, "serviços.xlsx"))
    comps = c.parse_composicoes(os.path.join(DATA_DIR, "composições serviços.xlsx"))
    insumos2, comps2, report = c.build_catalog(ins, srv, comps)
    assert len(comps2) == 440
    valid = {i.codigo for i in insumos2}
    for cp in comps2:
        for it in cp.itens:
            assert it.codigo_insumo in valid            # no broken FK
    assert len(insumos2) > len(ins)
    assert report["lump_sum"] == 132
    assert report["variants"] >= 1
    # FIDELITY: every detailed composition reproduces the Sienge total
    assert report["sanity_fail"] == 0, f"{report['sanity_fail']} composições divergem do Sienge"
    # no duplicate códigos in the final insumo catalog
    codes = [i.codigo for i in insumos2]
    assert len(codes) == len(set(codes))
    # umbrella code 1562 (Aço) must be split into per-bitola insumos
    descs = [i.descricao for i in insumos2]
    assert any("6,3mm" in d for d in descs), "faltou insumo aço bitola 6,3mm"
    assert any("8mm" in d for d in descs), "faltou insumo aço bitola 8mm"
    assert report["umbrella_codes"] >= 10
    assert report["sanity_fail"] == 0

@skip_no_data
def test_write_composicoes(tmp_path):
    ins = c.parse_insumos(os.path.join(DATA_DIR, "insumos.xlsx"))
    srv = c.parse_servicos(os.path.join(DATA_DIR, "serviços.xlsx"))
    comps = c.parse_composicoes(os.path.join(DATA_DIR, "composições serviços.xlsx"))
    insumos2, comps2, _ = c.build_catalog(ins, srv, comps)
    out = str(tmp_path / "OUT_COMP.xlsx")
    template = os.path.join(DATA_DIR, "Thozen - Electra_COMPOSIÇÕES.xlsx")
    c.write_composicoes(comps2, {i.codigo: i for i in insumos2}, template, out)
    wb = openpyxl.load_workbook(out, read_only=True, data_only=True)
    ws = wb.active
    rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0] is not None]
    distinct = {r[0] for r in rows}
    assert len(distinct) == 440
    assert ws["A1"].value == "<codigo>"
    sample = rows[0]
    assert sample[13] == "Villa Vauban"   # col N tabela_insumo
    assert sample[14] == "INSUMO"          # col O tipo_associacao
    wb.close()

def test_norm_unit():
    assert c.norm_unit("Metro Quadrado") == c.norm_unit("m²")
    assert c.norm_unit("Unidade") == c.norm_unit("un")
    assert c.norm_unit("Tonelada") != c.norm_unit("kg")

def test_best_match():
    insumos = [c.Insumo(10, "ACO 6.3MM", "kg", 5.0, "Material"),
               c.Insumo(11, "PORCELANATO ACETINADO", "m2", 50.0, "Material")]
    code, score = c.best_match("Aço 6.3mm", insumos)
    assert code == 10 and score > 0.8

@skip_no_data
def test_parse_fornecedores():
    forn = c.parse_fornecedores(os.path.join(DATA_DIR, "CTN&RDO -Fornecedores-Insumos.xlsx"))
    assert len(forn) >= 50
    assert any(f["item"].startswith("Parede diafragma") for f in forn)

def test_aplicar_de_para(tmp_path):
    insumos = [c.Insumo(10, "ACO 6.3MM", "kg", 5.0, "Material"),
               c.Insumo(11, "PORCELANATO", "m2", 50.0, "Material")]
    dp = str(tmp_path / "dp.xlsx")
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["item_fornecedor","desc_insumo_candidato","codigo_insumo","score",
               "preco_atual","preco_novo","unid_fornecedor","unid_insumo",
               "flag_unidade_divergente","fornecedor","aplicar?"])
    ws.append(["Aço 6.3mm","ACO 6.3MM",10,0.95,5.0,6.5,"kg","kg","","Giassi","S"])
    ws.append(["Porcelanato","PORCELANATO",11,0.9,50.0,60.0,"m2","m2","","Eliane","N"])
    wb.save(dp); wb.close()
    n = c.aplicar_de_para(dp, insumos)
    assert n == 1
    by = {i.codigo: i for i in insumos}
    assert by[10].valor == 6.5 and by[10].fonte_preco == "Giassi"
    assert by[11].valor == 50.0   # not applied (aplicar?=N)

@skip_no_data
def test_headers_validated_real_files():
    # should NOT raise on the real files
    c.parse_insumos(os.path.join(DATA_DIR, "insumos.xlsx"))
    c.parse_servicos(os.path.join(DATA_DIR, "serviços.xlsx"))
    c.parse_composicoes(os.path.join(DATA_DIR, "composições serviços.xlsx"))

def test_assert_header_raises():
    import pytest as _pt
    with _pt.raises(ValueError):
        c._assert_header(["WRONG", "x"], [(0, "Código")], "teste")

def test_gerar_de_para_gating(tmp_path):
    insumos = [c.Insumo(10, "ACO 6.3MM", "kg", 5.0, "Material"),
               c.Insumo(11, "PORCELANATO ACETINADO", "m2", 50.0, "Material")]
    forn = [
        {"item": "ACO 6.3MM", "custo": 6.5, "unidade": "kg", "fornecedor": "Giassi"},      # match, unit ok, price ok -> S
        {"item": "PORCELANATO ACETINADO", "custo": 60.0, "unidade": "Tonelada", "fornecedor": "Eliane"},  # unit divergent -> N + SIM
        {"item": "ACO 6.3MM", "custo": None, "unidade": "kg", "fornecedor": "X"},           # no price -> N
    ]
    out = str(tmp_path / "dp.xlsx")
    c.gerar_de_para(forn, insumos, out)
    import openpyxl as ox
    ws = ox.load_workbook(out, data_only=True).active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert rows[0][10] == "S"
    assert rows[1][10] == "N" and rows[1][8] == "SIM"
    assert rows[2][10] == "N"

def test_match_norm_numbers():
    # comma/dot in gauges unify; punctuation dropped
    assert "6.3mm" in c._match_norm("Aço 6,3mm")
    assert "6.3mm" in c._match_norm("Aço / CA - 50  6.3mm")

def test_best_match_bitola():
    insumos = [c.Insumo(10, "Aço / CA - 50  6,3mm", "kg", 5.0, "Material"),
               c.Insumo(11, "Aço / CA - 50  8mm", "kg", 4.9, "Material"),
               c.Insumo(12, "União PPR 63mm", "un", 8.5, "Material")]
    code, score = c.best_match("Aço 8mm", insumos)
    assert code == 11, f"esperava insumo 8mm, veio {code}"

def test_is_apply():
    assert c._is_apply("Sim") and c._is_apply("S") and c._is_apply("sim ")
    assert not c._is_apply("N") and not c._is_apply("") and not c._is_apply(None)

def test_gerar_de_para_preserva_edicoes(tmp_path):
    insumos = [c.Insumo(10, "ACO 6.3MM", "kg", 5.0, "Material"),
               c.Insumo(11, "PORCELANATO", "m2", 50.0, "Material")]
    forn = [{"item": "Porcelanato", "custo": 60.0, "unidade": "Tonelada", "fornecedor": "Eliane"}]
    out = str(tmp_path / "dp.xlsx")
    # 1ª geração: unidade divergente -> auto "N"
    c.gerar_de_para(forn, insumos, out)
    import openpyxl as ox
    ws = ox.load_workbook(out, data_only=True).active
    assert list(ws.iter_rows(min_row=2, values_only=True))[0][10] == "N"
    # usuário marca "Sim" e salva
    wb = ox.load_workbook(out); wb.active.cell(2, 11, "Sim"); wb.save(out); wb.close()
    # 2ª geração preservando edições
    c.gerar_de_para(forn, insumos, out, prev_path=out)
    ws = ox.load_workbook(out, data_only=True).active
    assert list(ws.iter_rows(min_row=2, values_only=True))[0][10] == "S", "edição do usuário não preservada"

def test_gerar_de_para_exclui_lumpsum(tmp_path):
    insumos = [c.Insumo(10, "Orçamento de Obra", "vb", 6276.18, "Outro", fonte_preco="SIENGE-SERVICO"),
               c.Insumo(11, "SIKA TOP FLEX FIBRAS", "cx", 110.0, "Material")]
    forn = [{"item": "SIKA TOP FLEX FIBRAS", "custo": 110.0, "unidade": "Caixa", "fornecedor": "SIKA"}]
    out = str(tmp_path / "dp.xlsx")
    c.gerar_de_para(forn, insumos, out)
    import openpyxl as ox
    row = list(ox.load_workbook(out, data_only=True).active.iter_rows(min_row=2, values_only=True))[0]
    assert row[2] == 11, "deveria casar com o insumo real, não a verba lump-sum"

def test_gerar_relatorio(tmp_path):
    report = {"detalhadas": 308, "lump_sum": 132, "placeholders": 0,
              "price_overrides": 31, "variants": 36,
              "sanity": [(5893, 87.21, 87.21, 0.0), (6334, 571.47, 571.47, 0.0)],
              "sanity_fail": 0}
    out = str(tmp_path / "rel.md")
    c.gerar_relatorio(report, n_insumos=1520, n_comps=440, out_path=out)
    txt = open(out, encoding="utf-8").read()
    assert "1520" in txt and "440" in txt
    assert "308" in txt and "132" in txt
    assert "Sanity check" in txt

def test_parse_extras_missing(tmp_path):
    assert c.parse_extras(str(tmp_path / "nope.csv")) == []

def test_parse_extras_reads(tmp_path):
    p = tmp_path / "insumos-extras.csv"
    p.write_text("descricao,unidade,valor,tipo,fonte_preco\n"
                 "Madeira Caixaria Mista,m3,810,Material,MADEIREIRA OTTO\n", encoding="utf-8")
    ex = c.parse_extras(str(p))
    assert len(ex) == 1 and ex[0]["descricao"] == "Madeira Caixaria Mista"

@skip_no_data
def test_build_catalog_extras():
    ins = c.parse_insumos(os.path.join(DATA_DIR, "insumos.xlsx"))
    srv = c.parse_servicos(os.path.join(DATA_DIR, "serviços.xlsx"))
    comps = c.parse_composicoes(os.path.join(DATA_DIR, "composições serviços.xlsx"))
    extras = [{"descricao": "Madeira Caixaria Mista", "unidade": "m3", "valor": "810", "tipo": "Material", "fonte_preco": "MADEIREIRA OTTO"}]
    base_n = len(c.build_catalog(ins, srv, comps)[0])
    # re-parse: build_catalog mutates comps (remaps cp.itens) — fresh inputs avoid the reuse artifact
    ins = c.parse_insumos(os.path.join(DATA_DIR, "insumos.xlsx"))
    comps = c.parse_composicoes(os.path.join(DATA_DIR, "composições serviços.xlsx"))
    insumos2, comps2, report = c.build_catalog(ins, srv, comps, extras=extras)
    assert report["extras"] == 1
    assert len(insumos2) == base_n + 1
    m = [i for i in insumos2 if i.descricao == "Madeira Caixaria Mista"]
    assert len(m) == 1 and m[0].valor == 810.0 and m[0].fonte_preco == "MADEIREIRA OTTO"
    # códigos únicos preservados
    codes = [i.codigo for i in insumos2]
    assert len(codes) == len(set(codes))

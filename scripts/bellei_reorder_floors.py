"""
Reordena pavimentos nas abas do BSR_Extracao_Quantidades.xlsx:
1. Inverte ordem: Reservatório no topo, Subsolo na última linha
2. Adiciona somatório por tipo a partir da linha 40 (com fórmulas SUMIF)
3. Mantém linha de TOTAL com fórmula SUM
"""

import openpyxl
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from copy import copy
import re
import shutil
import os

INPUT = r'planejamento/OBRA - BELLEI/BSR_Extracao_Quantidades.xlsx'
OUTPUT = r'planejamento/OBRA - BELLEI/BSR_Extracao_Quantidades_v2.xlsx'

# Copia o arquivo para não sobrescrever o original
shutil.copy2(INPUT, OUTPUT)

wb = load_workbook(OUTPUT)

# ─── Ordem canônica dos pavimentos (topo → base) ───────────────────────────
# Sheets com nomenclatura G1/G2 (alvenarias, instalações, etc.)
CANONICAL_G = [
    'Reservatório',
    'Barrilete',
    'Cobertura',
    '24° ANDAR', '23° ANDAR', '22° ANDAR', '21° ANDAR', '20° ANDAR',
    '19° ANDAR', '18° ANDAR', '17° ANDAR',
    'Garden',
    '15° ANDAR', '14° ANDAR', '13° ANDAR', '12° ANDAR', '11° ANDAR',
    '10° ANDAR', '9° ANDAR', '8° ANDAR', '7° ANDAR', '6° ANDAR',
    'Lazer',
    'G4', 'G3', 'G2', 'G1',
    'Térreo',
    'Subsolo',
]

# Sheets SUPRAESTRUTURA usa numeração sequencial 1-27
CANONICAL_S = (
    ['SUBSOLO', 'TÉRREO'] +  # serão invertidos para ficarem no final
    [f'{i}\u00b0 ANDAR' for i in range(1, 28)]
)
# Ordem final SUPRAESTRUTURA (topo → base)
CANONICAL_S_ORDERED = (
    [f'{i}\u00b0 ANDAR' for i in range(27, 0, -1)] +
    ['TÉRREO', 'SUBSOLO']
)

# ─── Tipos de pavimento para o somatório ────────────────────────────────────
# Para sheets G1/G2:
TIPOS_G = [
    ('Subsolo',      ['Subsolo']),
    ('Térreo',       ['Térreo']),
    ('G1',           ['G1']),
    ('G2',           ['G2']),
    ('G3',           ['G3']),
    ('G4',           ['G4']),
    ('Lazer',        ['Lazer']),
    ('Tipo A (x5)',  ['6\u00b0 ANDAR', '8\u00b0 ANDAR', '10\u00b0 ANDAR', '12\u00b0 ANDAR', '14\u00b0 ANDAR']),
    ('Tipo B (x5)',  ['7\u00b0 ANDAR', '9\u00b0 ANDAR', '11\u00b0 ANDAR', '13\u00b0 ANDAR', '15\u00b0 ANDAR']),
    ('Garden',       ['Garden']),
    ('Tipo C (x4)',  ['17\u00b0 ANDAR', '19\u00b0 ANDAR', '21\u00b0 ANDAR', '23\u00b0 ANDAR']),
    ('Tipo D (x4)',  ['18\u00b0 ANDAR', '20\u00b0 ANDAR', '22\u00b0 ANDAR', '24\u00b0 ANDAR']),
    ('Cobertura',    ['Cobertura']),
    ('Barrilete',    ['Barrilete']),
    ('Reservatório', ['Reservatório']),
]

# Para SUPRAESTRUTURA (numeração sequencial):
TIPOS_S = [
    ('Subsolo',      ['SUBSOLO']),
    ('Térreo',       ['TÉRREO']),
    ('G1 (1° Pav)',  ['1\u00b0 ANDAR']),
    ('G2 (2° Pav)',  ['2\u00b0 ANDAR']),
    ('G3 (3° Pav)',  ['3\u00b0 ANDAR']),
    ('G4 (4° Pav)',  ['4\u00b0 ANDAR']),
    ('Lazer (5° Pav)', ['5\u00b0 ANDAR']),
    ('Tipo A (x5)',  ['6\u00b0 ANDAR', '8\u00b0 ANDAR', '10\u00b0 ANDAR', '12\u00b0 ANDAR', '14\u00b0 ANDAR']),
    ('Tipo B (x5)',  ['7\u00b0 ANDAR', '9\u00b0 ANDAR', '11\u00b0 ANDAR', '13\u00b0 ANDAR', '15\u00b0 ANDAR']),
    ('Garden (16° Pav)', ['16\u00b0 ANDAR']),
    ('Tipo C (x4)',  ['17\u00b0 ANDAR', '19\u00b0 ANDAR', '21\u00b0 ANDAR', '23\u00b0 ANDAR']),
    ('Tipo D (x4)',  ['18\u00b0 ANDAR', '20\u00b0 ANDAR', '22\u00b0 ANDAR', '24\u00b0 ANDAR']),
    ('Cobertura (25° Pav)', ['25\u00b0 ANDAR']),
    ('Barrilete (26° Pav)', ['26\u00b0 ANDAR']),
    ('Reservatório (27° Pav)', ['27\u00b0 ANDAR']),
]

def copy_cell_style(src, dst):
    """Copia estilo de uma célula para outra."""
    if src.has_style:
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)

def get_row_data(ws, row_idx):
    """Retorna lista de (valor, célula) para cada coluna da linha."""
    return [(ws.cell(row_idx, col).value, ws.cell(row_idx, col)) 
            for col in range(1, ws.max_column + 1)]

def write_row_data(ws, row_idx, row_data, copy_styles=True):
    """Escreve lista de (valor, src_cell) em uma linha."""
    for col_idx, (val, src_cell) in enumerate(row_data, 1):
        dst = ws.cell(row_idx, col_idx)
        dst.value = val
        if copy_styles and src_cell is not None:
            copy_cell_style(src_cell, dst)

def normalize_name(name):
    """Normaliza nome do pavimento para comparação."""
    return str(name).strip().upper() if name else ''

def process_sheet(ws, canonical_order, tipos):
    """Processa uma sheet: reordena pavimentos e adiciona somatório."""
    max_col = ws.max_column
    
    # ── Ler todas as linhas de dados (row 5 em diante) ──────────────────────
    data_start = 5
    
    # Identificar linha TOTAL e linha de fim dos dados
    total_row_idx = None
    last_data_row = data_start - 1
    
    all_rows = {}  # row_idx → (floor_name, row_data)
    
    for r in range(data_start, ws.max_row + 1):
        val = ws.cell(r, 1).value
        if val is None:
            continue
        row_data = get_row_data(ws, r)
        if str(val).strip().upper() == 'TOTAL':
            total_row_idx = r
        else:
            all_rows[r] = (val, row_data)
            last_data_row = r
    
    if not all_rows:
        return
    
    # ── Separar pavimentos "canônicos" de linhas especiais ──────────────────
    canonical_upper = {normalize_name(n): n for n in canonical_order}
    
    floor_rows = {}    # floor_name → row_data
    special_rows = []  # list of (floor_name, row_data) - ordem original
    
    for r, (name, row_data) in sorted(all_rows.items()):
        if normalize_name(name) in canonical_upper:
            floor_rows[normalize_name(name)] = (name, row_data)
        else:
            special_rows.append((name, row_data))
    
    # ── Definir nova ordem dos pavimentos ───────────────────────────────────
    ordered_floors = []
    for canonical_name in canonical_order:
        key = normalize_name(canonical_name)
        if key in floor_rows:
            ordered_floors.append(floor_rows[key])
    
    # Pavimentos que não estão no canônico mas existem na sheet
    for r, (name, row_data) in sorted(all_rows.items()):
        if normalize_name(name) not in canonical_upper and (name, row_data) not in special_rows:
            special_rows.append((name, row_data))
    
    # ── Reescrever linhas na nova ordem ─────────────────────────────────────
    write_row = data_start
    
    # Pavimentos em ordem invertida
    floor_name_to_new_row = {}  # para as fórmulas SUMIF depois
    for name, row_data in ordered_floors:
        write_row_data(ws, write_row, row_data, copy_styles=True)
        floor_name_to_new_row[normalize_name(name)] = write_row
        write_row += 1
    
    # Linhas especiais (não-pavimento)
    for name, row_data in special_rows:
        write_row_data(ws, write_row, row_data, copy_styles=True)
        write_row += 1
    
    last_data_new = write_row - 1
    
    # ── Limpar linhas que ficaram sobrando (se a sheet encolheu) ────────────
    if total_row_idx and total_row_idx > last_data_new + 1:
        for r in range(last_data_new + 1, total_row_idx):
            for col in range(1, max_col + 1):
                ws.cell(r, col).value = None
    
    # ── Escrever linha TOTAL logo abaixo dos dados ──────────────────────────
    total_row_new = last_data_new + 1
    
    # Estilo da linha TOTAL (copiar do total original se existia)
    total_style_row = total_row_idx if total_row_idx else None
    
    ws.cell(total_row_new, 1).value = 'TOTAL'
    if total_style_row:
        copy_cell_style(ws.cell(total_style_row, 1), ws.cell(total_row_new, 1))
    
    for col in range(2, max_col + 1):
        col_letter = get_column_letter(col)
        formula = f'=SUM({col_letter}{data_start}:{col_letter}{last_data_new})'
        ws.cell(total_row_new, col).value = formula
        if total_style_row:
            copy_cell_style(ws.cell(total_style_row, col), ws.cell(total_row_new, col))
    
    # Limpar linha total antiga se ficou em outro lugar
    if total_row_idx and total_row_idx != total_row_new:
        for col in range(1, max_col + 1):
            ws.cell(total_row_idx, col).value = None
    
    # ── Somatório por tipo (a partir da linha 40) ───────────────────────────
    # Linha 40 ou 2 linhas após o TOTAL, o que for maior
    tipos_start = max(40, total_row_new + 2)
    
    # Cabeçalho da seção
    header_row = tipos_start
    ws.cell(header_row, 1).value = 'REPLICAÇÕES POR TIPO'
    # Copiar estilo do cabeçalho da sheet (linha 1)
    copy_cell_style(ws.cell(1, 1), ws.cell(header_row, 1))
    # Preencher colunas do cabeçalho com os nomes dos serviços (linha 3)
    for col in range(2, max_col + 1):
        ws.cell(header_row, col).value = ws.cell(3, col).value
        copy_cell_style(ws.cell(1, col), ws.cell(header_row, col))
    
    tipo_row = header_row + 1
    tipo_first = tipo_row  # para a linha de total dos tipos
    
    # Faixa de busca para SUMIF (todos os pavimentos)
    lookup_range = f'$A${data_start}:$A${last_data_new}'
    
    for tipo_label, floor_names in tipos:
        ws.cell(tipo_row, 1).value = tipo_label
        # Copiar estilo de uma linha de dados da sheet
        copy_cell_style(ws.cell(data_start, 1), ws.cell(tipo_row, 1))
        
        for col in range(2, max_col + 1):
            col_letter = get_column_letter(col)
            data_range = f'{col_letter}${data_start}:{col_letter}${last_data_new}'
            
            if len(floor_names) == 1:
                formula = f'=SUMIF({lookup_range},"{floor_names[0]}",{data_range})'
            else:
                parts = [f'SUMIF({lookup_range},"{fn}",{data_range})' for fn in floor_names]
                formula = '=' + '+'.join(parts)
            
            ws.cell(tipo_row, col).value = formula
            copy_cell_style(ws.cell(data_start, col), ws.cell(tipo_row, col))
        
        tipo_row += 1
    
    # Linha de total dos tipos
    tipo_last = tipo_row - 1
    ws.cell(tipo_row, 1).value = 'TOTAL POR TIPO'
    if total_style_row:
        copy_cell_style(ws.cell(total_style_row, 1), ws.cell(tipo_row, 1))
    
    for col in range(2, max_col + 1):
        col_letter = get_column_letter(col)
        formula = f'=SUM({col_letter}{tipo_first}:{col_letter}{tipo_last})'
        ws.cell(tipo_row, col).value = formula
        if total_style_row:
            copy_cell_style(ws.cell(total_style_row, col), ws.cell(tipo_row, col))
    
    print(f'  OK Pavimentos reordenados: {len(ordered_floors)} | Especiais: {len(special_rows)} | Tipos: {len(tipos)}')

# ─── Processar cada sheet ───────────────────────────────────────────────────
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    
    # Só processar sheets com header "Pavimento" na célula A1
    if ws.cell(1, 1).value != 'Pavimento':
        continue
    
    print(f'\nProcessando: {sheet_name}')
    
    if sheet_name == 'SUPRAESTRUTURA':
        process_sheet(ws, CANONICAL_S_ORDERED, TIPOS_S)
    else:
        process_sheet(ws, CANONICAL_G, TIPOS_G)

wb.save(OUTPUT)
print(f'\nArquivo salvo: {OUTPUT}')

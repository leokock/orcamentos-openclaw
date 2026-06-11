# converter_sienge_visus.py
"""Converte relatórios Sienge (insumos/serviços/composições) para o formato de
importação do AltoQi Visus, e atualiza preços via tabela de fornecedores.
Uso: py -3.10 -X utf8 converter_sienge_visus.py --etapa all --dry-run
Ver SPEC: specs/SPEC-sienge-visus-villa-vauban.md
"""
import argparse, os, re, unicodedata, shutil
import csv as _csv
from dataclasses import dataclass, field
import openpyxl

# --- column maps (0-indexed) ---
INS_COD, INS_DESC, INS_UNID, INS_PRECO, INS_DATA, INS_ATIVO = 0, 1, 6, 9, 10, 11
SRV_COD, SRV_DESC, SRV_UNID, SRV_TOTAL = 0, 1, 4, 13
CMP_TIPO, CMP_COD, CMP_DESC, CMP_UNID, CMP_QTD, CMP_PU, CMP_TOTAL = 0, 1, 2, 6, 8, 10, 18
FOR_ITEM, FOR_CUSTO, FOR_UNID, FOR_FORN, FOR_CONTATO = 1, 2, 3, 4, 5
GRUPO_TIPO = {"01":"Mao_obra","02":"Material","03":"Equipamento","04":"Terceirizados","05":"Outro","06":"Material"}
TABELA_INSUMO = "Villa Vauban"
SCORE_AUTO = 0.88
DATA_DIR = r"G:\Drives compartilhados\03 CTN Projetos\2. Projetos em Andamento\RDO Empreendimentos\Villa Vauban\04. Custo\03.4 Custo - Fechamento\Importação Insumos e Composições Visus"
TEMPLATE_INSUMOS = "Thozen - Electra_INSUMOS.xlsx"
TEMPLATE_COMPOS = "Thozen - Electra_COMPOSIÇÕES.xlsx"


@dataclass
class Insumo:
    codigo: int
    descricao: str
    unidade: str
    valor: float
    tipo: str
    data_update: str = ""
    ativo: bool = True
    fonte_preco: str = ""


@dataclass
class CompItem:
    codigo_insumo: int
    descricao: str
    unidade: str
    multiplicador: float
    pu: float


@dataclass
class Composicao:
    codigo: int
    descricao: str
    unidade: str
    grupo_servico: str
    servico: str
    itens: list = field(default_factory=list)


def derive_tipo(grupo: str) -> str:
    m = re.match(r"\s*(\d{1,2})", str(grupo or ""))
    if not m:
        return "Outro"
    return GRUPO_TIPO.get(m.group(1).zfill(2), "Outro")


def parse_ativo(v) -> bool:
    if v is None:
        return True
    return str(v).strip().lower().startswith("s")


def normalize_desc(s) -> str:
    s = str(s or "")
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def _assert_header(cells, checks, ctx):
    """Valida (substring, case/accent-insensitive) que os rótulos esperados estão
    nas colunas certas do cabeçalho. Lança ValueError se algo divergir."""
    for idx, expected in checks:
        got = normalize_desc(cells[idx]) if idx < len(cells) else ""
        if normalize_desc(expected) not in got:
            raise ValueError(
                f"{ctx}: cabeçalho inesperado (col {idx} esperava '{expected}', "
                f"achou '{cells[idx] if idx < len(cells) else None}')")


def to_float(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    sv = str(v).strip()
    if re.search(r"\d,\d", sv):
        sv = sv.replace(".", "").replace(",", ".")
    if sv in ("", "-"):
        return None
    try:
        return float(sv)
    except ValueError:
        return None


from difflib import SequenceMatcher

UNIT_ALIAS = {
    "m2": "m2", "m²": "m2", "metro quadrado": "m2",
    "m3": "m3", "m³": "m3", "metro cubico": "m3", "metro cúbico": "m3",
    "m": "m", "metro": "m", "metro linear": "m",
    "un": "un", "unidade": "un", "pc": "un", "peca": "un",
    "kg": "kg", "kilograma": "kg", "quilograma": "kg",
    "t": "t", "tonelada": "t",
    "vb": "vb", "1 verba": "vb", "verba": "vb",
    "cx": "cx", "caixa": "cx", "sc": "sc", "saco": "sc",
    "gl": "gl", "galao": "gl", "lata": "lata",
}


def norm_unit(u):
    k = normalize_desc(u)
    k = re.sub(r"\s*\(.*?\)", "", k).strip()   # drop "(20Kg)" etc.
    return UNIT_ALIAS.get(k, k)


def _match_norm(s):
    s = normalize_desc(s).replace(",", ".")
    s = re.sub(r"[^a-z0-9. ]", " ", s)
    return [t for t in s.split() if t]


def best_match(item, insumos):
    """Melhor insumo por sobreposição de tokens (com número normalizado) + similaridade."""
    q = _match_norm(item)
    qset = set(q)
    if not qset:
        return None, 0.0
    qjoin = " ".join(q)
    best, best_s = None, 0.0
    for i in insumos:
        c = _match_norm(i.descricao)
        cset = set(c)
        if not cset:
            continue
        cover = len(qset & cset) / len(qset)
        seq = SequenceMatcher(None, qjoin, " ".join(c)).ratio()
        s = 0.65 * cover + 0.35 * seq
        if s > best_s:
            best, best_s = i, round(s, 3)
    return (best.codigo if best else None), best_s


_APPLY_VALUES = {"s", "sim", "y", "yes", "x", "1", "true", "aplicar"}


def _is_apply(v):
    return str(v or "").strip().lower() in _APPLY_VALUES


def parse_fornecedores(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    out = []
    for row in ws.iter_rows(min_row=6, values_only=True):
        item = row[FOR_ITEM]
        if not item or not str(item).strip():
            continue
        out.append({
            "item": str(item).replace("\n", " ").strip(),
            "custo": to_float(row[FOR_CUSTO]),
            "unidade": str(row[FOR_UNID] or "").strip(),
            "fornecedor": str(row[FOR_FORN] or "").strip(),
        })
    wb.close()
    return out


def gerar_de_para(fornecedores, insumos, out_path, prev_path=None):
    # pool de match exclui verbas lump-sum (descrições de serviço poluem o match)
    pool = [i for i in insumos if i.fonte_preco != "SIENGE-SERVICO"]
    by_code = {i.codigo: i for i in insumos}
    # preservar decisões anteriores (coluna aplicar?) por (item_fornecedor, fornecedor)
    prev = {}
    if prev_path and os.path.exists(prev_path):
        try:
            pwb = openpyxl.load_workbook(prev_path, data_only=True)
            pws = pwb.active
            for row in pws.iter_rows(min_row=2, values_only=True):
                if row and row[0]:
                    prev[(str(row[0]).strip(), str(row[9] or "").strip())] = row[10]
            pwb.close()
        except Exception:
            prev = {}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "de-para"
    headers = ["item_fornecedor", "desc_insumo_candidato", "codigo_insumo", "score",
               "preco_atual", "preco_novo", "unid_fornecedor", "unid_insumo",
               "flag_unidade_divergente", "fornecedor", "aplicar?"]
    ws.append(headers)
    for f in fornecedores:
        code, score = best_match(f["item"], pool)
        ins = by_code.get(code)
        unit_div = norm_unit(f["unidade"]) != norm_unit(ins.unidade) if ins else True
        preco_ok = f["custo"] is not None
        auto = "S" if (score >= SCORE_AUTO and not unit_div and preco_ok) else "N"
        prior = prev.get((str(f["item"]).strip(), str(f["fornecedor"]).strip()))
        aplicar = "S" if _is_apply(prior) else auto    # preserva apenas decisões de aplicar do usuário
        ws.append([f["item"], ins.descricao if ins else "", code, score,
                   ins.valor if ins else "", f["custo"] if preco_ok else "",
                   f["unidade"], ins.unidade if ins else "",
                   "SIM" if unit_div else "", f["fornecedor"], aplicar])
    wb.save(out_path)
    wb.close()


def aplicar_de_para(de_para_path, insumos):
    """Lê o de-para aprovado, aplica preço novo + fonte nos insumos com aplicar?=S.
    Muta a lista `insumos` in place; retorna nº de updates."""
    wb = openpyxl.load_workbook(de_para_path, read_only=True, data_only=True)
    ws = wb.active
    by_code = {i.codigo: i for i in insumos}
    n = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if _is_apply(row[10]):
            code = int(row[2]); novo = to_float(row[5])
            if code in by_code and novo is not None:
                by_code[code].valor = novo
                by_code[code].fonte_preco = str(row[9] or "").strip()
                n += 1
    wb.close()
    return n


def parse_insumos(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    out, cur_grupo = [], None
    for row in ws.iter_rows(min_row=1, values_only=True):
        a = row[INS_COD]
        if a == "Grupo":
            cur_grupo = row[2]
        elif a == "Código":
            _assert_header(row, [(INS_COD, "Código"), (INS_DESC, "Descrição"),
                                 (INS_UNID, "Unidade"), (INS_PRECO, "Preço")],
                           "insumos.xlsx")
            continue
        elif a in (None, "Família", "Obra", "BDI"):
            continue
        elif isinstance(a, (int, float)):
            out.append(Insumo(
                codigo=int(a),
                descricao=str(row[INS_DESC] or "").strip(),
                unidade=str(row[INS_UNID] or "").strip(),
                valor=to_float(row[INS_PRECO]) or 0.0,
                tipo=derive_tipo(cur_grupo),
                data_update=str(row[INS_DATA] or ""),
                ativo=parse_ativo(row[INS_ATIVO]),
            ))
    wb.close()
    return out


def parse_servicos(path):
    """Retorna dict {codigo:int -> {descricao, unidade, total}}."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    out = {}
    for row in ws.iter_rows(min_row=1, values_only=True):
        if row[SRV_COD] == "Código":
            _assert_header(row, [(SRV_COD, "Código"), (SRV_DESC, "Descrição"),
                                 (SRV_UNID, "Unidade")], "serviços.xlsx")
            continue
        if isinstance(row[SRV_COD], (int, float)):
            out[int(row[SRV_COD])] = {
                "descricao": str(row[SRV_DESC] or "").strip(),
                "unidade": str(row[SRV_UNID] or "").strip(),
                "total": to_float(row[SRV_TOTAL]) or 0.0,
            }
    wb.close()
    return out


def _split_code_name(v):
    """'6682 - APROVAÇÃO...' -> (6682, 'APROVAÇÃO...')"""
    s = str(v or "")
    m = re.match(r"\s*(\d+)\s*-\s*(.*)", s)
    if m:
        return int(m.group(1)), m.group(2).strip()
    return None, s.strip()


def parse_composicoes(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    comps, cur = [], None
    cur_etapa = cur_sub = ""
    for row in ws.iter_rows(min_row=8, values_only=True):
        a = row[CMP_TIPO]
        if a == "Tipo":
            _assert_header(row, [(CMP_TIPO, "Tipo"), (CMP_COD, "Código"),
                                 (CMP_DESC, "Descrição"), (CMP_QTD, "Quantidade")],
                           "composições serviços.xlsx")
            continue
        if a == "Etapa":
            _, cur_etapa = _split_code_name(row[3])
        elif a == "Subetapa":
            _, cur_sub = _split_code_name(row[3])
        elif a == "Serviço":
            code, name = _split_code_name(row[3])
            unid = next((str(v).strip() for v in row[14:] if isinstance(v, str) and v.strip()), "vb")
            cur = Composicao(codigo=code, descricao=name, unidade=unid,
                             grupo_servico=cur_etapa, servico=cur_sub, itens=[])
            comps.append(cur)
        elif a in ("MC", "MO") and cur is not None and isinstance(row[CMP_COD], (int, float)):
            cur.itens.append(CompItem(
                codigo_insumo=int(row[CMP_COD]),
                descricao=str(row[CMP_DESC] or "").strip(),
                unidade=str(row[CMP_UNID] or "").strip(),
                multiplicador=to_float(row[CMP_QTD]) or 1.0,
                pu=to_float(row[CMP_PU]) or 0.0,
            ))
    wb.close()
    return [x for x in comps if x.itens]


def parse_extras(path):
    """Lê insumos extras (manuais) de um CSV. Retorna [] se o arquivo não existe.
    Colunas: descricao, unidade, valor, tipo, fonte_preco."""
    if not path or not os.path.exists(path):
        return []
    out = []
    with open(path, encoding="utf-8-sig", newline="") as fh:
        for row in _csv.DictReader(fh):
            if row.get("descricao"):
                out.append({k: (v or "").strip() for k, v in row.items()})
    return out


def build_catalog(insumos, servicos, comps_detalhadas, extras=None):
    """Fidelidade Sienge: o preço do insumo nas composições vem do PU congelado.
    Códigos 'guarda-chuva' (1 código Sienge usado com várias descrições, ex. Aço
    em várias bitolas) são desmembrados por (código, descrição): cada descrição
    vira um insumo próprio nomeado pela descrição específica. Reproduz 100% os
    totais do Sienge. Serviços sem breakdown viram composição de 1 verba.
    Retorna (insumos_final, composicoes_final, report)."""
    from collections import Counter, defaultdict
    by_code = {i.codigo: i for i in insumos}
    insumos_out = list(insumos)
    next_code = max(by_code) + 1
    report = {"detalhadas": len(comps_detalhadas), "lump_sum": 0, "lump_sum_zero": 0,
              "price_overrides": 0, "variants": 0, "umbrella_codes": 0,
              "sanity": [], "sanity_fail": 0, "extras": 0}

    # 1) Usos por código: (ndesc, pu) -> freq; guarda desc/unid representativos
    usages = defaultdict(list)
    for cp in comps_detalhadas:
        for it in cp.itens:
            usages[it.codigo_insumo].append(
                (normalize_desc(it.descricao), round(it.pu, 4), it.descricao, it.unidade))

    # 2) Resolver código efetivo por (orig, ndesc, pu)
    eff_code = {}
    for orig, uses in usages.items():
        key_count = Counter((nd, pu) for nd, pu, od, un in uses)
        rep_desc, rep_unid = {}, {}
        for nd, pu, od, un in uses:
            rep_desc.setdefault((nd, pu), od)
            rep_unid.setdefault((nd, pu), un)
        keys = sorted(key_count, key=lambda k: (-key_count[k], -k[1]))
        ndesc_count = Counter(nd for nd, pu in keys)

        def label(nd, pu):
            d = rep_desc[(nd, pu)] or f"Item {orig}"
            return f"{d} (preço {pu})" if ndesc_count[nd] > 1 else d

        base = by_code.get(orig)
        if base is None:
            nd0, pu0 = keys[0]
            base = Insumo(codigo=orig, descricao=label(nd0, pu0), unidade=rep_unid[(nd0, pu0)] or "un",
                          valor=pu0, tipo="Outro", fonte_preco="SIENGE-COMPOSICAO")
            by_code[orig] = base
            insumos_out.append(base)

        if len(keys) == 1:
            nd, pu = keys[0]
            if abs((base.valor or 0) - pu) > 0.01:
                base.valor = pu
                base.fonte_preco = base.fonte_preco or "SIENGE-COMPOSICAO"
                report["price_overrides"] += 1
            eff_code[(orig, nd, pu)] = orig
        else:
            report["umbrella_codes"] += 1
            nd, pu = keys[0]                       # primário repurposa o código original
            base.descricao = label(nd, pu)
            base.unidade = rep_unid[(nd, pu)] or base.unidade
            base.valor = pu
            base.fonte_preco = base.fonte_preco or "SIENGE-COMPOSICAO"
            report["price_overrides"] += 1
            eff_code[(orig, nd, pu)] = orig
            for nd, pu in keys[1:]:                # demais descrições viram insumos novos
                v = Insumo(codigo=next_code, descricao=label(nd, pu),
                           unidade=rep_unid[(nd, pu)] or base.unidade, valor=pu,
                           tipo=base.tipo, fonte_preco="SIENGE-COMPOSICAO-VAR")
                by_code[next_code] = v
                insumos_out.append(v)
                eff_code[(orig, nd, pu)] = next_code
                next_code += 1
                report["variants"] += 1

    # 3) Remap das composições detalhadas + sanity check
    comps_out, detail_codes = [], set()
    for cp in comps_detalhadas:
        detail_codes.add(cp.codigo)
        cp.itens = [CompItem(codigo_insumo=eff_code[(it.codigo_insumo, normalize_desc(it.descricao), round(it.pu, 4))],
                             descricao=it.descricao, unidade=it.unidade,
                             multiplicador=it.multiplicador, pu=it.pu) for it in cp.itens]
        calc = sum(it.multiplicador * by_code[it.codigo_insumo].valor for it in cp.itens)
        sienge = servicos.get(cp.codigo, {}).get("total")
        if sienge is not None:
            diff = round(calc - sienge, 2)
            report["sanity"].append((cp.codigo, round(calc, 2), round(sienge, 2), diff))
            if abs(diff) > 0.01:
                report["sanity_fail"] += 1
        comps_out.append(cp)

    # 4) Lump-sum
    for code, s in servicos.items():
        if code in detail_codes:
            continue
        if not s["total"]:
            report["lump_sum_zero"] += 1
        verba = Insumo(codigo=next_code, descricao=s["descricao"], unidade=s["unidade"] or "vb",
                       valor=s["total"], tipo="Outro", fonte_preco="SIENGE-SERVICO")
        by_code[next_code] = verba
        insumos_out.append(verba)
        comps_out.append(Composicao(codigo=code, descricao=s["descricao"], unidade=s["unidade"] or "vb",
                                    grupo_servico="LUMP-SUM", servico=s["descricao"],
                                    itens=[CompItem(codigo_insumo=next_code, descricao=s["descricao"],
                                                    unidade=s["unidade"] or "vb", multiplicador=1.0, pu=s["total"])]))
        next_code += 1
        report["lump_sum"] += 1

    # 5) Extras manuais (persistem entre regenerações do catálogo)
    for ex in (extras or []):
        ins_ex = Insumo(codigo=next_code, descricao=ex.get("descricao", ""),
                        unidade=ex.get("unidade") or "un", valor=to_float(ex.get("valor")) or 0.0,
                        tipo=ex.get("tipo") or "Material", fonte_preco=ex.get("fonte_preco") or "EXTRA-MANUAL")
        by_code[next_code] = ins_ex
        insumos_out.append(ins_ex)
        next_code += 1
        report["extras"] += 1

    return insumos_out, comps_out, report


def gerar_relatorio(report, n_insumos, n_comps, out_path):
    diffs = [s for s in report.get("sanity", []) if abs(s[3]) > 0.01]
    lines = [
        "# Relatório de Conversão Sienge → Visus — Villa Vauban",
        "",
        f"- Insumos finais: **{n_insumos}** (base + variantes de preço + verbas lump-sum)",
        f"- Composições: **{n_comps}** ({report.get('detalhadas',0)} detalhadas + {report.get('lump_sum',0)} lump-sum)",
        f"- Preços-mestre corrigidos pelo PU congelado: **{report.get('price_overrides',0)}**",
        f"- Variantes de preço criadas: **{report.get('variants',0)}**",
        f"- Lump-sum com preço R$ 0 (revisar): **{report.get('lump_sum_zero', 0)}**",
        "",
        "## Sanity check (preço composição calc vs total Sienge)",
        f"- Composições conferidas: {len(report.get('sanity', []))}",
        f"- Com diferença > R$ 0,01: **{report.get('sanity_fail', len(diffs))}**",
        "",
    ]
    if diffs:
        lines.append("| Código | Calc | Sienge | Diff |")
        lines.append("|---|---:|---:|---:|")
        for code, calc, sienge, diff in sorted(diffs, key=lambda x: -abs(x[3]))[:30]:
            lines.append(f"| {code} | {calc} | {sienge} | {diff} |")
    else:
        lines.append("Todas as composições detalhadas reproduzem o total do Sienge (diff ≤ R$ 0,01). ✅")
    with open(out_path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))


def _clone_template(template_path, out_path, data_start=5):
    """Copia o template, apaga linhas de dados (>= data_start), retorna (wb, ws)."""
    shutil.copyfile(template_path, out_path)
    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    if ws.max_row >= data_start:
        ws.delete_rows(data_start, ws.max_row - data_start + 1)
    return wb, ws


def write_insumos(insumos, template_path, out_path):
    wb, ws = _clone_template(template_path, out_path)
    r = 5
    for i in insumos:
        ws.cell(r, 1, i.codigo)
        ws.cell(r, 2, i.descricao)
        ws.cell(r, 3, i.unidade)
        ws.cell(r, 4, i.valor)
        ws.cell(r, 5, i.fonte_preco)
        ws.cell(r, 6, i.tipo)
        ws.cell(r, 7, i.data_update)
        ws.cell(r, 8, i.ativo)
        r += 1
    wb.save(out_path)
    wb.close()


def write_composicoes(comps, insumos_by_code, template_path, out_path):
    wb, ws = _clone_template(template_path, out_path)
    r = 5
    for cp in comps:
        preco = sum(it.multiplicador * insumos_by_code[it.codigo_insumo].valor for it in cp.itens)
        for it in cp.itens:
            ins = insumos_by_code[it.codigo_insumo]
            ws.cell(r, 1, cp.codigo)
            ws.cell(r, 2, cp.descricao)
            ws.cell(r, 3, cp.unidade)
            ws.cell(r, 6, round(preco, 5))
            ws.cell(r, 7, 1)
            ws.cell(r, 8, 0)
            ws.cell(r, 9, 0)
            ws.cell(r, 10, cp.grupo_servico)
            ws.cell(r, 11, cp.servico)
            ws.cell(r, 12, it.codigo_insumo)
            ws.cell(r, 13, it.multiplicador)
            ws.cell(r, 14, TABELA_INSUMO)
            ws.cell(r, 15, "INSUMO")
            ws.cell(r, 16, ins.unidade)
            ws.cell(r, 17, ins.descricao)
            ws.cell(r, 18, ins.valor)
            r += 1
    wb.save(out_path)
    wb.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--etapa", choices=["insumos", "composicoes", "precos", "all"], default="all")
    ap.add_argument("--data-dir", default=DATA_DIR)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--aplicar-precos", action="store_true")
    args = ap.parse_args()
    d = args.data_dir
    extras = parse_extras(os.path.join(d, "insumos-extras.csv"))
    if args.etapa in ("insumos", "all"):
        ins = parse_insumos(os.path.join(d, "insumos.xlsx"))
        print(f"[INSUMOS] {len(ins)} insumos parseados")
        if not args.dry_run:
            write_insumos(ins, os.path.join(d, TEMPLATE_INSUMOS),
                          os.path.join(d, "Villa-Vauban_INSUMOS.xlsx"))
            print("  -> Villa-Vauban_INSUMOS.xlsx escrito")

    if args.etapa in ("composicoes", "all"):
        ins = parse_insumos(os.path.join(d, "insumos.xlsx"))
        srv = parse_servicos(os.path.join(d, "serviços.xlsx"))
        comps = parse_composicoes(os.path.join(d, "composições serviços.xlsx"))
        insumos2, comps2, report = build_catalog(ins, srv, comps, extras=extras)
        print(f"[COMPOSIÇÕES] {len(comps2)} composições ({report['detalhadas']} detalhadas + {report['lump_sum']} lump-sum); "
              f"{report['variants']} variantes de preço; {report['price_overrides']} preços-mestre corrigidos; "
              f"insumos finais={len(insumos2)}; sanity_fail={report['sanity_fail']}")
        if not args.dry_run:
            write_insumos(insumos2, os.path.join(d, TEMPLATE_INSUMOS), os.path.join(d, "Villa-Vauban_INSUMOS.xlsx"))
            write_composicoes(comps2, {i.codigo: i for i in insumos2},
                              os.path.join(d, TEMPLATE_COMPOS), os.path.join(d, "Villa-Vauban_COMPOSIÇÕES.xlsx"))
            print("  -> Villa-Vauban_INSUMOS.xlsx + Villa-Vauban_COMPOSIÇÕES.xlsx escritos")
            gerar_relatorio(report, n_insumos=len(insumos2), n_comps=len(comps2),
                            out_path=os.path.join(d, "relatorio-conversao.md"))
            print("  -> relatorio-conversao.md escrito")

    if args.etapa in ("precos", "all") and not args.aplicar_precos:
        ins = parse_insumos(os.path.join(d, "insumos.xlsx"))
        srv = parse_servicos(os.path.join(d, "serviços.xlsx"))
        comps = parse_composicoes(os.path.join(d, "composições serviços.xlsx"))
        insumos2, _, _ = build_catalog(ins, srv, comps, extras=extras)
        forn = parse_fornecedores(os.path.join(d, "CTN&RDO -Fornecedores-Insumos.xlsx"))
        dp = os.path.join(d, "de-para-precos-fornecedores.xlsx")
        print(f"[PREÇOS] {len(forn)} itens fornecedor; gerando de-para (catálogo completo, preservando edições)...")
        if not args.dry_run:
            gerar_de_para(forn, insumos2, dp, prev_path=dp)
            print("  -> de-para-precos-fornecedores.xlsx (revisar coluna aplicar? antes de --aplicar-precos)")

    if args.aplicar_precos:
        ins = parse_insumos(os.path.join(d, "insumos.xlsx"))
        srv = parse_servicos(os.path.join(d, "serviços.xlsx"))
        comps = parse_composicoes(os.path.join(d, "composições serviços.xlsx"))
        insumos2, comps2, _ = build_catalog(ins, srv, comps, extras=extras)   # catálogo completo p/ casar códigos
        n = aplicar_de_para(os.path.join(d, "de-para-precos-fornecedores.xlsx"), insumos2)
        print(f"[PREÇOS] {n} preços aplicados")
        if not args.dry_run:
            write_insumos(insumos2, os.path.join(d, TEMPLATE_INSUMOS), os.path.join(d, "Villa-Vauban_INSUMOS.xlsx"))
            write_composicoes(comps2, {i.codigo: i for i in insumos2},
                              os.path.join(d, TEMPLATE_COMPOS), os.path.join(d, "Villa-Vauban_COMPOSIÇÕES.xlsx"))
            print("  -> INSUMOS + COMPOSIÇÕES regravados com preços aplicados")


if __name__ == "__main__":
    main()

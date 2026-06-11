import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

# --- LEITURA ---
wb_src = openpyxl.load_workbook(
    'planejamento/OBRA - Elizabeth/06. Dimensionamento de recursos/CTN-GSL_EZB - Orçamento Executivo_R03.xlsx',
    read_only=True, data_only=True
)
ws_src = wb_src['IMPERMEABILIZAÇÃO']

data = []
for row in ws_src.iter_rows(values_only=True):
    if row[3] != 'SERVIÇO':
        continue
    multiplicador = row[0]
    pavimento = row[2]
    descricao = str(row[5]) if row[5] else ''
    unid = row[6]
    qtd = row[7] if row[7] else 0
    sistema = str(row[8]) if row[8] else 'N/A'
    ambiente = str(row[9]) if row[9] else ''
    tipo_surface = str(row[10]) if row[10] else ''

    tipo = 'OUTRO'
    if tipo_surface.strip().startswith('PISO'):
        tipo = 'PISO'
    elif tipo_surface.strip().startswith('PAREDE'):
        tipo = 'PAREDE'
    elif tipo_surface.strip().startswith('TETO'):
        tipo = 'TETO'
    elif descricao.startswith('PISO:'):
        tipo = 'PISO'
    elif descricao.startswith('PAREDE:'):
        tipo = 'PAREDE'
    elif descricao.startswith('TETO:'):
        tipo = 'TETO'

    # 'Tipo x8' e 'Penthouse x22' sao LABELS que identificam o grupo do andar.
    # Todos os andares ja estao listados individualmente no quantitativo — mult=1 sempre.
    mult = 1
    mult_label = str(multiplicador) if multiplicador else 'Unico'

    data.append({
        'pavimento': pavimento,
        'multiplicador': mult_label,
        'mult': mult,
        'descricao': descricao,
        'tipo': tipo,
        'qtd_unit': float(qtd),
        'unid': unid,
        'sistema': sistema,
        'ambiente': ambiente,
        'qtd_total': float(qtd) * mult
    })

# --- NOVO WORKBOOK ---
wb = openpyxl.Workbook()

# Cores
BLUE_DARK = PatternFill("solid", fgColor="1F3864")
BLUE_MID  = PatternFill("solid", fgColor="2F5597")
BLUE_LIGHT= PatternFill("solid", fgColor="D6E4F0")
PISO_FILL = PatternFill("solid", fgColor="E2EFDA")
PAREDE_FILL = PatternFill("solid", fgColor="FCE4D6")
TETO_FILL = PatternFill("solid", fgColor="FFF2CC")
TOTAL_FILL = PatternFill("solid", fgColor="BDD7EE")
WHITE_FONT = Font(color="FFFFFF", bold=True)
BOLD = Font(bold=True)

thin = Side(border_style="thin", color="AAAAAA")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def header_cell(ws, row, col, value, fill, font=None, align='center'):
    c = ws.cell(row=row, column=col, value=value)
    c.fill = fill
    c.font = font or Font(bold=True)
    c.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
    c.border = BORDER
    return c

def data_cell(ws, row, col, value, fill=None, bold=False, fmt=None, align='left'):
    c = ws.cell(row=row, column=col, value=value)
    if fill:
        c.fill = fill
    c.font = Font(bold=bold)
    c.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
    c.border = BORDER
    if fmt:
        c.number_format = fmt
    return c

# ===================== ABA 1: RESUMO =====================
ws1 = wb.active
ws1.title = "RESUMO"

ws1.merge_cells('A1:D1')
c = ws1.cell(1, 1, "13. IMPERMEABILIZAÇÃO E TRATAMENTOS — EZB R03")
c.fill = BLUE_DARK
c.font = Font(color="FFFFFF", bold=True, size=13)
c.alignment = Alignment(horizontal='center', vertical='center')
ws1.row_dimensions[1].height = 30

ws1.merge_cells('A2:D2')
c2 = ws1.cell(2, 1, "Extração por superfície: PISO | PAREDE | TETO")
c2.fill = BLUE_MID
c2.font = Font(color="FFFFFF", size=10)
c2.alignment = Alignment(horizontal='center')

headers = ['Superfície', 'Qtd (m²)', '% do Total']
for i, h in enumerate(headers, 1):
    header_cell(ws1, 3, i, h, BLUE_MID, WHITE_FONT, 'center')

total_piso   = sum(d['qtd_total'] for d in data if d['tipo'] == 'PISO')
total_parede = sum(d['qtd_total'] for d in data if d['tipo'] == 'PAREDE')
total_teto   = sum(d['qtd_total'] for d in data if d['tipo'] == 'TETO')
total_geral  = total_piso + total_parede + total_teto

rows_resumo = [
    ('PISO',   total_piso,   PISO_FILL),
    ('PAREDE', total_parede, PAREDE_FILL),
    ('TETO',   total_teto,   TETO_FILL),
    ('TOTAL',  total_geral,  TOTAL_FILL),
]
for i, (surf, qtd, fill) in enumerate(rows_resumo, 4):
    data_cell(ws1, i, 1, surf, fill, bold=True, align='center')
    data_cell(ws1, i, 2, round(qtd, 2), fill, fmt='#,##0.00', align='right')
    pct = qtd / total_geral if total_geral else 0
    data_cell(ws1, i, 3, pct, fill, fmt='0.0%', align='right')

ws1.column_dimensions['A'].width = 15
ws1.column_dimensions['B'].width = 16
ws1.column_dimensions['C'].width = 14

# ===================== ABA 2: POR PAVIMENTO =====================
ws2 = wb.create_sheet("POR PAVIMENTO")

ws2.merge_cells('A1:E1')
c = ws2.cell(1, 1, "IMPERMEABILIZAÇÃO — Quantitativo por Pavimento")
c.fill = BLUE_DARK
c.font = Font(color="FFFFFF", bold=True, size=12)
c.alignment = Alignment(horizontal='center', vertical='center')
ws2.row_dimensions[1].height = 25

hdrs = ['Pavimento', 'Multiplicador', 'PISO (m²)', 'PAREDE (m²)', 'TETO (m²)', 'TOTAL (m²)']
for i, h in enumerate(hdrs, 1):
    header_cell(ws2, 2, i, h, BLUE_MID, WHITE_FONT, 'center')

pav_groups = defaultdict(lambda: {'PISO': 0.0, 'PAREDE': 0.0, 'TETO': 0.0, 'mult': ''})
for d in data:
    pav = d['pavimento']
    pav_groups[pav][d['tipo']] += d['qtd_total']
    pav_groups[pav]['mult'] = d['multiplicador']

PAV_ORDER = [
    '01_1º Pav. (Térreo)', '02_Mez. 1º Pav. / G1', '03_2º Pav. / G2',
    '04_Mez. 2º Pav. / G3', '05_3º Pav. / G4', '06_Mez. 3º Pav. / G5',
    '07_4º Pav. - Lazer / G6', '08_5º Pav. - Lazer', '09_6º Pav. - Tipo Dif.',
    '10_7º Pav. - Tipo', '11_8º Pav. - Tipo', '12_9º Pav. - Tipo',
    '13_10º Pav. - Tipo', '14_11º Pav. - Tipo', '15_12º Pav. - Tipo',
    '16_13º Pav. - Tipo', '17_14º Pav. - Penthouse',
]

sorted_pavs = sorted(pav_groups.keys(), key=lambda x: str(x))

row_idx = 3
sum_piso = sum_parede = sum_teto = 0
for pav in sorted_pavs:
    g = pav_groups[pav]
    piso   = round(g['PISO'], 2)
    parede = round(g['PAREDE'], 2)
    teto   = round(g['TETO'], 2)
    total  = round(piso + parede + teto, 2)
    sum_piso += piso; sum_parede += parede; sum_teto += teto
    alt_fill = BLUE_LIGHT if row_idx % 2 == 0 else None
    data_cell(ws2, row_idx, 1, pav,   alt_fill, align='left')
    data_cell(ws2, row_idx, 2, g['mult'], alt_fill, align='center')
    data_cell(ws2, row_idx, 3, piso,   PISO_FILL,   fmt='#,##0.00', align='right')
    data_cell(ws2, row_idx, 4, parede, PAREDE_FILL, fmt='#,##0.00', align='right')
    data_cell(ws2, row_idx, 5, teto,   TETO_FILL,   fmt='#,##0.00', align='right')
    data_cell(ws2, row_idx, 6, total,  TOTAL_FILL,  fmt='#,##0.00', align='right', bold=True)
    row_idx += 1

# Linha de total
grand = round(sum_piso + sum_parede + sum_teto, 2)
data_cell(ws2, row_idx, 1, 'TOTAL GERAL', TOTAL_FILL, bold=True, align='center')
data_cell(ws2, row_idx, 2, '', TOTAL_FILL)
data_cell(ws2, row_idx, 3, round(sum_piso, 2),   TOTAL_FILL, bold=True, fmt='#,##0.00', align='right')
data_cell(ws2, row_idx, 4, round(sum_parede, 2), TOTAL_FILL, bold=True, fmt='#,##0.00', align='right')
data_cell(ws2, row_idx, 5, round(sum_teto, 2),   TOTAL_FILL, bold=True, fmt='#,##0.00', align='right')
data_cell(ws2, row_idx, 6, grand,                TOTAL_FILL, bold=True, fmt='#,##0.00', align='right')

ws2.column_dimensions['A'].width = 36
ws2.column_dimensions['B'].width = 16
ws2.column_dimensions['C'].width = 14
ws2.column_dimensions['D'].width = 14
ws2.column_dimensions['E'].width = 12
ws2.column_dimensions['F'].width = 14

# ===================== ABA 3: POR SISTEMA =====================
ws3 = wb.create_sheet("POR SISTEMA")

ws3.merge_cells('A1:F1')
c = ws3.cell(1, 1, "IMPERMEABILIZAÇÃO — Quantitativo por Sistema")
c.fill = BLUE_DARK
c.font = Font(color="FFFFFF", bold=True, size=12)
c.alignment = Alignment(horizontal='center', vertical='center')
ws3.row_dimensions[1].height = 25

hdrs3 = ['Sistema', 'PISO (m²)', 'PAREDE (m²)', 'TETO (m²)', 'TOTAL (m²)', '% Total']
for i, h in enumerate(hdrs3, 1):
    header_cell(ws3, 2, i, h, BLUE_MID, WHITE_FONT, 'center')

sis_groups = defaultdict(lambda: {'PISO': 0.0, 'PAREDE': 0.0, 'TETO': 0.0})
for d in data:
    if d['tipo'] in ['PISO', 'PAREDE', 'TETO']:
        sis_groups[d['sistema']][d['tipo']] += d['qtd_total']

row_idx = 3
for sis in sorted(sis_groups.keys()):
    g = sis_groups[sis]
    piso   = round(g['PISO'], 2)
    parede = round(g['PAREDE'], 2)
    teto   = round(g['TETO'], 2)
    total  = round(piso + parede + teto, 2)
    pct    = total / total_geral if total_geral else 0
    alt_fill = BLUE_LIGHT if row_idx % 2 == 0 else None
    data_cell(ws3, row_idx, 1, f'Sist. {sis}', alt_fill)
    data_cell(ws3, row_idx, 2, piso,   PISO_FILL,   fmt='#,##0.00', align='right')
    data_cell(ws3, row_idx, 3, parede, PAREDE_FILL, fmt='#,##0.00', align='right')
    data_cell(ws3, row_idx, 4, teto,   TETO_FILL,   fmt='#,##0.00', align='right')
    data_cell(ws3, row_idx, 5, total,  TOTAL_FILL,  fmt='#,##0.00', align='right', bold=True)
    data_cell(ws3, row_idx, 6, pct,    alt_fill,    fmt='0.0%', align='right')
    row_idx += 1

data_cell(ws3, row_idx, 1, 'TOTAL GERAL', TOTAL_FILL, bold=True)
data_cell(ws3, row_idx, 2, round(total_piso, 2),   TOTAL_FILL, bold=True, fmt='#,##0.00', align='right')
data_cell(ws3, row_idx, 3, round(total_parede, 2), TOTAL_FILL, bold=True, fmt='#,##0.00', align='right')
data_cell(ws3, row_idx, 4, round(total_teto, 2),   TOTAL_FILL, bold=True, fmt='#,##0.00', align='right')
data_cell(ws3, row_idx, 5, round(total_geral, 2),  TOTAL_FILL, bold=True, fmt='#,##0.00', align='right')
data_cell(ws3, row_idx, 6, 1.0,                    TOTAL_FILL, bold=True, fmt='0.0%', align='right')

ws3.column_dimensions['A'].width = 18
ws3.column_dimensions['B'].width = 14
ws3.column_dimensions['C'].width = 14
ws3.column_dimensions['D'].width = 12
ws3.column_dimensions['E'].width = 14
ws3.column_dimensions['F'].width = 10

# ===================== ABA 4: DETALHE COMPLETO =====================
ws4 = wb.create_sheet("DETALHE")

ws4.merge_cells('A1:G1')
c = ws4.cell(1, 1, "IMPERMEABILIZAÇÃO — Detalhe por Item")
c.fill = BLUE_DARK
c.font = Font(color="FFFFFF", bold=True, size=12)
c.alignment = Alignment(horizontal='center', vertical='center')
ws4.row_dimensions[1].height = 25

hdrs4 = ['Pavimento', 'Multiplicador', 'Superfície', 'Sistema', 'Ambiente', 'Qtd Unitária (m²)', 'Qtd Total (m²)']
for i, h in enumerate(hdrs4, 1):
    header_cell(ws4, 2, i, h, BLUE_MID, WHITE_FONT, 'center')

fill_map = {'PISO': PISO_FILL, 'PAREDE': PAREDE_FILL, 'TETO': TETO_FILL}
for i, d in enumerate(data, 3):
    fill = fill_map.get(d['tipo'])
    data_cell(ws4, i, 1, d['pavimento'], fill)
    data_cell(ws4, i, 2, d['multiplicador'], fill, align='center')
    data_cell(ws4, i, 3, d['tipo'], fill, align='center')
    data_cell(ws4, i, 4, d['sistema'], fill, align='center')
    data_cell(ws4, i, 5, d['ambiente'], fill)
    data_cell(ws4, i, 6, round(d['qtd_unit'], 2), fill, fmt='#,##0.00', align='right')
    data_cell(ws4, i, 7, round(d['qtd_total'], 2), fill, fmt='#,##0.00', align='right')

ws4.column_dimensions['A'].width = 36
ws4.column_dimensions['B'].width = 16
ws4.column_dimensions['C'].width = 10
ws4.column_dimensions['D'].width = 10
ws4.column_dimensions['E'].width = 18
ws4.column_dimensions['F'].width = 18
ws4.column_dimensions['G'].width = 16

# --- SALVAR ---
out_path = 'planejamento/OBRA - Elizabeth/06. Dimensionamento de recursos/EZB_R03_Imper_PISO_PAREDE.xlsx'
wb.save(out_path)
print(f"Arquivo salvo: {out_path}")
print(f"PISO:   {total_piso:.2f} m2")
print(f"PAREDE: {total_parede:.2f} m2")
print(f"TETO:   {total_teto:.2f} m2")
print(f"TOTAL:  {total_geral:.2f} m2")

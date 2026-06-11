#!/usr/bin/env python3
"""Aplica overrides manuais na aba Gerenciamento do paramétrico Placon V2.

Overrides pedidos (Leo 2026-04-22):
  - Mestre obras:        9.940 -> 8.000
  - Encarregado:         8.000 -> 6.000
  - Estagiário:          remover (PU=0 e tachar)
  - Téc. Segurança:      remover (PU=0 e tachar)
  - Vigilância:         15.261 -> 5.000

Modifica a coluna E (PU) das linhas correspondentes da aba 'Gerenciamento'.
Preserva todas as fórmulas de totalização que dependem dessas células.
"""
import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

sys.stdout.reconfigure(encoding="utf-8")

XLSX = Path.home() / "orcamentos-openclaw" / "base" / "pacotes" / "placon-arminio-tavares" / "parametrico-placon-arminio-tavares.xlsx"

# (row, novo_pu, nota)
OVERRIDES = [
    (9,  8000,  "Mestre obras — override Leo 22/04 (de R$ 9.940)"),
    (10, 6000,  "Encarregado — override Leo 22/04 (de R$ 8.000)"),
    (11, 0,     "Estagiário — REMOVIDO por Leo 22/04"),
    (12, 0,     "Téc. Segurança — REMOVIDO por Leo 22/04"),
    (15, 5000,  "Vigilância — override Leo 22/04 (de R$ 15.261)"),
]

STRIKE_FILL = PatternFill(start_color="FDEDEC", end_color="FDEDEC", fill_type="solid")
OVERRIDE_FILL = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")


def main():
    print(f"[info] abrindo: {XLSX.name}")
    wb = load_workbook(XLSX)
    ws = wb["Gerenciamento"]

    print("[info] aplicando overrides:")
    for row, novo_pu, nota in OVERRIDES:
        desc = ws.cell(row, 2).value
        pu_antigo = ws.cell(row, 5).value
        print(f"  r{row} | {desc:30s} | PU {pu_antigo} -> {novo_pu}  [{nota}]")

        # Col E = PU
        cell_pu = ws.cell(row, 5, novo_pu)
        cell_pu.number_format = '#,##0.00'

        if novo_pu == 0:
            # Visual: linha tachada + fundo vermelho claro
            for col in range(1, 8):
                c = ws.cell(row, col)
                c.font = Font(strike=True, size=9, color="888888", name="Arial")
                c.fill = STRIKE_FILL
        else:
            # Visual: célula PU destacada como override manual
            cell_pu.font = Font(bold=True, size=9, color="E67E22", name="Arial")
            cell_pu.fill = OVERRIDE_FILL

    # Adiciona nota na célula A2 (logo abaixo do título Gerenciamento)
    existing = ws.cell(2, 1).value or ""
    nota_geral = "OVERRIDES Leo 22/04: Mestre R$8k, Encarregado R$6k, Vigil R$5k; Estagiário e Téc.Seg. removidos"
    ws.cell(2, 1, nota_geral)
    ws.cell(2, 1).font = Font(italic=True, size=8, color="E67E22", name="Arial")

    wb.save(XLSX)
    print(f"[ok] xlsx salvo")


if __name__ == "__main__":
    main()

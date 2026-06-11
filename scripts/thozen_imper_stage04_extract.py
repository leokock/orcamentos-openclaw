#!/usr/bin/env python3
from __future__ import annotations

import re
from collections import OrderedDict, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BASE = Path("planejamento/OBRA - Thozen/06. Dimensionamento de recursos")
ORC = BASE / "CTN-TZN_ELT_EX_ORÇAMENTO -R00 .xlsx"
OUT = Path("output/TZN_Extracao_Quantidades_04_IMPERMEABILIZACAO_MODELO_IA.xlsx")


PAV_ORDER = [
    "TORRE A",
    "_RESERVATÓRIO",
    "_COBERTURA",
    "_24° PAVTO TIPO",
    "_23° PAVTO TIPO",
    "_22° PAVTO TIPO",
    "_21° PAVTO TIPO",
    "_20° PAVTO TIPO",
    "_19° PAVTO TIPO",
    "_18° PAVTO TIPO",
    "_17° PAVTO TIPO",
    "_16° PAVTO TIPO",
    "_15° PAVTO TIPO",
    "_14° PAVTO TIPO",
    "_13° PAVTO TIPO",
    "_12° PAVTO TIPO",
    "_11° PAVTO TIPO",
    "_10° PAVTO TIPO",
    "_9° PAVTO TIPO",
    "_8° PAVTO TIPO",
    "_7° PAVTO TIPO",
    "_6° PAVTO TIPO",
    "_5° PAVTO TIPO",
    "_4° PAVTO TIPO",
    "_3° PAVTO TIPO",
    "_2° PAVTO TIPO",
    "_1° PAVTO TIPO",
    "TORRE B",
    "_RESERVATÓRIO",
    "_COBERTURA",
    "_24° PAVTO TIPO",
    "_23° PAVTO TIPO",
    "_22° PAVTO TIPO",
    "_21° PAVTO TIPO",
    "_20° PAVTO TIPO",
    "_19° PAVTO TIPO",
    "_18° PAVTO TIPO",
    "_17° PAVTO TIPO",
    "_16° PAVTO TIPO",
    "_15° PAVTO TIPO",
    "_14° PAVTO TIPO",
    "_13° PAVTO TIPO",
    "_12° PAVTO TIPO",
    "_11° PAVTO TIPO",
    "_10° PAVTO TIPO",
    "_9° PAVTO TIPO",
    "_8° PAVTO TIPO",
    "_7° PAVTO TIPO",
    "_6° PAVTO TIPO",
    "_5° PAVTO TIPO",
    "_4° PAVTO TIPO",
    "_3° PAVTO TIPO",
    "_2° PAVTO TIPO",
    "_1° PAVTO TIPO",
    "EMBASAMENTO",
    "_PAVTO DIFERENCIADO / LAZER",
    "_GARAGEM G5",
    "_GARAGEM G4",
    "_GARAGEM G3",
    "_GARAGEM G2",
    "_GARAGEM G1",
    "_TERREO",
]


def clean(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\xa0", " ")).strip()


def strip_code(text: str) -> str:
    return re.sub(r"^\d+(?:\.\d+)*\.?\s*", "", clean(text)).strip()


def to_float(value):
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def norm_service(text: str) -> str:
    text = strip_code(text).upper()
    text = text.replace(":", " ")
    text = re.sub(r"[^A-Z0-9ÁÉÍÓÚÂÊÔÃÕÇ ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def service_group(text: str) -> str:
    up = norm_service(text)
    if "PAREDE" in up:
        return "Parede"
    if "PISO" in up:
        return "Piso"
    if "FORRO" in up:
        return "Forro"
    return strip_code(text)


def pav_label(raw: str) -> str:
    text = clean(raw)
    text = re.sub(r"^\d+_", "", text)
    text = re.sub(r"^(\d+)_", "", text)
    text = re.sub(r"^(\d+)°", r"\1°", text.strip())
    text = text.replace("PVTO", "PAVTO")
    text = text.replace("PAVIMENTO", "PAVTO")
    text = text.replace("PAVTO DIFERENCIADO", "PAVTO DIFERENCIADO / LAZER")
    text = text.replace("TÉRREO", "TERREO")
    text = text.replace("° PAVTO TIPO", "° PAVTO TIPO")
    if text.upper() in {"TERREO", "TÉRREO"}:
        return "_TERREO"
    if text.upper().startswith("GARAGEM"):
        return "_" + text.upper()
    if "DIFERENCIADO" in text.upper() or "LAZER" in text.upper():
        return "_PAVTO DIFERENCIADO / LAZER"
    if "COBERTURA" in text.upper():
        return "_COBERTURA"
    if "RESERV" in text.upper():
        return "_RESERVATÓRIO"
    return "_" + text.upper().lstrip()


def parse_ger_exec():
    wb = openpyxl.load_workbook(ORC, data_only=False)
    ws = wb["GER_EXEC_CTN"]
    rows = list(ws.iter_rows(values_only=True))
    records = []
    current_cc = current_etapa = current_subetapa = ""
    in_imper = False

    for row in rows[17:]:
        level = clean(row[1]).upper()
        item = clean(row[0])
        code = clean(row[2])
        desc = clean(row[3])
        unit = clean(row[4])
        qty = row[5]
        note = clean(row[9]) if len(row) > 9 else ""
        if level == "CÉLULA CONSTRUTIVA":
            current_cc = desc
            in_imper = desc.startswith("04.") or "IMPERMEABILIZA" in desc.upper()
            current_etapa = current_subetapa = ""
        elif level == "ETAPA" and in_imper:
            current_etapa = desc
            current_subetapa = ""
        elif level == "SUBETAPA" and in_imper:
            current_subetapa = desc
            if desc and clean(row[4]) and row[5] not in (None, ""):
                records.append(
                    {
                        "item": item,
                        "code": code,
                        "cc": current_cc,
                        "etapa": current_etapa,
                        "subetapa": desc,
                        "servico": desc,
                        "unit": unit,
                        "qty": qty,
                        "note": note,
                    }
                )
        elif level == "SERVIÇO" and in_imper and desc:
            records.append(
                {
                    "item": item,
                    "code": code,
                    "cc": current_cc,
                    "etapa": current_etapa,
                    "subetapa": current_subetapa,
                    "servico": desc,
                    "unit": unit,
                    "qty": qty,
                    "note": note,
                }
            )
    return records


def parse_visus():
    wb = openpyxl.load_workbook(ORC, data_only=True)
    ws = wb["VISUS - IMPER"]
    rows = list(ws.iter_rows(values_only=True))
    pav = ""
    service = ""
    buckets = defaultdict(lambda: OrderedDict())
    detailed = []

    for row in rows[11:]:
        item = clean(row[0])
        desc = clean(row[1])
        unit = clean(row[2])
        qty = to_float(row[3])
        if not item or not desc:
            continue
        depth = len([part for part in item.rstrip(".").split(".") if part])
        if depth <= 2 and not unit and qty is None:
            pav = pav_label(desc)
            continue
        if depth == 3 and not unit and qty is None:
            service = service_group(desc)
            continue
        if unit and qty is not None and pav and service:
            ambiente = strip_code(desc).upper()
            key = (service, ambiente, unit)
            if pav not in buckets[key]:
                buckets[key][pav] = 0.0
            buckets[key][pav] += qty
            detailed.append((pav, service, ambiente, unit, qty))
    return buckets, detailed


def pick_eap_record(records, grupo, ambiente):
    text = f"{grupo} {ambiente}".upper()
    candidates = []
    if "ESQUADRIA" in text or "PEITOR" in text:
        candidates = ["04.007"]
    elif "FOSSO" in text or "ELEVADOR" in text:
        candidates = ["04.002"]
    elif "VIGA BALDRAME" in text or "BALDRAME" in text:
        candidates = ["04.003"]
    elif any(k in text for k in ["PISCINA", "SPA", "LAVA PÉS", "LAVA PES", "ESPELHO"]):
        candidates = ["04.005"]
    elif "FLOREIRA" in text:
        candidates = ["04.009"]
    elif "RESERV" in text:
        candidates = ["04.010"]
    elif "TERRAÇO" in text or "TERRACO" in text or "ÁREA TÉCNICA" in text or "AREA TECNICA" in text:
        candidates = ["04.004"]
    elif any(k in text for k in ["BANHEIRO", "BOX", "COZINHA", "ÁREA DE SERVIÇO", "AREA DE SERVICO"]):
        candidates = ["04.006"]
    elif "FORRO" in text:
        candidates = ["04.006"]

    for prefix in candidates:
        for rec in records:
            if rec["etapa"].startswith(prefix):
                return rec
    return records[0] if records else {"etapa": "04. IMPERMEABILIZAÇÃO", "servico": "", "subetapa": ""}


def style(cell, fill=None, font=None, align=None):
    cell.border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    if align:
        cell.alignment = align


def build():
    ger = parse_ger_exec()
    visus, detailed = parse_visus()
    pav_rows = list(OrderedDict((pav, None) for pav, *_ in detailed).keys())

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "04 IMPERMEABILIZAÇÃO"
    resumo = wb.create_sheet("Resumo")
    fonte = wb.create_sheet("Fonte VISUS")

    fill_cc = PatternFill("solid", fgColor="2F5496")
    fill_sub = PatternFill("solid", fgColor="4472C4")
    fill_serv = PatternFill("solid", fgColor="B4C6E7")
    fill_visus = PatternFill("solid", fgColor="D9EAD3")
    fill_unit = PatternFill("solid", fgColor="C6EFCE")
    fill_total = PatternFill("solid", fgColor="FFFF00")
    fill_note = PatternFill("solid", fgColor="FCD5B4")
    fill_zebra = PatternFill("solid", fgColor="F2F2F2")
    white = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    black = Font(name="Calibri", size=10, color="000000")
    bold = Font(name="Calibri", size=10, bold=True, color="000000")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    cols = list(visus.keys())
    for r, label in enumerate(["CÉLULA", "SUBETAPA", "SERVIÇO", "VISUS", "UNIDADE"], 1):
        ws.cell(r, 1, label)
        style(ws.cell(r, 1), [fill_cc, fill_sub, fill_serv, fill_visus, fill_unit][r - 1], white if r < 3 else bold, center)

    for c, (grupo, ambiente, unit) in enumerate(cols, 2):
        eap = pick_eap_record(ger, grupo, ambiente)
        values = ["04. IMPERMEABILIZAÇÃO", eap["etapa"], eap["servico"], f"{grupo}: {ambiente}", f"({unit})"]
        fills = [fill_cc, fill_sub, fill_serv, fill_visus, fill_unit]
        fonts = [white, white, black, black, bold]
        for r, value in enumerate(values, 1):
            ws.cell(r, c, value)
            style(ws.cell(r, c), fills[r - 1], fonts[r - 1], left if r in (3, 4) else center)

    row_idx = 6
    for label in pav_rows:
        ws.cell(row_idx, 1, label)
        style(ws.cell(row_idx, 1), fill_zebra if row_idx % 2 else None, black, left)
        for c, key in enumerate(cols, 2):
            val = visus[key].get(label)
            if val:
                ws.cell(row_idx, c, val)
                ws.cell(row_idx, c).number_format = "#,##0.00"
            style(ws.cell(row_idx, c), fill_zebra if row_idx % 2 else None, black, center)
        row_idx += 1

    total_row = row_idx
    ws.cell(total_row, 1, "TOTAL")
    style(ws.cell(total_row, 1), fill_total, bold, left)
    for c in range(2, len(cols) + 2):
        letter = get_column_letter(c)
        ws.cell(total_row, c, f"=SUM({letter}6:{letter}{total_row-1})")
        ws.cell(total_row, c).number_format = "#,##0.00"
        style(ws.cell(total_row, c), fill_total, bold, center)

    note_row = total_row + 1
    ws.cell(note_row, 1, "OBSERVAÇÃO / FONTE")
    style(ws.cell(note_row, 1), fill_note, bold, left)
    for c in range(2, len(cols) + 2):
        ws.cell(note_row, c, "VISUS - IMPER")
        style(ws.cell(note_row, c), fill_note, black, left)

    ws.freeze_panes = "B6"
    ws.column_dimensions["A"].width = 26
    for c in range(2, len(cols) + 2):
        ws.column_dimensions[get_column_letter(c)].width = 18
    for r in range(1, 6):
        ws.row_dimensions[r].height = 34 if r in (2, 3, 4) else 24

    resumo_headers = ["Fonte", "Registros", "Total m²"]
    for c, h in enumerate(resumo_headers, 1):
        resumo.cell(1, c, h)
        style(resumo.cell(1, c), fill_cc, white, center)
    total_visus = sum(q for *_rest, q in detailed)
    resumo_rows = [
        ("GER_EXEC_CTN - 04. IMPERMEABILIZAÇÃO", len(ger), ""),
        ("VISUS - IMPER", len(detailed), total_visus),
        ("Colunas geradas no modelo IA", len(cols), ""),
    ]
    for r, row in enumerate(resumo_rows, 2):
        for c, value in enumerate(row, 1):
            resumo.cell(r, c, value)
            if c == 3 and isinstance(value, (int, float)):
                resumo.cell(r, c).number_format = "#,##0.00"
            style(resumo.cell(r, c), None, black, left if c == 1 else center)
    resumo.column_dimensions["A"].width = 42
    resumo.column_dimensions["B"].width = 14
    resumo.column_dimensions["C"].width = 16

    headers = ["pavimento", "grupo_visus", "ambiente", "unidade", "quantidade"]
    for c, h in enumerate(headers, 1):
        fonte.cell(1, c, h)
        style(fonte.cell(1, c), fill_cc, white, center)
    for r, row in enumerate(detailed, 2):
        for c, value in enumerate(row, 1):
            fonte.cell(r, c, value)
            if c == 5:
                fonte.cell(r, c).number_format = "#,##0.00"
            style(fonte.cell(r, c), None, black, left if c in (1, 2, 3) else center)
    for col, width in {"A": 24, "B": 18, "C": 28, "D": 12, "E": 14}.items():
        fonte.column_dimensions[col].width = width

    wb.save(OUT)
    print(OUT)
    print(f"servicos_ger_exec={len(ger)}")
    print(f"registros_visus={len(detailed)}")
    print(f"colunas_modelo_ia={len(cols)}")
    print(f"total_visus_m2={total_visus:.2f}")


if __name__ == "__main__":
    build()

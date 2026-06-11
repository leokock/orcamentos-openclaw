#!/usr/bin/env python3
r"""parse_custos.py — Fase 1 (SPEC base-custos-historica).

Extrai de UM workbook/PDF de orçamento:

  (a) PROJETO : ac / total / padrao(None aqui) + macrogrupos E disciplinas CRUS
                (mesmo shape do indices-executivo/*.json — duas visões independentes)
  (b) ITENS   : codigo / descricao / unidade / qtd / pu / total / aba + macrogrupo
  (c) INSUMOS : descricao / categoria(material|mao_obra|equipamento|outros) /
                fornecedor / consumo / pu  (composição/CPU)

Reúso (sem reescrever):
  - extract_composicoes.detect_header / classify / to_number / is_composition_aba
  - processar_executivo.safe_float / DISCIPLINE_TO_MACROGROUP / DISCIPLINE_SHEETS /
    mapear_disciplina / extrair_macrogrupos (visão "macrogrupos" crua)

Determinístico + FLAG: onde a coluna for ambígua (código vs fornecedor), aplica a
heurística do código hierárquico `01.001.001.001` e marca o resíduo em
`item["_flags"]` / `result["ambiguities"]` pra o JULGAMENTO Claude entrar na fase real.

NÃO escreve nada. Devolve dict. run_selftest.py grava o staging.

Formatos suportados:
  - Modelo Cartesian `Ger_Executivo` (Cincatarina rico / GDI resumo) — header linha 7
  - PDF AltoQi/SIENGE (TecVerde) — via pdfplumber extract_tables
  - Fallback genérico — reusa estratégias do processar_executivo (multi-aba/EAP/Sienge/ABC)

Uso CLI:
    python parse_custos.py "<arquivo.xlsx|.pdf>"
    python parse_custos.py "<arquivo>" --ac 655.19 --json
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import traceback
import unicodedata
from pathlib import Path

# --- importar os módulos reusáveis do scripts/ pai ---
_SCRIPTS = Path(__file__).resolve().parent.parent
if str(_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS))

import extract_composicoes as ec       # detect_header, classify, to_number, normalize
import processar_executivo as pe       # safe_float, DISCIPLINE_TO_MACROGROUP, etc.

from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# Regex / constantes
# ---------------------------------------------------------------------------

# Código hierárquico do modelo Cartesian/AltoQi: 01.001.001.001 (2-4 níveis)
CODIGO_HIER_RE = re.compile(r"^\d{1,3}(?:\.\d{1,3}){1,4}$")
# Código SINAPI / fonte: 87905, 7.1.2, CPU1, COMP-12 etc. (qualquer coisa alfanumérica curta)
CODIGO_FONTE_RE = re.compile(r"^[A-Z]*\d[\w\-./]{0,12}$", re.IGNORECASE)

# Níveis textuais do modelo Cartesian (coluna NIVEL)
NIVEL_FOLHA = {"servico", "serviço"}
NIVEL_AGREGADOR = {"celula construtiva", "célula construtiva", "etapa",
                   "subetapa", "unidade construtiva"}

MACRO_LABELS = sorted(set(pe.DISCIPLINE_TO_MACROGROUP.values()))


def _norm(s) -> str:
    s = str(s or "").lower().strip()
    s = unicodedata.normalize("NFKD", s)
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def _str(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s and s.lower() != "none" else None


def _is_codigo_hier(v) -> bool:
    s = _str(v)
    return bool(s) and bool(CODIGO_HIER_RE.match(s))


def _looks_codigo(v) -> bool:
    s = _str(v)
    if not s:
        return False
    return bool(CODIGO_HIER_RE.match(s) or CODIGO_FONTE_RE.match(s))


# ---------------------------------------------------------------------------
# Reconciliação de coluna (BUG de coluna no fallback — racitec/rozzo/grupo-eis)
# ---------------------------------------------------------------------------
# Os "Gerenciamento Executivo" paramétricos têm layout esparso/irregular por bloco
# (ex.: 'Descrição | Área subsolo | Altura | Volume total | Preço unitário | Valor
#  total | Observação'). O detector por keyword do processar_executivo casa o
# substring 'TOTAL' em 'Volume total' (uma coluna de QUANTIDADE) ANTES de 'Valor
# total', gravando qtd/peso na coluna total. A verdade física é: total ~= pu*qtd.
# Aqui detectamos a coluna total POR RECONCILIAÇÃO, não por nome.

def _recon_ok(pu, qtd, total, tol=0.02) -> bool:
    """pu*qtd reconcilia com total dentro de tolerância relativa."""
    if pu is None or qtd is None or total is None:
        return False
    if total == 0:
        return abs(pu * qtd) <= tol * max(1.0, abs(pu) + abs(qtd))
    return abs(pu * qtd - total) <= tol * abs(total)


# Texto que parece UNIDADE (não confundir com pu/qtd que são números)
_UNIDADE_RE = re.compile(
    r"^(m2|m3|m²|m³|m|ml|kg|t|un|und|unid|pç|pc|vb|cj|conj|gb|h|dia|mes|"
    r"l|pt|jg|par|rolo|saco|sc|brc|cento|milheiro)\.?$",
    re.IGNORECASE,
)


def _looks_unidade(v) -> bool:
    s = _str(v)
    if not s:
        return False
    return bool(_UNIDADE_RE.match(s.strip()))


# BDI típico Cartesian: 'Valor total' = pu*qtd*(1+BDI). Banda larga pra reconhecer
# que uma coluna É a coluna total (vs ser a própria qtd/pu). 1.0 (sem BDI) .. ~1.45.
_BDI_LO, _BDI_HI = 0.99, 1.50


def _resolver_triple_por_reconciliacao(num_cells, tol=0.02):
    """Dado [(col_idx, valor_float), ...] de uma linha, acha (qtd, pu, total) onde a
    coluna total é o 'Valor total' do item (= pu*qtd, possivelmente com BDI).

    Estratégia determinística (sem confiar em nome de coluna):
      - candidato a total = cada coluna, preferindo MAIOR valor (Valor total domina);
      - pra cada total, testar pares (qtd, pu) das DEMAIS colunas tais que
        total / (pu*qtd) caia na banda de BDI [0.99 .. 1.50] (i.e. total ~= pu*qtd
        ou pu*qtd + BDI). Isso evita casar a coluna qtd como total (qtd*pu >> qtd).
      - desambiguar: preferir o fit mais próximo de pu*qtd (menor BDI implícito).
    Retorna (qtd, pu, total) ou None.
    """
    if len(num_cells) < 3:
        return None
    by_total = sorted(num_cells, key=lambda c: (-abs(c[1]), -c[0]))
    melhor = None
    melhor_dist = None
    for t_idx, t_val in by_total:
        if t_val == 0:
            continue
        restantes = [(i, v) for (i, v) in num_cells if i != t_idx]
        for a in range(len(restantes)):
            for b in range(len(restantes)):
                if a == b:
                    continue
                qtd_idx, qtd_val = restantes[a]
                pu_idx, pu_val = restantes[b]
                if qtd_val == 0 or pu_val == 0:
                    continue
                base = pu_val * qtd_val
                if base == 0:
                    continue
                ratio = t_val / base
                if _BDI_LO <= ratio <= _BDI_HI:
                    # dist do fit ideal (sem BDI). qtd deve ser a coluna mais à esquerda
                    # tipicamente; usamos |ratio-1| + leve preferência a qtd<pu? não —
                    # apenas o menor BDI implícito ganha (fit exato pu*qtd vence).
                    dist = abs(ratio - 1.0)
                    if melhor_dist is None or dist < melhor_dist:
                        melhor_dist = dist
                        melhor = (qtd_val, pu_val, t_val)
                        if dist <= tol:   # fit perfeito — não precisa procurar mais
                            return melhor
    return melhor


# ---------------------------------------------------------------------------
# Mapa descrição -> macrogrupo (quando não há código; heurística determinística)
# Reusa DISCIPLINE_SHEETS->DISCIPLINE_TO_MACROGROUP do processar_executivo.
# ---------------------------------------------------------------------------

def macrogrupo_de_texto(desc: str):
    """Mapeia uma descrição/seção a um macrogrupo Cartesian por keyword.
    Retorna (macrogrupo|None, confianca 'alta'|'media'|None).
    """
    if not desc:
        return None, None
    up = str(desc).upper().strip()
    for key, disc_label in pe.DISCIPLINE_SHEETS.items():
        if key in up:
            mg = pe.DISCIPLINE_TO_MACROGROUP.get(disc_label)
            if mg:
                return mg, "media"
    return None, None


# ---------------------------------------------------------------------------
# Detector de modelo Cartesian Ger_Executivo (Cincatarina / GDI)
# ---------------------------------------------------------------------------

GER_EXEC_SHEET_NAMES = {"ger_executivo", "gerenciamento executivo", "ger executivo",
                        "gerenciamento_executivo", "orcamento executivo",
                        "orçamento executivo"}

# header keywords -> campo lógico (procurados na linha de cabeçalho ~7)
GER_HEADER_KW = {
    "nivel": ["nivel", "nível"],
    "item": ["item"],
    "descricao": ["descricao", "descrição"],
    "unidade": ["unidade", "und", "un."],
    "qtd": ["quant", "qtd"],
    "pu": ["preco un", "preço un", "preco unit", "preço unit", "custo unit", "p. unit"],
    "pu_bdi": ["preco un bdi", "preço un bdi", "preco un. bdi", "preço un. bdi"],
    "total": ["total"],
    "codigo": ["cod.", "código", "codigo", "cód"],
    "origem": ["origem"],
}


def _norm_header(s) -> str:
    """Como _norm, mas COLAPSA whitespace (inclui \n) em espaço único.
    Cabeçalhos quebram a célula em 2 linhas ('Preço un\\nBDI') — sem colapsar,
    a keyword 'preco un bdi' (com espaço) não casaria com 'preco un\\nbdi'."""
    return re.sub(r"\s+", " ", _norm(s)).strip()


def _find_ger_header(ws, max_rows=15):
    """Acha a linha de cabeçalho do Ger_Executivo e mapeia campos->col_idx (0-based).
    Retorna (header_row_1based, col_map) ou (None, None).

    Ordem dos campos em GER_HEADER_KW importa: 'pu' é testado antes de 'pu_bdi',
    mas a célula 'Preço un' (col O) e 'Preço un BDI' (col P) são colunas distintas;
    cada campo só ocupa 1 coluna, então col O -> pu e col P -> pu_bdi.
    """
    for ri, row in enumerate(ws.iter_rows(min_row=1, max_row=max_rows, values_only=True), start=1):
        if not row:
            continue
        col_map = {}
        for ci, cell in enumerate(row):
            n = _norm_header(cell)
            if not n:
                continue
            for field, kws in GER_HEADER_KW.items():
                if field in col_map:
                    continue
                # casar exato-ish: a célula deve ser curta e bater keyword
                for kw in kws:
                    if n == kw or (len(n) <= 18 and kw in n):
                        col_map[field] = ci
                        break
        # cabeçalho válido = tem NIVEL + Descricao + (Total ou PU) + Item
        if ("nivel" in col_map and "descricao" in col_map
                and ("total" in col_map or "pu" in col_map) and "item" in col_map):
            return ri, col_map
    return None, None


def _is_ger_executivo(ws) -> bool:
    name = _norm(ws.title)
    if name in GER_EXEC_SHEET_NAMES or "ger_exec" in name or "executivo" in name:
        ri, cm = _find_ger_header(ws)
        return ri is not None
    return False


def parse_ger_executivo(ws, ac=None):
    """Parse do modelo Cartesian Ger_Executivo (Cincatarina cols G-U / GDI cols G-P).

    Linha-folha = NIVEL == 'SERVIÇO'. Agregadores (CÉLULA/ETAPA/SUBETAPA) viram
    macrogrupo corrente (1º nível textual mapeável). Devolve itens + macrogrupos
    derivados dos agregadores de topo.
    """
    header_row, cm = _find_ger_header(ws)
    if header_row is None:
        return None

    itens = []
    ambiguities = []
    # macrogrupos derivados da árvore (nome cru do agregador de nível mais alto)
    macro_atual = None
    macro_totais = {}        # {nome_cru: total_acumulado}  -> visão "macrogrupos"
    disc_totais = {}         # {macrogrupo_canonico: {'total':..,'n_itens':..}} -> visão "disciplinas"

    c_nivel = cm.get("nivel")
    c_item = cm.get("item")
    c_desc = cm.get("descricao")
    c_unid = cm.get("unidade")
    c_qtd = cm.get("qtd")
    c_pu = cm.get("pu")
    c_pu_bdi = cm.get("pu_bdi")
    c_total = cm.get("total")
    c_cod = cm.get("codigo")

    # PU "oficial" do item = o que reconcilia com o Total (pu*qtd==total).
    # Cincatarina (Tipo A) tem 2 colunas: 'Preço un' (custo, pré-BDI) E 'Preço un BDI'
    # (preço vendido, == Total/qtd). GDI (Tipo E) tem só 'Preço un' (já == Total/qtd).
    # Regra: se existe coluna BDI, ela é o PU do item (preserva 'Preço un' como pu_sem_bdi).
    c_pu_oficial = c_pu_bdi if c_pu_bdi is not None else c_pu

    def cell(row, idx):
        return row[idx] if (idx is not None and idx < len(row)) else None

    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not row or all(c is None or c == "" for c in row):
            continue
        nivel = _norm(cell(row, c_nivel))
        desc = _str(cell(row, c_desc))
        if not desc:
            continue

        item_cod = _str(cell(row, c_item))     # código hierárquico 01.001...
        fonte_cod = _str(cell(row, c_cod))     # código SINAPI/fonte (col G no CIN)

        # Agregador de topo: define macrogrupo corrente
        if nivel in NIVEL_AGREGADOR:
            mg, _ = macrogrupo_de_texto(desc)
            # nível mais alto da árvore = CÉLULA CONSTRUTIVA / ETAPA — usamos a primeira
            # descrição mapeável como macrogrupo cru
            if nivel in ("celula construtiva", "célula construtiva", "etapa"):
                macro_atual = desc
                tot = pe.safe_float(cell(row, c_total))
                if tot:
                    macro_totais[desc] = macro_totais.get(desc, 0.0) + tot
            continue

        if nivel not in NIVEL_FOLHA:
            # linha sem nível reconhecível mas com descrição — só conta se tiver número
            tot_chk = pe.safe_float(cell(row, c_total))
            qtd_chk = pe.safe_float(cell(row, c_qtd))
            if not (tot_chk and qtd_chk):
                continue

        # --- linha-folha (SERVIÇO) ---
        qtd = pe.safe_float(cell(row, c_qtd))
        unid = _str(cell(row, c_unid))
        pu = pe.safe_float(cell(row, c_pu_oficial))   # PU que reconcilia com o Total
        pu_sem_bdi = pe.safe_float(cell(row, c_pu)) if c_pu_bdi is not None else None
        total = pe.safe_float(cell(row, c_total))

        # desambiguação código: item hierárquico é a chave; fonte é SINAPI
        codigo = item_cod if _is_codigo_hier(item_cod) else (item_cod or fonte_cod)
        flags = []
        if item_cod and not _is_codigo_hier(item_cod):
            flags.append("codigo_item_nao_hierarquico")
        if not codigo:
            flags.append("sem_codigo")

        mg_canon, mg_conf = macrogrupo_de_texto(macro_atual or "")
        it = {
            "codigo": codigo,
            "descricao": desc,
            "unidade": pe.normalizar_unidade(unid) if unid else None,
            "qtd": qtd,
            "pu": pu if pu else (total / qtd if (total and qtd) else None),
            "pu_sem_bdi": pu_sem_bdi,   # custo pré-BDI (só Tipo A com coluna BDI)
            "total": total,
            "aba": ws.title,
            "macrogrupo": mg_canon,
            "macrogrupo_cru": macro_atual,
            "nivel": nivel,
        }
        if flags:
            it["_flags"] = flags
            ambiguities.append({"aba": ws.title, "descricao": desc[:60], "flags": flags})
        itens.append(it)

        if mg_canon:
            d = disc_totais.setdefault(mg_canon, {"total": 0.0, "n_itens": 0, "rsm2": None})
            d["total"] += total or 0.0
            d["n_itens"] += 1

    # rsm2 por disciplina
    if ac:
        for d in disc_totais.values():
            d["rsm2"] = round(d["total"] / ac, 2) if d["total"] else None

    macrogrupos = {nome: {"valor": round(v, 2), "rsm2": round(v / ac, 2) if ac else None}
                   for nome, v in macro_totais.items()}

    return {
        "modelo": "ger_executivo",
        "aba": ws.title,
        "header_row": header_row,
        "col_map": cm,
        "itens": itens,
        "macrogrupos": macrogrupos,
        "disciplinas": disc_totais,
        "ambiguities": ambiguities,
    }


# ---------------------------------------------------------------------------
# Composições / Insumos (xlsx) — reusa extract_composicoes
# ---------------------------------------------------------------------------

def parse_composicoes_xlsx(wb) -> list[dict]:
    """Extrai insumos das abas de composição. Reusa detect_header/classify do ec.
    Detecta coluna 'fornecedor' quando presente (cotações). Marca ambíguos.
    """
    insumos = []
    for sn in wb.sheetnames:
        if not ec.is_composition_aba(sn) and "cotac" not in _norm(sn) and "cotaç" not in _norm(sn):
            continue
        ws = wb[sn]
        rows = []
        for r in ws.iter_rows(values_only=True):
            if any(c is not None and c != "" for c in r):
                rows.append(list(r))
            if len(rows) > 8000:
                break
        if not rows:
            continue
        header_idx, col_map = ec.detect_header(rows)
        if not col_map:
            continue

        # detectar coluna fornecedor (heurística por header)
        forn_idx = None
        for ci, cell in enumerate(rows[header_idx]):
            if "fornecedor" in ec.normalize(cell):
                forn_idx = ci
                break

        for row in rows[header_idx + 1:]:
            it = {}
            for field, idx in col_map.items():
                if idx < len(row):
                    v = row[idx]
                    if field in ("consumo", "pu", "total"):
                        it[field] = ec.to_number(v)
                    else:
                        it[field] = _str(v)
            desc = it.get("descricao") or ""
            if not desc or len(desc) < 3:
                continue
            has_value = any(it.get(k) not in (None, 0) for k in ("consumo", "pu", "total"))
            if not has_value:
                continue

            categoria = ec.classify(desc)
            fornecedor = _str(row[forn_idx]) if (forn_idx is not None and forn_idx < len(row)) else None

            flags = []
            # ambiguidade código x fornecedor: se 'codigo' não parece código, pode ser fornecedor
            cod = it.get("codigo")
            if cod and not _looks_codigo(cod) and not fornecedor:
                flags.append("codigo_pode_ser_fornecedor")
                fornecedor = cod  # palpite determinístico, marcado

            insumos.append({
                "descricao": desc,
                "categoria": categoria,
                "fornecedor": fornecedor,
                "unidade": it.get("unidade"),
                "consumo": it.get("consumo"),
                "pu": it.get("pu"),
                "total": it.get("total"),
                "aba": sn,
                **({"_flags": flags} if flags else {}),
            })
    return insumos


# ---------------------------------------------------------------------------
# PDF (TecVerde) — pdfplumber extract_tables
# ---------------------------------------------------------------------------

def _pt_number(s):
    """Converte número PT-BR '1.049,4800' / '155.468,28' -> float."""
    if s is None:
        return None
    if isinstance(s, (int, float)):
        return float(s)
    t = str(s).strip()
    if not t:
        return None
    t = re.sub(r"[^\d,.\-]", "", t)
    if not t:
        return None
    # PT-BR: ponto=milhar, vírgula=decimal
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


PDF_HEADER_HINTS = ["descrição", "descricao", "preço total", "preco total",
                    "preço unitário", "preco unitario", "quantidade", "código", "codigo"]


def parse_pdf_orcamento(pdf_path: Path) -> dict:
    """Extrai itens do PDF de orçamento AltoQi/SIENGE (Tipo C, TecVerde).

    Estratégia: page.extract_tables() — reagrupa colunas melhor que extract_text.
    Linhas-folha têm Un./Qtd/PU/Total; agregadoras só Total. Marca ambiguidade
    de coluna quando o header não casa direito.
    """
    import pdfplumber

    itens = []
    ambiguities = []
    texto_total = []
    total_obra = None
    n_pages = 0
    n_tables_total = 0
    pdf_macro_atual = (None, None)  # (macrogrupo_canonico, descricao_crua) corrente

    try:
        with pdfplumber.open(str(pdf_path)) as pdf:
            n_pages = len(pdf.pages)
            for pi, page in enumerate(pdf.pages):
                txt = page.extract_text() or ""
                texto_total.append(txt)
                # "Total da obra 155.468,28"
                m = re.search(r"total\s+da\s+obra[^\d]{0,10}([\d.]+,\d{2})", txt, re.IGNORECASE)
                if m:
                    total_obra = _pt_number(m.group(1))

                tables = page.extract_tables() or []
                n_tables_total += len(tables)
                for tbl in tables:
                    if not tbl or len(tbl) < 2:
                        continue
                    # achar linha-header dentro da tabela
                    hdr_i = None
                    for i, trow in enumerate(tbl[:5]):
                        joined = _norm(" ".join(str(c or "") for c in trow))
                        if sum(1 for h in PDF_HEADER_HINTS if _norm(h) in joined) >= 2:
                            hdr_i = i
                            break
                    if hdr_i is None:
                        # tabela sem header reconhecível -> ambígua, pula com flag
                        ambiguities.append({"page": pi + 1, "motivo": "tabela sem header reconhecível"})
                        continue

                    hdr = [_norm(c) for c in tbl[hdr_i]]
                    cmap = _map_pdf_cols(hdr)
                    if "descricao" not in cmap:
                        ambiguities.append({"page": pi + 1, "motivo": "sem coluna descrição", "header": hdr})
                        continue

                    for trow in tbl[hdr_i + 1:]:
                        if not trow:
                            continue
                        def g(field):
                            idx = cmap.get(field)
                            return trow[idx] if (idx is not None and idx < len(trow)) else None
                        desc = _str(g("descricao"))
                        if not desc or len(desc) < 2:
                            continue
                        codigo = _str(g("codigo"))
                        qtd = _pt_number(g("qtd"))
                        unid = _str(g("unidade"))
                        pu = _pt_number(g("pu"))
                        total = _pt_number(g("total"))

                        # linha agregadora (macrogrupo/etapa/subetapa): tem código mas
                        # sem qtd E sem pu -> é subtotal/cabeça, não item-folha.
                        cod_niveis = codigo.count(".") + 1 if codigo else 0
                        if codigo and not (qtd or pu):
                            # nível 1 (ex.: "01 SERVIÇOS TÉCNICOS") define macrogrupo corrente
                            if cod_niveis == 1:
                                mg, _ = macrogrupo_de_texto(desc)
                                pdf_macro_atual = (mg, desc)
                            continue  # agregador — não entra como item

                        mg_canon, mg_cru = pdf_macro_atual
                        it = {
                            "codigo": codigo,
                            "descricao": desc,
                            "unidade": pe.normalizar_unidade(unid) if unid else None,
                            "qtd": qtd,
                            "pu": pu,
                            "total": total,
                            "aba": f"pdf_p{pi+1}",
                            "macrogrupo": mg_canon,
                            "macrogrupo_cru": mg_cru,
                        }
                        if not codigo:
                            it["_flags"] = ["sem_codigo"]
                        itens.append(it)
    except Exception as e:
        return {
            "modelo": "pdf_altoqi",
            "itens": [],
            "macrogrupos": {},
            "disciplinas": {},
            "total_obra": None,
            "ambiguities": [{"motivo": f"pdf error: {type(e).__name__}: {e}"}],
            "errors": [traceback.format_exc()[-400:]],
        }

    # STATUS: PDF Tipo C/E que tem páginas mas NÃO produziu nenhum item-folha
    # (tabela sem header reconhecível / PDF gráfico-apresentação / escaneado) NÃO deve
    # ser descartado em silêncio — marca 'pendente-ocr' pra entrar na fila de OCR/manual
    # (caso julio-e-kalil). Só marca quando há páginas (PDF real, não vazio/corrompido).
    status = None
    if not itens and n_pages > 0:
        status = "pendente-ocr"

    out = {
        "modelo": "pdf_altoqi",
        "itens": itens,
        "macrogrupos": {},      # PDF de orçamento não traz resumo de MG do projetista
        "disciplinas": {},
        "total_obra": total_obra,
        "ambiguities": ambiguities,
        "n_pages": n_pages,
        "n_tables": n_tables_total,
    }
    if status:
        out["status"] = status
        out["status_motivo"] = (
            f"PDF com {n_pages} pág / {n_tables_total} tabela(s) mas 0 item extraível "
            f"(header não reconhecido / gráfico / escaneado) — requer OCR ou fonte itemizada"
        )
    return out


def _map_pdf_cols(hdr: list[str]) -> dict:
    cmap = {}
    for ci, h in enumerate(hdr):
        if not h:
            continue
        if "descri" in h and "descricao" not in cmap:
            cmap["descricao"] = ci
        elif ("codigo" in h or "código" in h) and "codigo" not in cmap:
            cmap["codigo"] = ci
        elif ("preco total" in h or "preço total" in h or h == "total") and "total" not in cmap:
            cmap["total"] = ci
        elif ("preco unit" in h or "preço unit" in h or "p. unit" in h or "unitario" in h) and "pu" not in cmap:
            cmap["pu"] = ci
        elif ("quant" in h or "qtd" in h) and "qtd" not in cmap:
            cmap["qtd"] = ci
        elif (h in ("un", "un.", "und", "unid", "unidade")) and "unidade" not in cmap:
            cmap["unidade"] = ci
    return cmap


def parse_pdf_composicoes(pdf_path: Path) -> list[dict]:
    """Extrai insumos do Composições.pdf (Tipo C). Cada bloco: tabela de insumos
    com Tipo(MO/MC) · Código · Descrição · Un · Qtd · PU. Reagrupa via extract_tables.
    """
    import pdfplumber
    insumos = []
    try:
        with pdfplumber.open(str(pdf_path)) as pdf:
            for pi, page in enumerate(pdf.pages):
                for tbl in (page.extract_tables() or []):
                    if not tbl or len(tbl) < 2:
                        continue
                    hdr_i = None
                    for i, trow in enumerate(tbl[:4]):
                        joined = _norm(" ".join(str(c or "") for c in trow))
                        if "descri" in joined and ("tipo" in joined or "preco" in joined or "preço" in joined):
                            hdr_i = i
                            break
                    if hdr_i is None:
                        continue
                    hdr = [_norm(c) for c in tbl[hdr_i]]
                    cmap = _map_pdf_cols(hdr)
                    tipo_idx = next((ci for ci, h in enumerate(hdr) if h == "tipo"), None)
                    if "descricao" not in cmap:
                        continue
                    for trow in tbl[hdr_i + 1:]:
                        if not trow:
                            continue
                        desc = _str(trow[cmap["descricao"]]) if cmap["descricao"] < len(trow) else None
                        if not desc or len(desc) < 3:
                            continue
                        tipo_raw = _str(trow[tipo_idx]) if (tipo_idx is not None and tipo_idx < len(trow)) else None
                        categoria = _categoria_de_tipo(tipo_raw) or ec.classify(desc)
                        pu = _pt_number(trow[cmap["pu"]]) if "pu" in cmap and cmap["pu"] < len(trow) else None
                        consumo = _pt_number(trow[cmap["qtd"]]) if "qtd" in cmap and cmap["qtd"] < len(trow) else None
                        unid = _str(trow[cmap["unidade"]]) if "unidade" in cmap and cmap["unidade"] < len(trow) else None
                        insumos.append({
                            "descricao": desc,
                            "categoria": categoria,
                            "fornecedor": None,
                            "unidade": unid,
                            "consumo": consumo,
                            "pu": pu,
                            "total": None,
                            "aba": f"pdf_comp_p{pi+1}",
                        })
    except Exception:
        pass
    return insumos


def _categoria_de_tipo(tipo_raw):
    """Mapeia o código de tipo do AltoQi (MO/MC/EQ/...) à categoria canônica."""
    if not tipo_raw:
        return None
    t = _norm(tipo_raw)
    if t in ("mo",) or "mao" in t:
        return "mao_obra"
    if t in ("eq",) or "equip" in t:
        return "equipamento"
    if t in ("mc", "ma") or "materia" in t:
        return "material"
    return None


# ---------------------------------------------------------------------------
# AC extraction helper (reusa extrair_metadados do pe)
# ---------------------------------------------------------------------------

def _extract_ac_total(wb, filename):
    """Tenta AC/total via extrair_metadados (capa) do processar_executivo."""
    try:
        meta = pe.extrair_metadados(wb, filename or "arquivo.xlsx")
    except Exception:
        meta = {}
    return meta.get("ac"), meta.get("total")


# ---------------------------------------------------------------------------
# Orquestrador
# ---------------------------------------------------------------------------

def parse_workbook(path: str | Path, ac: float | None = None,
                   slug: str | None = None) -> dict:
    """Parse de UM arquivo (xlsx/xls/xlsb/pdf). Determinístico + flags.

    Retorna:
      {
        projeto: {slug, ac, total, padrao(None), macrogrupos{}, disciplinas{}, modelo},
        itens: [ ... ],
        insumos: [ ... ],
        ambiguities: [ ... ],
        errors: [ ... ],
      }
    """
    path = Path(path)
    ext = path.suffix.lower()
    result = {
        "arquivo": str(path),
        "projeto": {
            "slug": slug,
            "ac": ac, "total": None, "padrao": None,
            "macrogrupos": {}, "disciplinas": {}, "modelo": None,
            "status": None,
        },
        "itens": [],
        "insumos": [],
        "ambiguities": [],
        "errors": [],
    }

    # ---------------- PDF (Tipo C) ----------------
    if ext == ".pdf":
        low = _norm(path.name)
        if "composi" in low:
            result["insumos"] = parse_pdf_composicoes(path)
            result["projeto"]["modelo"] = "pdf_composicoes"
            return result
        parsed = parse_pdf_orcamento(path)
        result["projeto"]["modelo"] = parsed["modelo"]
        result["projeto"]["total"] = parsed.get("total_obra")
        result["itens"] = parsed["itens"]
        result["ambiguities"] += parsed.get("ambiguities", [])
        result["errors"] += parsed.get("errors", [])
        # PDF Tipo C/E sem tabela extraível -> status pendente-ocr (não descarta)
        if parsed.get("status"):
            result["projeto"]["status"] = parsed["status"]
            result["projeto"]["status_motivo"] = parsed.get("status_motivo")
        return result

    # ---------------- xlsb/xls (sem engine aqui) ----------------
    if ext in (".xlsb", ".xls"):
        result["errors"].append(
            f"{ext} requer engine externo (pyxlsb/Excel COM) — não suportado neste parser; "
            "use o .xlsx irmão ou os PDFs (Tipo C)."
        )
        return result

    # ---------------- xlsx/xlsm ----------------
    try:
        wb = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as e:
        result["errors"].append(f"open xlsx: {type(e).__name__}: {e}")
        return result

    try:
        # AC/total da capa, se não vier por parâmetro
        ac_capa, total_capa = _extract_ac_total(wb, path.name)
        if ac is None and ac_capa:
            ac = ac_capa
            result["projeto"]["ac"] = ac
        if total_capa:
            result["projeto"]["total"] = total_capa

        # (1) tentar modelo Ger_Executivo Cartesian
        ger = None
        for sn in wb.sheetnames:
            ws = wb[sn]
            if _is_ger_executivo(ws):
                ger = parse_ger_executivo(ws, ac)
                if ger and ger["itens"]:
                    break
                ger = None

        if ger:
            result["projeto"]["modelo"] = "ger_executivo"
            result["projeto"]["macrogrupos"] = ger["macrogrupos"]
            result["projeto"]["disciplinas"] = ger["disciplinas"]
            result["itens"] = ger["itens"]
            result["ambiguities"] += ger["ambiguities"]
            # total: soma dos itens-folha se não veio da capa
            if not result["projeto"]["total"]:
                s = sum(i["total"] for i in ger["itens"] if i.get("total"))
                result["projeto"]["total"] = round(s, 2) if s else None
        else:
            # (2) fallback: estratégias do processar_executivo (multi-aba/EAP/Sienge/ABC)
            result["projeto"]["modelo"] = "fallback_processar_executivo"
            _parse_fallback(path, wb, ac, result)

        # macrogrupos crus pela visão "resumo" do pe (independente) — só se vazio
        if not result["projeto"]["macrogrupos"]:
            try:
                mg = pe.extrair_macrogrupos(wb, ac)
                if mg:
                    result["projeto"]["macrogrupos"] = mg
            except Exception as e:
                result["errors"].append(f"extrair_macrogrupos: {e}")

        # (3) insumos / composições
        try:
            result["insumos"] = parse_composicoes_xlsx(wb)
        except Exception as e:
            result["errors"].append(f"composicoes: {e}")

    finally:
        try:
            wb.close()
        except Exception:
            pass

    return result


# Header de bloco: a coluna TOTAL é a que diz 'Valor total'/'Valor' (NÃO 'Volume
# total' — é justamente o que confundia o detector antigo). A coluna QTD é
# 'Quantidade'/'Qtd'. Detectamos por nome quando o bloco tem header explícito.
def _map_block_header(row):
    """Mapeia uma linha-header de bloco -> {'qtd': idx, 'total': idx} por nome.
    'total' exige 'valor' (ou ser exatamente 'total'/'valor total'); 'volume total',
    'área total' etc. NÃO contam como coluna de valor."""
    qtd_idx = total_idx = None
    for ci, v in enumerate(row):
        n = _norm_header(v)
        if not n:
            continue
        # QTD
        if qtd_idx is None and (n == "quantidade" or n == "qtd" or n.startswith("quant")):
            qtd_idx = ci
        # TOTAL = coluna de VALOR (não volume/área). 'valor total', 'valor', 'total'
        if total_idx is None and (
            "valor total" in n or n == "valor" or n == "total"
            or (n.startswith("valor") and "total" in n)
        ):
            total_idx = ci
    # total é obrigatório; qtd pode faltar (bloco parametrico só com 'Valor total' —
    # ex.: COMUNICAÇÃO 'CFTV | R$/AC | Valor total'). Nesse caso qtd=1, pu=total.
    if total_idx is not None:
        return {"qtd": qtd_idx, "total": total_idx}
    return None


def _is_block_header_row(row):
    """True se a linha parece header de bloco (tem 'Descrição' + 'Valor total')."""
    has_desc = has_val = False
    for v in row:
        n = _norm_header(v)
        if not n:
            continue
        if "descricao" in n or n == "item":
            has_desc = True
        if "valor total" in n or n == "valor":
            has_val = True
    return has_desc and has_val


def _extrair_sheet_reconciliada(ws):
    """Extrai itens de UMA aba-disciplina, achando a coluna TOTAL certa sem confundir
    'Valor total' com 'Volume total' (o bug). Estratégia em 2 trilhos por linha:

      1. HEADER DE BLOCO (preferido): rastreia o header corrente ('Descrição ... Valor
         total'); usa col 'Quantidade' + col 'Valor total' diretamente. PU = total/qtd.
         Robusto a 'Parâmetro' ser razão de consumo (racitec) — só importa qtd e total.
      2. RECONCILIAÇÃO (sem header): acha (qtd, pu, total) onde pu*qtd ~= total (±BDI).

    Robusto a layout esparso/irregular por bloco (Gerenciamento Executivo paramétrico:
    racitec/rozzo/grupo-eis). Linhas-header, sub-totais e linhas sem valor são puladas.
    """
    itens = []
    block_cols = None    # {'qtd','total'} do header de bloco corrente
    for row in ws.iter_rows(min_row=1, max_row=4000, values_only=True):
        if not row:
            continue

        # atualizar header de bloco corrente
        if _is_block_header_row(row):
            block_cols = _map_block_header(row)
            continue

        # descrição = 1ª célula de texto não-numérica (col A normalmente)
        desc = None
        desc_idx = None
        for ci, v in enumerate(row):
            s = _str(v)
            if s and pe.safe_float(v) is None and len(s) >= 3 and not _looks_unidade(v):
                desc = s
                desc_idx = ci
                break
        if not desc:
            continue
        up = desc.upper()
        if any(k in up for k in ("DESCRIÇÃO", "DESCRICAO", "TOTAL GERAL", "SUBTOTAL",
                                 "VALOR TOTAL", "SOMA", "ARQUIVO BASE")):
            continue
        if up.strip() in ("TOTAL", "TOTAIS", "DATA", "OBRA", "CLIENTE", "DOCUMENTO"):
            continue

        # unidade = 1ª célula que parece unidade
        unidade = None
        for v in row:
            if _looks_unidade(v):
                unidade = _str(v)
                break

        qtd = total = None
        pu_lido = None

        # --- trilho 1: header de bloco explícito ---
        if block_cols:
            qi, ti = block_cols["qtd"], block_cols["total"]
            tv = pe.safe_float(row[ti]) if (ti is not None and ti < len(row)) else None
            qv = pe.safe_float(row[qi]) if (qi is not None and qi < len(row)) else None
            if tv and tv > 0:
                if qv:
                    qtd, total = qv, tv
                else:
                    # bloco parametrico sem coluna Quantidade: 1 unidade, pu=total
                    qtd, total = 1.0, tv
                    if unidade is None:
                        unidade = "vb"

        # --- trilho 2: reconciliação pura (sem header confiável) ---
        if total is None:
            num_cells = []
            for ci, v in enumerate(row):
                if ci == desc_idx or _looks_unidade(v):
                    continue
                fv = pe.safe_float(v)
                if fv is not None:
                    num_cells.append((ci, fv))
            triple = _resolver_triple_por_reconciliacao(num_cells)
            if triple:
                qtd, pu_lido, total = triple

        if not (qtd and total) or total <= 0:
            continue

        # PU oficial = total/qtd (preço efetivo COM BDI; reconcilia 1:1 — mesma
        # convenção do extrair_itens_disciplina original). pu_lido (pré-BDI ou razão
        # de consumo) só é guardado como pu_sem_bdi quando é um preço plausível.
        pu = total / qtd
        pu_sem_bdi = None
        if pu_lido and 0.5 <= (pu / pu_lido) <= 1.6 and abs(pu_lido - pu) > 0.01 * max(1.0, pu):
            pu_sem_bdi = pu_lido

        it = {
            "descricao": desc,
            "unidade": pe.normalizar_unidade(unidade) if unidade else None,
            "qtd": qtd,
            "pu": pu,
            "total": total,
            "aba": ws.title,
        }
        if pu_sem_bdi:
            it["pu_sem_bdi"] = pu_sem_bdi
        itens.append(it)
    return itens


def _reconcile_frac(itens):
    """Fração [0..1] de soma(pu*qtd) que bate com soma(total) — métrica de saúde.
    Retorna (frac, n_com_triple). frac=1.0 quando vazio (nada a reconciliar)."""
    td = tc = 0.0
    n = 0
    for it in itens:
        pu, qtd, total = it.get("pu"), it.get("qtd"), it.get("total")
        if pu is None or qtd is None or total is None:
            continue
        n += 1
        td += total
        tc += pu * qtd
    if td == 0:
        return (1.0, n)
    frac = tc / td
    # distância simétrica de 1.0 (penaliza tanto inflar quanto encolher)
    return (1.0 - abs(frac - 1.0), n)


def _desc_set(itens):
    """Conjunto de descrições normalizadas (primeiras palavras) de um item-set —
    usado pra casar a aba reconciliada certa com o item-set original quando o label
    é sintético ('Insumos ABC') e NÃO mapeia pra um nome de aba real."""
    out = set()
    for it in itens:
        d = _norm(it.get("descricao"))
        if d:
            out.add(d[:40])
    return out


def _achar_aba_fonte_por_reconciliacao(wb, itens_orig, frac_orig, tol_recon=0.995):
    """Quando o label da disciplina é SINTÉTICO (ex.: 'Insumos ABC' do Strategy-4 do
    processar_executivo) e NÃO há nome de aba real pra re-extrair, varre TODAS as abas
    do workbook por reconciliação e devolve a melhor candidata.

    Critério (geral, sem hardcode de nome de aba):
      1. extrai a aba via _extrair_sheet_reconciliada (acha pu/qtd/total que fecham);
      2. exige reconciliar BEM (frac >= tol_recon) — senão a aba não é orçamento real;
      3. entre as que reconciliam, escolhe a que MELHOR reproduz o item-set original:
         maior sobreposição de descrições (Jaccard) e, em empate, contagem mais próxima.
    Isso garante que recuperamos exatamente a MESMA aba-fonte que o fallback leu
    (mesma realidade física do orçamento), só que com a coluna total correta.

    Retorna (sheet_name, novos_itens, frac_novo, score) ou None.
    """
    alvo = _desc_set(itens_orig)
    n_orig = len(itens_orig)
    melhor = None
    for sn in wb.sheetnames:
        try:
            novos = _extrair_sheet_reconciliada(wb[sn])
        except Exception:
            continue
        if not novos:
            continue
        frac_novo, n_rec = _reconcile_frac(novos)
        if frac_novo < tol_recon or n_rec == 0:
            continue
        # sobreposição de descrições com o item-set original
        ds = _desc_set(novos)
        inter = len(alvo & ds)
        union = len(alvo | ds) or 1
        jacc = inter / union
        # proximidade de contagem (1.0 quando idêntico)
        prox = 1.0 - (abs(len(novos) - n_orig) / max(len(novos), n_orig, 1))
        score = jacc * 0.75 + prox * 0.25
        cand = (sn, novos, frac_novo, score, jacc, len(novos))
        if melhor is None or score > melhor[3]:
            melhor = cand
    # só aceita se reproduz minimamente o original (evita pegar uma aba aleatória que
    # reconcilia mas é outra coisa); 0.30 de Jaccard OU contagem idêntica = mesma fonte
    if melhor and (melhor[4] >= 0.30 or melhor[5] == n_orig):
        return melhor[0], melhor[1], melhor[2], melhor[3]
    return None


def _reparar_itens_por_reconciliacao(path, raw_disciplinas, result):
    """Corrige o BUG de coluna: re-extrai por reconciliação as abas-disciplina cujo
    item-set do processar_executivo NÃO reconcilia (pu*qtd vs total).

    NÃO toca abas que já reconciliam (essege/holze/soles ficam idênticos). Re-abre o
    workbook só pra ler as abas problemáticas. Devolve a lista final de itens.
    """
    from openpyxl import load_workbook

    # itens originais por aba (do processar_executivo), no shape parse_custos
    orig_por_aba = {}
    for disc_label, d in raw_disciplinas.items():
        mg = pe.DISCIPLINE_TO_MACROGROUP.get(disc_label)
        bucket = orig_por_aba.setdefault(disc_label, [])
        for it in d.get("itens", []):
            bucket.append({
                "codigo": it.get("codigo"),
                "descricao": it.get("descricao"),
                "unidade": it.get("unidade"),
                "qtd": it.get("quantidade"),
                "pu": it.get("pu"),
                "total": it.get("total"),
                "aba": disc_label,
                "macrogrupo": mg,
                "macrogrupo_cru": disc_label,
            })

    # quais disciplinas NÃO reconciliam? (bug de coluna OU pu pré-BDI). Re-extrair por
    # reconciliação corrige ambos: acha a coluna 'Valor total' certa e deriva pu=total/qtd.
    # Limiar quase-perfeito (0.999) pega tanto o bug de coluna grave quanto o gap de BDI
    # (~5%); o guard adiante só aceita o reparo se reconciliar MELHOR (nunca regride).
    abas_quebradas = {}
    for disc_label, itens in orig_por_aba.items():
        frac, n = _reconcile_frac(itens)
        if n >= 1 and frac < 0.999:
            abas_quebradas[disc_label] = (frac, n)

    if not abas_quebradas:
        # nada a reparar — devolve os itens originais inalterados
        return [it for bucket in orig_por_aba.values() for it in bucket]

    # mapear label-canônico -> nome real da aba no workbook
    try:
        wb = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as e:
        result["errors"].append(f"reparo: open xlsx: {type(e).__name__}: {e}")
        return [it for bucket in orig_por_aba.values() for it in bucket]

    label_to_sheet = {}
    for sn in wb.sheetnames:
        lbl = pe.mapear_disciplina(sn)
        if lbl and lbl not in label_to_sheet:
            label_to_sheet[lbl] = sn

    finais = []
    reparos = []
    nao_reparadas = []   # disciplinas quebradas que NENHUMA aba reconciliou -> pendente-revisao
    for disc_label, itens in orig_por_aba.items():
        if disc_label not in abas_quebradas:
            finais.extend(itens)
            continue
        frac_orig, _ = abas_quebradas[disc_label]

        # (a) aba real pelo nome canônico da disciplina (essege/holze/soles etc.)
        sheet_name = label_to_sheet.get(disc_label)
        novos = None
        if sheet_name:
            try:
                novos = _extrair_sheet_reconciliada(wb[sheet_name])
            except Exception as e:
                result["errors"].append(f"reparo {disc_label}: {type(e).__name__}: {e}")
                novos = None

        frac_novo, n_novo = _reconcile_frac(novos) if novos else (0.0, 0)

        # (b) label SINTÉTICO ('Insumos ABC' do Strategy-4) OU a aba canônica não
        #     reconciliou: varre TODAS as abas por reconciliação e casa a fonte certa
        #     pela sobreposição de descrições com o item-set original. É isso que
        #     recupera all-empreiteira-* (Tipo B), onde a coluna %Part. foi lida como
        #     total e o label não bate com nome de aba nenhum.
        if not (novos and frac_novo > frac_orig):
            scan = _achar_aba_fonte_por_reconciliacao(wb, itens, frac_orig)
            if scan:
                sheet_name, novos_scan, frac_scan, _score = scan
                if frac_scan > frac_orig:
                    novos = novos_scan
                    frac_novo, n_novo = frac_scan, len(novos_scan)

        # aceita o reparo só se reconcilia melhor que o original (nunca regride)
        if novos and frac_novo > frac_orig:
            mg = pe.DISCIPLINE_TO_MACROGROUP.get(disc_label)
            for it in novos:
                it["macrogrupo"] = mg
                it["macrogrupo_cru"] = disc_label
                it["_flags"] = (it.get("_flags") or []) + ["coluna_total_reconciliada"]
            finais.extend(novos)
            reparos.append({"disciplina": disc_label, "aba": sheet_name,
                            "frac_antes": round(frac_orig, 3),
                            "frac_depois": round(frac_novo, 3),
                            "n_antes": len(itens), "n_depois": n_novo})
        else:
            # NÃO chuta: mantém o original mas registra que ficou quebrado. O gate
            # marca pendente-revisao (vide _parse_fallback) — não carrega lixo.
            finais.extend(itens)
            nao_reparadas.append({"disciplina": disc_label,
                                  "frac": round(frac_orig, 3), "n": len(itens)})
    wb.close()

    if reparos:
        result.setdefault("repairs", []).extend(reparos)
        for rp in reparos:
            result["ambiguities"].append({
                "tipo": "coluna_total_reconciliada", **rp,
            })
    if nao_reparadas:
        result.setdefault("_reparo_pendente", []).extend(nao_reparadas)
    return finais


def _parse_fallback(path, wb, ac, result):
    """Usa o processar_executivo.processar_executivo (estratégia em cascata) e
    converte a saída pro shape (itens + disciplinas + macrogrupos)."""
    try:
        meta, indices_data, raw_data = pe.processar_executivo(str(path),
                                                              slug=result["projeto"]["slug"])
    except Exception as e:
        result["errors"].append(f"fallback processar_executivo: {type(e).__name__}: {e}")
        return

    if indices_data.get("ac") and not result["projeto"]["ac"]:
        result["projeto"]["ac"] = indices_data["ac"]
    if indices_data.get("total"):
        result["projeto"]["total"] = indices_data["total"]
    result["projeto"]["macrogrupos"] = indices_data.get("macrogrupos", {}) or {}
    # disciplinas no shape do indices-executivo (label canônico -> total/n_itens)
    result["projeto"]["disciplinas"] = indices_data.get("disciplinas", {}) or {}

    # itens detalhados vêm do raw_data.disciplinas[label].itens — com REPARO de coluna
    # por reconciliação (corrige o bug 'Volume total'/'%Part.' lido como total em
    # racitec/rozzo/grupo-eis e all-empreiteira-*; abas que já reconciliam passam intactas).
    raw_disc = raw_data.get("disciplinas", {})
    itens_finais = _reparar_itens_por_reconciliacao(path, raw_disc, result)
    result["itens"].extend(itens_finais)

    # STATUS pendente-revisao (geral, não-hardcode): se o item-set final NÃO reconcilia
    # em massa (>=95% dos itens com triple pu/qtd/total e a soma fecha), o parser NÃO
    # deve deixar o gate carregar lixo. Acontece quando o fallback misparse a coluna e
    # NENHUMA aba do workbook reconciliou no reparo. Marca o motivo e deixa pra revisão.
    frac_final, n_final = _reconcile_frac(itens_finais)
    pct_final = round(100 * frac_final, 2) if n_final else None
    if result.get("_reparo_pendente") and n_final >= 1 and frac_final < 0.95:
        pend = result.pop("_reparo_pendente")
        discs = ", ".join(f"{p['disciplina']}(recon~{int(p['frac']*100)}%,n={p['n']})"
                          for p in pend)
        result["projeto"]["status"] = "pendente-revisao"
        result["projeto"]["status_motivo"] = (
            f"coluna total não reconciliou no reparo (final {pct_final}% em {n_final} itens; "
            f"disciplinas: {discs}) — nenhuma aba do workbook fechou pu*qtd~=total >=95%; "
            f"requer mapeamento manual de coluna"
        )
    else:
        # reparo resolveu (ou nunca quebrou) — limpa o marcador interno
        result.pop("_reparo_pendente", None)


def main():
    ap = argparse.ArgumentParser(description="Parse de workbook/PDF de custo (Fase 1)")
    ap.add_argument("arquivo")
    ap.add_argument("--ac", type=float, default=None)
    ap.add_argument("--slug", default=None)
    ap.add_argument("--json", action="store_true")
    args = ap.parse_args()

    r = parse_workbook(args.arquivo, ac=args.ac, slug=args.slug)

    if args.json:
        print(json.dumps(r, ensure_ascii=False, indent=2))
        return

    p = r["projeto"]
    print(f"Arquivo: {Path(r['arquivo']).name}")
    print(f"Modelo: {p['modelo']}  | AC: {p['ac']}  | Total: {p['total']}")
    print(f"Macrogrupos (cru): {len(p['macrogrupos'])}  | Disciplinas: {len(p['disciplinas'])}")
    print(f"Itens: {len(r['itens'])}  | Insumos: {len(r['insumos'])}")
    print(f"Ambiguidades: {len(r['ambiguities'])}  | Erros: {len(r['errors'])}")
    if r["errors"]:
        for e in r["errors"][:5]:
            print(f"  ERRO: {e}")
    print("\nAmostra de itens:")
    for it in r["itens"][:5]:
        print(f"  [{it.get('codigo')}] {str(it.get('descricao'))[:50]:<50} "
              f"{it.get('unidade')} q={it.get('qtd')} pu={it.get('pu')} "
              f"tot={it.get('total')} mg={it.get('macrogrupo')}")
    if r["insumos"]:
        print("\nAmostra de insumos:")
        for ix in r["insumos"][:5]:
            print(f"  [{ix['categoria']:<10}] {str(ix['descricao'])[:45]:<45} "
                  f"forn={ix.get('fornecedor')} cons={ix.get('consumo')} pu={ix.get('pu')}")


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
r"""_verify_pus.py — PROVA do auto-teste Fase 2 (NÃO faz parte do pipeline).

Para cada âncora: re-abre o workbook/PDF de origem de forma INDEPENDENTE do parser,
localiza 5 itens gravados no staging pelo CÓDIGO e confere o PU gravado vs a célula
REAL (sheet!cell=valor). Imprime resultado por item.

Uso: PYTHONIOENCODING=utf-8 python _verify_pus.py
"""
from __future__ import annotations
import json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

STAGING = Path(r"c:\Users\leona\openclaw\_local\base-custos-historica\staging")

# Picks reais (do manifest) — fonte de verdade pra reabrir
SOURCES = {
    "cincatarina": {
        "kind": "ger_xlsx",
        "wb": r"G:\Drives compartilhados\03 CTN Projetos\2. Projetos em Andamento\Cincatarina\04. Custo\3.6 Custo - Orçamento Executivo\CTN-MON_CIN - Orçamento Executivo_R00.xlsx",
        "sheet": "Ger_Executivo",
        # colunas REAIS conferidas manualmente no dump (1-based): G=cod-fonte, I=NIVEL,
        # J=Item(codigo hier), L=Descricao, M=Unid, N=Quant, O=Preco un, P=Preco un BDI, Q=Total
        "col_item": "J", "col_nivel": "I", "col_pu_oficial": "P", "col_pu_sem_bdi": "O",
        "col_qtd": "N", "col_total": "Q", "header_row": 7,
    },
    "gdi-santa-monica": {
        "kind": "ger_xlsx",
        "wb": r"G:\Drives compartilhados\03 CTN Projetos\2. Projetos em Andamento\GDI Empreendimentos\Santa Mônica\04. Custo\03.4 Custo - Fechamento\CTN_GDI-THP - Orçamento_R00.xls.xlsx",
        "sheet": "Ger_Executivo",
        # I=Item(codigo), H=NIVEL, J=Descricao, K=Unid, L=Quant, M=Preco un, N=Total
        "col_item": "I", "col_nivel": "H", "col_pu_oficial": "M", "col_pu_sem_bdi": None,
        "col_qtd": "L", "col_total": "N", "header_row": 7,
    },
    "tecverde-casa-c4e": {
        "kind": "pdf",
        "pdf": r"G:\Drives compartilhados\03 CTN Projetos\2. Projetos em Andamento\TecVerde\04. Custo\03.0 Custo - Orçamento Preliminar\R01\Relatorios\CTN-TCV-C4E - Orçamento-R01.pdf",
    },
}

COL = {}  # cache letter->idx0


def col_idx0(letter: str) -> int:
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(letter) - 1


def approx(a, b, tol=0.01):
    if a is None or b is None:
        return False
    return abs(float(a) - float(b)) <= tol * max(1.0, abs(float(b)))


def verify_ger(slug, cfg, n=5):
    """Confere PU de n itens xlsx Ger_Executivo lendo a célula REAL pelo codigo."""
    recs = [json.loads(l) for l in (STAGING / slug / "itens.jsonl").open(encoding="utf-8")]
    # escolher 5 espalhados que tenham pu+codigo hierárquico
    cands = [r for r in recs if r.get("pu") and r.get("codigo") and r["codigo"].count(".") == 3]
    picks = [cands[i] for i in (0, len(cands)//4, len(cands)//2, 3*len(cands)//4, len(cands)-1)][:n]

    wb = load_workbook(cfg["wb"], read_only=True, data_only=True)
    ws = wb[cfg["sheet"]]
    ci_item = col_idx0(cfg["col_item"])
    ci_pu = col_idx0(cfg["col_pu_oficial"])
    ci_total = col_idx0(cfg["col_total"])
    ci_qtd = col_idx0(cfg["col_qtd"])

    # indexar TODAS as linhas SERVIÇO por codigo -> (excel_row, valores reais)
    real = {}
    for ri, row in enumerate(ws.iter_rows(min_row=cfg["header_row"]+1, values_only=True), start=cfg["header_row"]+1):
        item = row[ci_item] if ci_item < len(row) else None
        if item is None:
            continue
        s = str(item).strip()
        if s.count(".") == 3:  # codigo de servico-folha
            real[s] = {
                "row": ri,
                "pu": row[ci_pu] if ci_pu < len(row) else None,
                "total": row[ci_total] if ci_total < len(row) else None,
                "qtd": row[ci_qtd] if ci_qtd < len(row) else None,
            }
    wb.close()

    results = []
    for r in picks:
        cod = r["codigo"]
        rr = real.get(cod)
        if not rr:
            results.append({"codigo": cod, "status": "NAO_ENCONTRADO_NA_CELULA"})
            continue
        cell_ref = f"{cfg['sheet']}!{cfg['col_pu_oficial']}{rr['row']}"
        match = approx(r["pu"], rr["pu"])
        results.append({
            "codigo": cod,
            "descricao": str(r.get("descricao"))[:40],
            "celula": cell_ref,
            "valor_celula": rr["pu"],
            "valor_staging": r["pu"],
            "match": match,
            # checagem extra de reconciliação na própria célula real
            "recon_real": approx((rr["pu"] or 0)*(rr["qtd"] or 0), rr["total"]),
        })
    return results


def verify_pdf(slug, cfg, n=5):
    """Confere PU de n itens do PDF re-extraindo a tabela do PDF de forma independente."""
    import pdfplumber, re
    recs = [json.loads(l) for l in (STAGING / slug / "itens.jsonl").open(encoding="utf-8")]
    cands = [r for r in recs if r.get("pu") and r.get("descricao")]
    picks = [cands[i] for i in (0, len(cands)//4, len(cands)//2, 3*len(cands)//4, len(cands)-1)][:n]

    def ptnum(s):
        if s is None:
            return None
        t = re.sub(r"[^\d,.\-]", "", str(s))
        if not t:
            return None
        if "," in t:
            t = t.replace(".", "").replace(",", ".")
        try:
            return float(t)
        except ValueError:
            return None

    # Fonte da verdade = LINHA DE TEXTO do PDF (extract_text), independente do parser
    # (que usa extract_tables). Cada item-folha AltoQi é uma linha:
    #   "18.000.003.001 Ligações Definitivas vb 1,0000 2.500,0000 2.500,00"
    def nrm(s):
        return re.sub(r"\s+", " ", str(s or "").strip().lower())

    # NB: a numeração AltoQi REINICIA por Unidade Construtiva — código não é único no
    # PDF (01.001.001.001 aparece em 'Levantamento Topografico' p1 E 'Serviço de Bota
    # fora' p2). Por isso a chave de unicidade do SPEC é (sha1,sheet,row), não código.
    # Aqui guardamos TODAS as linhas e desambiguamos por (código + descrição).
    all_lines = []  # (page, codigo, raw_line)
    cod_re = re.compile(r"^(\d{1,3}(?:\.\d{1,3}){3})\s+(.*)$")
    with pdfplumber.open(cfg["pdf"]) as pdf:
        for pi, page in enumerate(pdf.pages, start=1):
            for line in (page.extract_text() or "").splitlines():
                m = cod_re.match(line.strip())
                if m:
                    all_lines.append((pi, m.group(1), line.strip()))

    def find_line(cod, desc):
        dn = nrm(desc)[:30]
        # 1) código + descrição batem
        for pi, c, raw in all_lines:
            if c == cod and dn and dn in nrm(raw):
                return pi, raw
        # 2) só descrição
        for pi, c, raw in all_lines:
            if dn and dn in nrm(raw):
                return pi, raw
        # 3) só código (1ª ocorrência)
        for pi, c, raw in all_lines:
            if c == cod:
                return pi, raw
        return None

    results = []
    for r in picks:
        cod = r.get("codigo")
        entry = find_line(cod, r.get("descricao"))
        if not entry:
            results.append({"descricao": str(r["descricao"])[:40], "status": "NAO_ENCONTRADO_NO_PDF"})
            continue
        page, raw = entry
        # números na linha, em ordem: ... qtd pu total (3 últimos)
        nums = re.findall(r"\d[\d.]*,\d+", raw)
        vals = [ptnum(x) for x in nums]
        pu_na_linha = any(approx(r["pu"], x, tol=0.01) for x in vals if x is not None)
        tot_na_linha = r.get("total") is None or any(approx(r["total"], x, tol=0.01) for x in vals if x is not None)
        recon = (approx((r.get("pu") or 0)*(r.get("qtd") or 0), r.get("total"))
                 if r.get("pu") and r.get("qtd") and r.get("total") else None)
        match = pu_na_linha and tot_na_linha
        results.append({
            "descricao": str(r["descricao"])[:40],
            "celula": f"PDF p{page} linha '{raw[:70]}'",
            "valores_celula_row": vals,
            "valor_staging": r["pu"],
            "match": match,
            "recon_real": recon,
        })
    return results


def main():
    print("=" * 70)
    print("PROVA — 5 PUs por âncora: staging vs célula REAL (re-leitura independente)")
    print("=" * 70)
    all_ok = True
    for slug, cfg in SOURCES.items():
        print(f"\n### {slug} ({cfg['kind']})")
        if cfg["kind"] == "ger_xlsx":
            res = verify_ger(slug, cfg)
        else:
            res = verify_pdf(slug, cfg)
        for r in res:
            ok = r.get("match")
            all_ok = all_ok and bool(ok)
            mark = "OK " if ok else ("?? " if ok is None else "XX ")
            if cfg["kind"] == "ger_xlsx":
                print(f"  {mark}[{r.get('codigo')}] {r.get('descricao','')}")
                print(f"      {r.get('celula')} = {r.get('valor_celula')}  | staging.pu = {r.get('valor_staging')}"
                      f"  | recon_celula_real={r.get('recon_real')}")
            else:
                print(f"  {mark}{r.get('descricao','')}")
                print(f"      {r.get('celula')} nums={r.get('valores_celula_row')}  | staging.pu = {r.get('valor_staging')}")
            if r.get("status"):
                print(f"      STATUS: {r['status']}")
    print(f"\n=== TODAS AS PROVAS PASSARAM: {all_ok} ===")


if __name__ == "__main__":
    main()

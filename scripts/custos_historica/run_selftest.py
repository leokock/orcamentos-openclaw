#!/usr/bin/env python3
r"""run_selftest.py — Fase 1/2 (SPEC base-custos-historica, portão de auto-teste).

Dado o caminho de UMA pasta de custo, roda discovery + parse + data e grava
em staging local. ZERO escrita em Supabase/Drive — só staging.

Saídas em:
  c:\Users\leona\openclaw\_local\base-custos-historica\staging\<slug>\
    - projeto.json    : metadados do projeto (ac/total/padrao + macrogrupos + disciplinas crus)
    - itens.jsonl     : 1 item por linha (codigo/desc/un/qtd/pu/total/aba/macrogrupo/data)
    - insumos.jsonl   : 1 insumo por linha (desc/categoria/fornecedor/consumo/pu)
    - manifest.json    : discovery + folder_tipo + completude + data_base + contagens

folder_tipo (heurística do mapa de extração):
  A = rico (xlsx com Ger_Executivo + composições)   -> Cincatarina
  E = resumo (1 xlsx, modelo Ger_Executivo enxuto)  -> GDI/Santa Mônica
  C = PDF (export AltoQi/SIENGE)                     -> TecVerde
  B/D/F = outros (marcados; D vazio, F terceiros)

Completude (0-100 -> rico>=70 / medio / pobre), §5 do SPEC:
  grupo% (25) · PU+un% (25) · data confiável (20) · composição (15) · formato A/B-E/C (15)

Uso:
    python run_selftest.py "<pasta de custo>"
    python run_selftest.py "<pasta>" --slug cincatarina --ac 655.19
    python run_selftest.py --anchors      # roda as 3 pastas-âncora do SPEC
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from datetime import datetime, timezone, timedelta
from pathlib import Path

_HERE = Path(__file__).resolve().parent
if str(_HERE) not in sys.path:
    sys.path.insert(0, str(_HERE))

import discovery_custos as disc
import parse_custos as pc
import resolve_data_preco as rdp

BRT = timezone(timedelta(hours=-3))
STAGING = Path(r"c:\Users\leona\openclaw\_local\base-custos-historica\staging")

# Pastas-âncora do SPEC (Fase 2)
ANCHORS = {
    "cincatarina": {
        "folder": r"G:\Drives compartilhados\03 CTN Projetos\2. Projetos em Andamento\Cincatarina\04. Custo",
        "tipo_esperado": "A",
    },
    "gdi-santa-monica": {
        "folder": r"G:\Drives compartilhados\03 CTN Projetos\2. Projetos em Andamento\GDI Empreendimentos\Santa Mônica\04. Custo",
        "tipo_esperado": "E",
    },
    "tecverde-casa-c4e": {
        "folder": r"G:\Drives compartilhados\03 CTN Projetos\2. Projetos em Andamento\TecVerde\04. Custo",
        "tipo_esperado": "C",
    },
}


def _slug_from_folder(folder: Path) -> str:
    """Deriva slug de '.../<Cliente>/[<Obra>/]04. Custo'."""
    parts = [p for p in folder.parts]
    # remover o '04. Custo' / '03. Custo' final
    idx = len(parts) - 1
    if re.match(r"^\d+\.\s*custo", parts[idx], re.IGNORECASE):
        idx -= 1
    # cliente = parte antes da pasta-mãe 'Projetos em Andamento'
    try:
        anchor = next(i for i, p in enumerate(parts) if "projetos em andamento" in p.lower())
        comps = parts[anchor + 1: idx + 1]
    except StopIteration:
        comps = parts[max(0, idx - 1): idx + 1]
    s = "-".join(comps)
    s = unicodedata.normalize("NFKD", s.lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s or "projeto"


def classify_folder_tipo(manifest: dict, parsed: dict) -> tuple[str, str]:
    """Classifica A-F deterministicamente. Devolve (tipo, motivo).
    Claude refina os casos limítrofes na fase real (sobre manifesto+amostra)."""
    counts = manifest.get("counts", {})
    n_wb = counts.get("workbook_orcamento", 0)
    n_comp = counts.get("composicao", 0)
    n_pdf = counts.get("relatorio_pdf", 0)
    modelo = parsed["projeto"].get("modelo")
    pick_wb = manifest.get("picks", {}).get("workbook_orcamento")
    pick_ext = Path(pick_wb).suffix.lower() if pick_wb else None

    n_files = len(manifest.get("files", []))
    n_cot = counts.get("cotacao", 0)
    n_eap = counts.get("eap", 0)
    has_files = n_files > 0
    if not has_files:
        return "D", "pasta sem arquivos de orçamento (vazia)"

    # Tipo C: o pick de orçamento é PDF (export AltoQi)
    if pick_ext == ".pdf" or modelo in ("pdf_altoqi", "pdf_composicoes"):
        return "C", "orçamento principal em PDF (export AltoQi/SIENGE)"

    # A vs E (modelo Cartesian Ger_Executivo): A = pasta rica (vários arquivos de
    # apoio — cotações, EAP, composições, relatórios); E = 1 xlsx auto-contido.
    # Discriminador: presença de arquivos de apoio separados, não a aba CPU interna.
    if modelo == "ger_executivo":
        rica = (n_comp >= 1 or n_cot >= 1 or n_eap >= 1 or n_pdf >= 3 or n_files >= 5)
        if rica:
            return "A", (f"Ger_Executivo + pasta rica (files={n_files}, cot={n_cot}, "
                         f"eap={n_eap}, comp={n_comp}, pdf={n_pdf})")
        return "E", f"Ger_Executivo auto-contido (files={n_files}, sem apoio separado)"

    # Tipo B: xlsx mas não modelo Cartesian (outro layout reconhecido)
    if n_wb >= 1 and pick_ext in (".xlsx", ".xlsm"):
        return "B", "xlsx de orçamento não-Cartesian (fallback)"

    # Tipo F: só xls/xlsb antigo (provável terceiro/legado)
    if pick_ext in (".xls", ".xlsb"):
        return "F", "arquivo legado .xls/.xlsb (terceiros/legado)"

    return "B", "indeterminado — revisar"


def _reconcile_summary(itens: list, tol: float = 0.02) -> dict:
    """Saúde da extração: quanto de soma(pu*qtd) bate com soma(total).

    Retorna {pct, n, ok, bad}. pct = 100*soma(pu*qtd)/soma(total) (alvo >=95%).
    Itens sem o triple completo (pu/qtd/total) não entram na conta."""
    td = tc = 0.0
    n = ok = bad = 0
    for it in itens:
        pu, qtd, total = it.get("pu"), it.get("qtd"), it.get("total")
        if pu is None or qtd is None or total is None:
            continue
        n += 1
        td += total
        tc += pu * qtd
        if abs(pu * qtd - total) <= tol * max(1.0, abs(total)):
            ok += 1
        else:
            bad += 1
    pct = round(100 * tc / td, 2) if td else None
    return {"pct": pct, "n": n, "ok": ok, "bad": bad}


def score_completude(parsed: dict, data: dict, folder_tipo: str) -> dict:
    """Pontua 0-100 conforme §5. Retorna {score, tier, breakdown}."""
    itens = parsed.get("itens", [])
    n = len(itens)
    bd = {}

    # grupo% (25): fração de itens com macrogrupo OU código hierárquico
    if n:
        com_grupo = sum(1 for i in itens if i.get("macrogrupo") or pc._is_codigo_hier(i.get("codigo")))
        bd["grupo"] = round(25 * com_grupo / n, 1)
    else:
        bd["grupo"] = 0.0

    # PU+un% (25): fração de itens com pu E unidade
    if n:
        com_pu_un = sum(1 for i in itens if i.get("pu") and i.get("unidade"))
        bd["pu_un"] = round(25 * com_pu_un / n, 1)
    else:
        bd["pu_un"] = 0.0

    # data confiável (20): alta=20, media=12, baixa=5, nenhuma=0
    conf = data.get("confianca")
    bd["data"] = {"alta": 20, "media": 12, "baixa": 5, "nenhuma": 0}.get(conf, 0)

    # composição (15): tem insumos?
    n_insumos = len(parsed.get("insumos", []))
    bd["composicao"] = 15 if n_insumos > 0 else 0

    # formato (15): A=15, B/E=10, C=8, demais=3
    bd["formato"] = {"A": 15, "B": 10, "E": 10, "C": 8}.get(folder_tipo, 3)

    score = round(sum(bd.values()), 1)
    tier = "rico" if score >= 70 else ("medio" if score >= 40 else "pobre")
    return {"score": score, "tier": tier, "breakdown": bd,
            "n_itens": n, "n_insumos": n_insumos}


# Abas que carregam data-base / CUB / metadados (Cincatarina=CAPA C7/G26;
# GDI=BASES B11). Buscar essas PRIMEIRO e fundo (a data fica longe da linha 1).
CAPA_SHEET_KW = ["capa", "bases", "dados", "obra", "resumo", "ger_exec",
                 "gerenciamento", "orçamento", "orcamento"]


def _xlsx_capa_text(p: Path, max_chars: int) -> str:
    """Lê texto das abas de metadados de um xlsx (CAPA/BASES/DADOS), fundo o bastante
    pra alcançar a célula 'Data base' (ex.: CAPA!C7, BASES!B11, CUB G26)."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(str(p), read_only=True, data_only=True)
    except Exception:
        return ""
    out = []
    # priorizar abas de metadados, depois o resto (limitado)
    def _rank(sn):
        n = sn.lower()
        for i, kw in enumerate(CAPA_SHEET_KW):
            if kw in n:
                return i
        return 99
    sheets = sorted(wb.sheetnames, key=_rank)
    for sn in sheets[:6]:
        ws = wb[sn]
        out.append(f"=== {sn} ===")
        for row in ws.iter_rows(max_row=40, max_col=12, values_only=True):
            line = " | ".join(str(c) for c in row if c not in (None, ""))
            if line:
                out.append(line)
            if sum(len(x) for x in out) > max_chars:
                break
        if sum(len(x) for x in out) > max_chars:
            break
    wb.close()
    return "\n".join(out)[:max_chars]


def _read_capa_text(picks: dict, max_chars: int = 8000) -> str:
    """Lê um blob de texto da capa do melhor arquivo (pra regex de data-base/CUB).

    Para xlsx usa _xlsx_capa_text (abas CAPA/BASES, scan fundo). Para PDF usa pdfplumber.
    Reusa extrair_dados_entregas.ler_capa pro caso PDF/xls antigo."""
    target = (picks.get("workbook_orcamento") or picks.get("relatorio_pdf")
              or picks.get("composicao"))
    if not target:
        return ""
    p = Path(target)
    ext = p.suffix.lower()

    if ext in (".xlsx", ".xlsm"):
        txt = _xlsx_capa_text(p, max_chars)
        if txt:
            return txt

    if ext == ".pdf":
        try:
            import pdfplumber
            out = []
            with pdfplumber.open(str(p)) as pdf:
                for pg in pdf.pages[:5]:
                    t = pg.extract_text()
                    if t:
                        out.append(t)
            return "\n".join(out)[:max_chars]
        except Exception:
            pass

    # fallback genérico (xls/xlsb via pandas, etc.)
    try:
        import extrair_dados_entregas as ede
        return ede.ler_capa(p)[:max_chars]
    except Exception:
        return ""


def run_one(folder: str, slug: str | None = None, ac: float | None = None,
            tipo_esperado: str | None = None) -> dict:
    folder_p = Path(folder)
    if not slug:
        slug = _slug_from_folder(folder_p)

    out_dir = STAGING / slug
    out_dir.mkdir(parents=True, exist_ok=True)

    # 1. discovery
    manifest = disc.discover(folder, compute_sha1=True)

    pick_wb = manifest.get("picks", {}).get("workbook_orcamento")
    pick_mtime = None
    for f in manifest.get("files", []):
        if f["path"] == pick_wb:
            pick_mtime = f.get("mtime")
            break

    # 2. parse do workbook/PDF principal
    if pick_wb:
        parsed = pc.parse_workbook(pick_wb, ac=ac, slug=slug)
    else:
        parsed = {"arquivo": None, "projeto": {"slug": slug, "ac": ac, "total": None,
                  "padrao": None, "macrogrupos": {}, "disciplinas": {}, "modelo": None},
                  "itens": [], "insumos": [], "ambiguities": [], "errors": ["sem workbook"]}

    # 2b. composições de PDF separado (Tipo C tem Composições.pdf)
    for comp_path in manifest.get("picks", {}).get("composicao_all", []):
        if comp_path == pick_wb:
            continue
        if Path(comp_path).suffix.lower() == ".pdf":
            try:
                extra = pc.parse_pdf_composicoes(Path(comp_path))
                parsed["insumos"].extend(extra)
            except Exception as e:
                parsed["errors"].append(f"pdf comp {Path(comp_path).name}: {e}")
        elif Path(comp_path).suffix.lower() in (".xlsx", ".xlsm") and not parsed["insumos"]:
            # composição em xlsx separado do orçamento principal
            try:
                from openpyxl import load_workbook
                wbc = load_workbook(comp_path, read_only=True, data_only=True)
                parsed["insumos"].extend(pc.parse_composicoes_xlsx(wbc))
                wbc.close()
            except Exception as e:
                parsed["errors"].append(f"xlsx comp {Path(comp_path).name}: {e}")

    # 3. data-base
    capa_txt = _read_capa_text(manifest.get("picks", {}))
    data = rdp.resolve_data_preco(
        texto_capa=capa_txt,
        folder_path=folder,
        file_name=Path(pick_wb).name if pick_wb else None,
        file_mtime_iso=pick_mtime,
    )

    # 4. classificação + completude
    folder_tipo, tipo_motivo = classify_folder_tipo(manifest, parsed)
    completude = score_completude(parsed, data, folder_tipo)

    # 4b. reconciliação pu*qtd vs total (saúde da extração; alvo SPEC >=95%)
    reconcile = _reconcile_summary(parsed["itens"])

    # status do projeto (ex.: pendente-ocr de PDF Tipo C/E sem tabela extraível)
    status = parsed["projeto"].get("status")

    # ------- gravar staging -------
    proj = dict(parsed["projeto"])
    proj.update({
        "slug": slug,
        "folder": folder,
        "folder_tipo": folder_tipo,
        "data_base": data.get("data_base"),
        "data_base_origem": data.get("origem"),
        "data_base_confianca": data.get("confianca"),
        "data_base_ambiguo": data.get("ambiguo"),
        "vintage": data.get("vintage"),
        "revisao": data.get("revisao"),
        "completude": completude,
    })
    (out_dir / "projeto.json").write_text(
        json.dumps(proj, ensure_ascii=False, indent=2), encoding="utf-8")

    # itens.jsonl — carimba data + fonte por item
    src_sha1 = None
    for f in manifest.get("files", []):
        if f["path"] == pick_wb:
            src_sha1 = f.get("sha1")
            break
    with (out_dir / "itens.jsonl").open("w", encoding="utf-8") as fh:
        for row_i, it in enumerate(parsed["itens"]):
            rec = dict(it)
            rec.update({
                "slug": slug,
                "source_sha1": src_sha1,
                "source_sheet": it.get("aba"),
                "source_row": row_i,
                "data_base": data.get("data_base"),
                "data_base_origem": data.get("origem"),
                "data_base_confianca": data.get("confianca"),
                "vintage": data.get("vintage"),
                "folder_tipo": folder_tipo,
                "revisao": data.get("revisao"),
            })
            fh.write(json.dumps(rec, ensure_ascii=False) + "\n")

    # insumos.jsonl
    with (out_dir / "insumos.jsonl").open("w", encoding="utf-8") as fh:
        for ix in parsed["insumos"]:
            rec = dict(ix)
            rec.update({
                "slug": slug,
                "data_base": data.get("data_base"),
                "vintage": data.get("vintage"),
            })
            fh.write(json.dumps(rec, ensure_ascii=False) + "\n")

    # manifest.json (discovery + tudo)
    full_manifest = {
        "slug": slug,
        "folder": folder,
        "scanned_at": datetime.now(BRT).isoformat(timespec="seconds"),
        "folder_tipo": folder_tipo,
        "folder_tipo_motivo": tipo_motivo,
        "tipo_esperado": tipo_esperado,
        "tipo_match": (tipo_esperado is None or tipo_esperado == folder_tipo),
        "modelo_parser": parsed["projeto"].get("modelo"),
        "status": status,
        "status_motivo": parsed["projeto"].get("status_motivo"),
        "ac": proj.get("ac"),
        "total": proj.get("total"),
        "data_base": data,
        "completude": completude,
        "reconcile": reconcile,
        "repairs": parsed.get("repairs", []),
        "counts": {
            "files": len(manifest.get("files", [])),
            "itens": len(parsed["itens"]),
            "insumos": len(parsed["insumos"]),
            "ambiguities": len(parsed.get("ambiguities", [])),
            "errors": len(parsed.get("errors", [])),
            "por_papel": manifest.get("counts", {}),
        },
        "picks": manifest.get("picks", {}),
        "discovery_files": manifest.get("files", []),
        "ambiguities": parsed.get("ambiguities", []),
        "errors": parsed.get("errors", []) + manifest.get("errors", []),
    }
    (out_dir / "manifest.json").write_text(
        json.dumps(full_manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    return full_manifest


def _print_summary(m: dict):
    print(f"\n=== {m['slug']} ===")
    print(f"  Pasta: {m['folder']}")
    print(f"  folder_tipo: {m['folder_tipo']} ({m['folder_tipo_motivo']})"
          + (f"  [esperado {m['tipo_esperado']} -> {'OK' if m['tipo_match'] else 'DIVERGE'}]"
             if m.get('tipo_esperado') else ""))
    print(f"  Modelo parser: {m['modelo_parser']}  | AC: {m['ac']}  | Total: {m['total']}")
    if m.get("status"):
        print(f"  STATUS: {m['status']} — {m.get('status_motivo')}")
    rc = m.get("reconcile") or {}
    if rc.get("n"):
        print(f"  Reconcile pu*qtd vs total: {rc.get('pct')}%  (ok={rc['ok']}/{rc['n']}, bad={rc['bad']})"
              + (f"  | reparos={len(m.get('repairs', []))}" if m.get('repairs') else ""))
    db = m["data_base"]
    print(f"  Data-base: {db['data_base']} (origem={db['origem']}, conf={db['confianca']}, "
          f"ambiguo={db['ambiguo']}, vintage={db['vintage']})")
    cp = m["completude"]
    print(f"  Completude: {cp['score']} ({cp['tier']})  {cp['breakdown']}")
    c = m["counts"]
    print(f"  Arquivos={c['files']}  itens={c['itens']}  insumos={c['insumos']}  "
          f"ambig={c['ambiguities']}  erros={c['errors']}")
    print(f"  Por papel: {c['por_papel']}")
    if m["errors"]:
        for e in m["errors"][:5]:
            print(f"    ERRO: {e}")


def main():
    ap = argparse.ArgumentParser(description="run_selftest — discovery+parse+data -> staging")
    ap.add_argument("folder", nargs="?", help="caminho da pasta de custo")
    ap.add_argument("--slug", default=None)
    ap.add_argument("--ac", type=float, default=None)
    ap.add_argument("--tipo-esperado", default=None)
    ap.add_argument("--anchors", action="store_true", help="rodar as 3 pastas-âncora do SPEC")
    args = ap.parse_args()

    print(f"Staging: {STAGING}")

    if args.anchors:
        results = []
        for slug, info in ANCHORS.items():
            try:
                m = run_one(info["folder"], slug=slug, tipo_esperado=info["tipo_esperado"])
                results.append(m)
                _print_summary(m)
            except Exception as e:
                import traceback
                print(f"\n=== {slug} === FALHOU: {type(e).__name__}: {e}")
                print(traceback.format_exc()[-800:])
        print("\n--- RESUMO ÂNCORAS ---")
        for m in results:
            print(f"  {m['slug']:<22} tipo={m['folder_tipo']} "
                  f"itens={m['counts']['itens']:>5} insumos={m['counts']['insumos']:>5} "
                  f"data={m['data_base']['data_base']} compl={m['completude']['score']}")
        return

    if not args.folder:
        ap.error("informe a pasta de custo ou use --anchors")

    m = run_one(args.folder, slug=args.slug, ac=args.ac, tipo_esperado=args.tipo_esperado)
    _print_summary(m)
    print(f"\nStaging gravado em: {STAGING / m['slug']}")


if __name__ == "__main__":
    main()
